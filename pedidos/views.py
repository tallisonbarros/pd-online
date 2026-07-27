import csv
import json
import math
import unicodedata
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation
from urllib.error import HTTPError, URLError
from urllib.parse import quote, urlencode
from urllib.request import Request, urlopen

from django.contrib.admin.views.decorators import staff_member_required
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.conf import settings
from django.core.cache import cache
from django.core.files.base import ContentFile
from django.core import signing
from django.db import transaction
from django.db.models import Count, Max, Q, Sum
from django.db.models.functions import ExtractHour, TruncDate
from django.http import HttpResponse, HttpResponseBadRequest, JsonResponse
from django.urls import reverse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.dateparse import parse_date, parse_datetime, parse_time
from django.utils.timesince import timesince
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_GET, require_POST
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .api_serializers import serialize_pedido_api, serialize_pedido_summary_api
from .contabil_services import (
    atualizar_movimentacao_caixa,
    atualizar_movimentacao_conta,
    banco_padrao,
    caixa_diario_context,
    conta_diario_context,
    criar_movimentacao_caixa,
    criar_movimentacao_conta,
    duplicar_movimentacao_caixa,
    duplicar_movimentacao_conta,
    excluir_movimentacao_caixa,
    excluir_movimentacao_caixa_pedido,
    excluir_movimentacao_conta,
    money_label,
    signed_money_label,
    sync_movimentacao_caixa_pedido,
    terminal_padrao,
)
from .dashboard import get_dashboard_diaria
from .forms import AdicionalForm, BebidaForm, PratoForm, PratoProntoForm
from .image_optimization import optimized_menu_image_url
from .legacy_import import build_legacy_import_preview, import_clean_legacy_orders, money_decimal as legacy_money_decimal
from .models import AccessEvent, Adicional, BancoConta, Bebida, CategoriaMovimentacaoCaixa, CategoriaMovimentacaoConta, Cliente, ClienteTokenConflito, ConfiguracaoEntrega, Cupom, DataFechada, DisponibilidadeTurnoDia, DisponibilidadeTurnoItem, EnderecoCliente, FaixaFrete, ItemPedido, MovimentacaoCaixa, MovimentacaoConta, Pedido, PedidoApiKey, PedidoListaImpressao, Prato, ResumoOperacionalDia, TerminalCaixa, TurnoAtendimento, TurnoPratoPadrao
from .order_services import (
    create_order_items_from_payload,
    inherit_customer_from_known_tokens,
    money_decimal,
    normalize_coupon_code,
    normalize_phone,
    recalculate_order_totals,
    reprice_order_items_from_catalog,
    replace_order_items,
    serialize_editor_catalog,
    sync_customer_from_order,
    validar_cupom,
)
from .utils import closed_dates_map, next_open_date, relative_day_label
from .turno_services import (
    WEEKDAYS as TURNO_WEEKDAYS,
    ensure_disponibilidade_turno,
    ensure_turnos_padrao,
    liberar_estoque_pedido,
    preco_turno_prato,
    resolve_turno_publico,
    sincronizar_disponibilidade_automatica,
    turno_agendado_no_dia,
    turno_label_pedido,
    validar_turno_para_pedido,
)

RECURRENT_CUSTOMER_EXCLUDED_STATUSES = [Pedido.Status.RASCUNHO, Pedido.Status.CANCELADO]

WEEKDAYS = ["seg", "ter", "qua", "qui", "sex", "sab", "dom"]
ORDER_HISTORY_COOKIE = "prato_delivery_orders"
ORDER_HISTORY_COOKIE_MAX_AGE = 60 * 60 * 24 * 180
WEEKDAY_LABELS = {
    "seg": "SEGUNDA",
    "ter": "TERÇA",
    "qua": "QUARTA",
    "qui": "QUINTA",
    "sex": "SEXTA",
    "sab": "SABADO",
    "dom": "DOMINGO",
}
OVERPASS_API_URL = "https://overpass-api.de/api/interpreter"
RIO_VERDE_CENTER = {"lat": -17.7923, "lon": -50.9192}
RIO_VERDE_BBOX = "-51.0500,-17.9500,-50.7500,-17.6500"  # minLon,minLat,maxLon,maxLat
OSRM_ROUTE_BASE_URL = "https://router.project-osrm.org/route/v1/driving/"
ATENDENTE_GROUP_NAME = "Atendente"
DIRETOR_GROUP_NAME = "Diretor"
GERENTE_GROUP_NAME = "Gerente"


def healthz(request):
    return HttpResponse("ok", content_type="text/plain")

RIO_VERDE_BAIRROS_OFICIAIS = [
    "Anhanguera", "Area Rural de Rio Verde", "Cesar Bastos", "Ceu Azul", "Cidade Empresarial Nova Alianca",
    "Conjunto Mauricio Arantes", "Conjunto Morada do Sol", "Conjunto Vila Verde", "Distrito Agroindustrial (DARV)",
    "Eldorado", "Eldorado Prolongamento", "Jardim Adriana", "Jardim America", "Jardim Bela Vista", "Jardim Brasilia",
    "Jardim Cruvinel", "Jardim das Margaridas", "Jardim Diniz", "Jardim Eleonora", "Jardim Floresta", "Jardim Goias",
    "Jardim Marconal", "Jardim Mondale", "Jardim Neves", "Jardim Presidente", "Jardim Sao Tomaz", "Liberdade",
    "Lindolfina", "Loteamento Gameleira", "Maristela", "Martins", "Medeiros", "Nova Vila Maria", "Odilia",
    "Paraguassu", "Parque Bandeirante", "Parque Betel", "Parque das Acacias", "Parque das Laranjeiras",
    "Parque das Paineiras", "Parque Dom Miguel", "Parque dos Buritis", "Parque dos Girassois", "Parque dos Jatobas",
    "Popular", "Presidente Nasser", "Primavera", "Residencial Agua Santa", "Residencial Araguaia",
    "Residencial Arco Iris", "Residencial Atalaia", "Residencial Canaa", "Residencial Dona Iza",
    "Residencial dos Buritis", "Residencial Gameleira", "Residencial Green Park", "Residencial Interlagos",
    "Residencial Jardim Campestre", "Residencial Jardim Helena", "Residencial Maranata",
    "Residencial Nilson Veloso", "Residencial Parque dos Ipes", "Residencial Recanto do Bosque",
    "Residencial Solar dos Ataides", "Residencial Tocantins", "Residencial Veneza",
    "Residencial Villagio Terra Cota", "Santo Agostinho", "Santo Antonio de Lisboa", "Sao Felipe", "Sao Joao",
    "Sao Joaquim", "Serra Dourada", "Setor Alvorada", "Setor Central", "Setor Dona Gercina",
    "Setor dos Funcionarios", "Setor Industrial", "Setor Morada do Sol", "Setor Oeste", "Setor Pauzanes",
    "Setor Santa Luzia", "Setor Universitario", "Solar Campestre", "Solar Monte Siao", "Vila Amalia",
    "Vila Andre Luiz", "Vila Baylao", "Vila Borges", "Vila Carolina", "Vila Dinara", "Vila Dona Auta",
    "Vila Gomes", "Vila Maria", "Vila Mariana", "Vila Meneses", "Vila Miafiori", "Vila Modelo", "Vila Morais",
    "Vila Mutirao", "Vila Olinda", "Vila Promissao", "Vila Renovacao", "Vila Rocha", "Vila Rosalina",
    "Vila Santa Barbara", "Vila Santa Cruz", "Vila Santo Andre", "Vila Santo Antonio", "Vila Serpro",
    "Vitoria Regia",
]


def _prato_dias_disponiveis(prato):
    if not prato.dias_disponiveis.strip():
        return set(WEEKDAYS)
    return {dia.strip().lower() for dia in prato.dias_disponiveis.split(",") if dia.strip()}


def prato_disponível_no_dia(prato, weekday_key):
    return weekday_key in _prato_dias_disponiveis(prato)


def prato_disponível_hoje(prato):
    hoje = WEEKDAYS[timezone.localtime().weekday()]
    return prato_disponível_no_dia(prato, hoje)


def _resolve_cardapio_pratos(config=None, now=None):
    config = config or ConfiguracaoEntrega.get_solo()
    current = now or timezone.localtime()
    fechamento = getattr(config, "horario_fechamento", None)
    start_offset = 1 if fechamento and current.time() >= fechamento else 0
    current_date = current.date()
    closed_dates = closed_dates_map(current_date + timedelta(days=start_offset), days=14)
    active_pratos = list(Prato.objects.filter(ativo=True, exclusivo_prato_pronto=False))

    for offset in range(start_offset, start_offset + 7):
        target_date = current_date + timedelta(days=offset)
        if target_date in closed_dates:
            continue
        weekday_index = target_date.weekday()
        weekday_key = WEEKDAYS[weekday_index]
        pratos = [prato for prato in active_pratos if prato_disponível_no_dia(prato, weekday_key)]
        if pratos:
            is_today = offset == 0
            return {
                "pratos": pratos,
                "weekday_key": weekday_key,
                "is_today": is_today,
                "day_offset": offset,
                "target_date": target_date,
                "title_lines": ["PRATO", "DO DIA"] if is_today else ["PRATO", "DE", WEEKDAY_LABELS[weekday_key]],
                "empty_label": "hoje" if is_today else f"para {WEEKDAY_LABELS[weekday_key].lower()}",
            }

    target_date = next_open_date(current_date + timedelta(days=start_offset), closed_dates=closed_dates)
    day_offset = (target_date - current_date).days
    weekday_key = WEEKDAYS[target_date.weekday()]
    return {
        "pratos": [],
        "weekday_key": weekday_key,
        "is_today": day_offset == 0,
        "day_offset": day_offset,
        "target_date": target_date,
        "title_lines": ["PRATO", "DO DIA"] if day_offset == 0 else ["PRATO", "DE", WEEKDAY_LABELS[weekday_key]],
        "empty_label": "hoje" if day_offset == 0 else f"para {WEEKDAY_LABELS[weekday_key].lower()}",
    }


def _cardapio_status_tag(config, cardapio_context, turno_context=None, now=None):
    abertura = getattr(config, "horario_abertura", None)
    fechamento = getattr(config, "horario_fechamento", None)
    current = now or timezone.localtime()
    current_time = current.time()
    current_date = current.date()

    if turno_context and turno_context.get("is_open"):
        fim_prato_pronto = turno_context.get("fim")
        return {
            "label": "Prato Pronto",
            "time": (
                f"Aberto até {fim_prato_pronto.strftime('%H:%M')}"
                if fim_prato_pronto
                else "Aberto agora"
            ),
        }

    if turno_context and turno_context.get("owns_menu_context") and turno_context.get("inicio"):
        inicio_prato_pronto = turno_context["inicio"]
        fim_prato_pronto = turno_context.get("fim")
        return {
            "label": "Prato Pronto",
            "time": (
                f"{inicio_prato_pronto.strftime('%H:%M')} às {fim_prato_pronto.strftime('%H:%M')}"
                if fim_prato_pronto
                else inicio_prato_pronto.strftime("%H:%M")
            ),
        }

    if not abertura or not fechamento:
        return None

    opening_range = f"{abertura.strftime('%H:%M')} às {fechamento.strftime('%H:%M')}"
    is_today = bool(cardapio_context.get("is_today"))
    day_offset = int(cardapio_context.get("day_offset") or 0)
    target_date = cardapio_context.get("target_date") or (current_date + timedelta(days=day_offset))
    weekday_key = cardapio_context.get("weekday_key")
    weekday_label = WEEKDAY_LABELS.get(weekday_key, "HOJE")
    closed_dates = closed_dates_map(current_date, days=max(day_offset + 1, 2))

    if is_today and abertura <= current_time < fechamento:
        return {"label": "Aberto agora", "time": f"até {fechamento.strftime('%H:%M')}"}
    if is_today and current_time < abertura:
        return {"label": "Abre hoje", "time": opening_range}
    if day_offset == 1:
        return {"label": "Amanhã", "time": opening_range}
    return {"label": weekday_label, "time": opening_range}


def _cardapio_closed_notice(config, cardapio_context, now=None):
    abertura = getattr(config, "horario_abertura", None)
    fechamento = getattr(config, "horario_fechamento", None)
    if not abertura or not fechamento:
        return None

    current = now or timezone.localtime()
    current_time = current.time()
    current_date = current.date()
    day_offset = int(cardapio_context.get("day_offset") or 0)
    target_date = cardapio_context.get("target_date") or (current_date + timedelta(days=day_offset))
    closed_dates = closed_dates_map(current_date, days=max(day_offset + 1, 2))
    tomorrow = current_date + timedelta(days=1)

    closed_date = None
    if current_date in closed_dates:
        closed_date = current_date
    elif current_time >= fechamento and tomorrow in closed_dates and target_date > tomorrow:
        closed_date = tomorrow
    if not closed_date:
        return None

    return {
        "message": f"Fechado {relative_day_label(closed_date, current_date)}"
    }


def _cart_closed_notice(config=None, now=None):
    config = config or ConfiguracaoEntrega.get_solo()
    abertura = getattr(config, "horario_abertura", None)
    fechamento = getattr(config, "horario_fechamento", None)
    if not abertura or not fechamento:
        return None
    current = now or timezone.localtime()
    current_time = current.time()
    current_date = current.date()
    closed_dates = closed_dates_map(current_date, days=14)
    is_open = abertura <= current_time < fechamento and current_date not in closed_dates
    if is_open:
        return None
    opening_date = current_date if current_date in closed_dates or current_time < abertura else current_date + timedelta(days=1)
    closed_reference_date = opening_date if opening_date in closed_dates else None
    opening_date = next_open_date(opening_date, closed_dates=closed_dates)
    opening_day = relative_day_label(opening_date, current_date)
    if closed_reference_date:
        closed_day = relative_day_label(closed_reference_date, current_date)
        message = f"Pedido antecipado: {closed_day} estaremos fechados. Nossa equipe revisa {opening_day} às {abertura.strftime('%H:%M')} da manhã."
    else:
        message = f"Pedido antecipado: finalize com calma, e nossa equipe revisa {opening_day} às {abertura.strftime('%H:%M')} da manhã."
    return {
        "opening_time": abertura.strftime("%H:%M"),
        "opening_day": opening_day,
        "message": message,
    }


def _menu_image_fallback(image_field):
    placeholder = settings.STATIC_URL + "img/placeholder-prato.svg"
    if not image_field:
        return placeholder
    try:
        if not Path(image_field.path).exists():
            return placeholder
    except (NotImplementedError, ValueError):
        return placeholder
    return image_field.url


def serializar_prato(prato, disponibilidade_item=None):
    preco = disponibilidade_item.preco if disponibilidade_item else prato.preco_site_resolvido
    fallback_image = _menu_image_fallback(prato.imagem)
    payload = {
        "id": prato.id,
        "nome": prato.nome,
        "descricao": prato.descricao,
        "variacoes": prato.variacoes,
        "preco": f"{preco:.2f}" if preco is not None else "",
        "preco_formatado": f"R$ {preco:.2f}".replace(".", ",") if preco is not None else "",
        "imagem": optimized_menu_image_url(prato.imagem, fallback_image),
    }
    if disponibilidade_item:
        payload.update(
            {
                "disponibilidade_turno_item_id": disponibilidade_item.id,
                "quantidade_restante": disponibilidade_item.quantidade_restante,
                "esgotado": disponibilidade_item.esgotado,
            }
        )
    return payload


def serializar_bebida(bebida):
    preco = bebida.preco_site_resolvido
    fallback_image = _menu_image_fallback(bebida.imagem)
    return {
        "id": bebida.id,
        "nome": bebida.nome,
        "descricao": bebida.descricao,
        "preco": f"{preco:.2f}" if preco is not None else "",
        "preco_formatado": f"R$ {preco:.2f}".replace(".", ",") if preco is not None else "",
        "imagem": optimized_menu_image_url(bebida.imagem, fallback_image),
    }


def serializar_adicional(adicional):
    preco = adicional.preco_site_resolvido
    fallback_image = _menu_image_fallback(adicional.imagem)
    return {
        "id": adicional.id,
        "nome": adicional.nome,
        "descricao": adicional.descricao,
        "preco": f"{preco:.2f}" if preco is not None else "",
        "preco_formatado": f"R$ {preco:.2f}".replace(".", ",") if preco is not None else "",
        "imagem": optimized_menu_image_url(adicional.imagem, fallback_image),
    }


def user_is_atendente(user):
    if not getattr(user, "is_authenticated", False):
        return False
    return user.groups.filter(name=ATENDENTE_GROUP_NAME).exists()


def user_is_gerente(user):
    if not getattr(user, "is_authenticated", False):
        return False
    return user.groups.filter(name=GERENTE_GROUP_NAME).exists()


def user_is_diretor(user):
    if not getattr(user, "is_authenticated", False):
        return False
    return user.groups.filter(name=DIRETOR_GROUP_NAME).exists()


def montar_mensagem_whatsapp(pedido):
    linhas = [
        "*PRATO-DELIVERY*",
        f"Pedido #{pedido.numero}",
        f"*Status:* {pedido.status_label_contextual}",
    ]
    if pedido.is_prato_pronto:
        linhas.append("*Produto:* Prato Pronto")
        if pedido.orientacao_turno_snapshot:
            linhas.append(f"*Conservação:* {pedido.orientacao_turno_snapshot}")
    linhas.extend(["", f"*Cliente:* {pedido.nome_cliente}"])
    if pedido.telefone:
        linhas.append(f"*Telefone:* {pedido.telefone}")
    linhas.append(f"*Endereço:* {pedido.endereco}")
    if pedido.lote_quadra:
        linhas.append(f"*Lote/Quadra:* {pedido.lote_quadra}")
    if pedido.complemento:
        linhas.append(f"*Complemento:* {pedido.complemento}")
    if pedido.ponto_referencia:
        linhas.append(f"*Ponto de referência:* {pedido.ponto_referencia}")
    linhas.extend(
        [
            f"*Talheres:* {'Sim' if pedido.enviar_talheres else 'Não'}",
            "",
            "*Itens:*",
        ]
    )
    for item in pedido.itens.all():
        nome_item = item.nome_prato_snapshot
        if item.variacao_nome_snapshot:
            nome_item = f"{nome_item} - {item.variacao_nome_snapshot}"
        if item.classificacao_saida != ItemPedido.ClassificacaoSaida.VENDIDA:
            nome_item = f"{nome_item} ({item.get_classificacao_saida_display()})"
        linhas.append(
            f"- {item.quantidade}x {nome_item} | R$ {item.subtotal:.2f}".replace(".", ",")
        )
        if item.observacao:
            linhas.append(f"  Obs: {item.observacao}")
    if pedido.observacao_geral:
        linhas.extend(["", f"*Observação geral:* {pedido.observacao_geral}"])
    linhas.extend(["", f"*Pagamento:* {pedido.get_forma_pagamento_display()}"])
    linhas.append(f"*{pedido.pagamento_copia_status}*")
    if pedido.forma_pagamento == Pedido.FormaPagamento.PIX:
        pix_chave = _safe_text(getattr(ConfiguracaoEntrega.get_solo(), "pix_chave", ""))
        if pix_chave:
            linhas.append(f"*Chave Pix:* {pix_chave}")
    linhas.extend(["", f"*Entrega:* {_pedido_frete_line_value(pedido)}"])
    if pedido.promocao_desconto and pedido.promocao_desconto > 0:
        descricao = pedido.promocao_descricao or "Promoção especial"
        linhas.append(f"*{descricao}:* - R$ {pedido.promocao_desconto:.2f}".replace(".", ","))
    if pedido.cupom_desconto and pedido.cupom_desconto > 0:
        linhas.append(f"*Desconto:* - R$ {pedido.cupom_desconto:.2f}".replace(".", ","))
    linhas.append(f"*Total:* R$ {pedido.total:.2f}".replace(".", ","))
    return "\n".join(linhas)


def montar_mensagem_entregador(pedido):
    linhas = [
        f"Pedido #{pedido.numero} - {pedido.nome_cliente}",
        f"Endereço: {pedido.endereco}",
    ]
    linhas.append(f"Pagamento: {pedido.get_forma_pagamento_display()}")
    status_pagamento = (
        f"COBRAR DO CLIENTE: {_money_line_value(pedido.total)}"
        if pedido.pagamento_na_entrega and not pedido.pagamento_recebido
        else pedido.pagamento_copia_status
    )
    linhas.append(f"*{status_pagamento}*")
    if pedido.lote_quadra:
        linhas.append(f"Lote/Quadra: {pedido.lote_quadra}")
    if pedido.complemento:
        linhas.append(f"Complemento: {pedido.complemento}")
    if pedido.ponto_referencia:
        linhas.append(f"Referência: {pedido.ponto_referencia}")
    if pedido.google_maps_route_url:
        linhas.append(f"Rota: {pedido.google_maps_route_url}")
        if pedido.canal in {Pedido.Canal.BALCAO, Pedido.Canal.IFOOD}:
            linhas.append("Rota aproximada, confira o endereco escrito.")
    return "\n".join(linhas)


def _pedido_confirmacao_entrega_line(pedido):
    if pedido.tipo_coleta == Pedido.TipoColeta.RETIRADA:
        return "Retirada no local"
    partes = [f"Endereco: {pedido.endereco}"]
    if pedido.lote_quadra:
        partes.append(f"Lote/Quadra: {pedido.lote_quadra}")
    if pedido.complemento:
        partes.append(f"Complemento: {pedido.complemento}")
    if pedido.ponto_referencia:
        partes.append(f"Referencia: {pedido.ponto_referencia}")
    return "\n".join(partes)


def _money_line_value(value):
    return f"R$ {value:.2f}".replace(".", ",")


def _pedido_frete_line_value(pedido):
    return "gratis" if pedido.frete_gratis else _money_line_value(pedido.valor_frete)


def _pedido_frete_display_value(pedido):
    return "Gratis" if pedido.frete_gratis else _money_line_value(pedido.valor_frete)


def _pedido_distancia_display_value(pedido):
    distancia = f"{pedido.distancia_km:.2f}".replace(".", ",")
    if pedido.frete_gratis and pedido.tipo_coleta == Pedido.TipoColeta.ENTREGA:
        return f"Entrega gratis - distancia {distancia} km"
    return f"{distancia} km"


def _post_bool(value):
    return _safe_text(value).lower() in {"1", "true", "on", "sim", "yes"}


def _pedido_confirmacao_status_final(pedido):
    if pedido.tipo_coleta == Pedido.TipoColeta.RETIRADA:
        return "Ja vamos iniciar o preparo e te avisamos por aqui quando estiver pronto para retirada."
    return "Ja vamos iniciar o preparo e te avisamos por aqui quando sair para entrega."


def _pedido_confirmacao_pagamento_line(pedido):
    return f"Pagamento: {pedido.pagamento_copia_status}"


def _pedido_confirmacao_item_unit_value(item):
    if item.quantidade:
        return _money_line_value(item.subtotal / Decimal(item.quantidade))
    return _money_line_value(item.subtotal)


def _pedido_confirmacao_desconto_lines(pedido):
    linhas = []
    if pedido.promocao_desconto and pedido.promocao_desconto > 0:
        descricao = pedido.promocao_descricao or "Promocao especial"
        linhas.append(f"{descricao}: - {_money_line_value(pedido.promocao_desconto)}")
    if pedido.cupom_desconto and pedido.cupom_desconto > 0:
        descricao = f"Cupom {pedido.cupom_codigo}" if pedido.cupom_codigo else "Cupom"
        linhas.append(f"{descricao}: - {_money_line_value(pedido.cupom_desconto)}")
    return linhas


def montar_mensagem_confirmacao_pedido(pedido):
    nome_cliente = _safe_text(pedido.nome_cliente)
    saudacao = f"Oi, {nome_cliente}!" if nome_cliente.casefold() not in {"", "cliente"} else "Oi!"
    if pedido.canal == Pedido.Canal.BALCAO:
        itens = list(pedido.itens.all())
        linhas = [
            f"Pedido #{pedido.numero} confirmado.",
            "",
            "Itens:",
        ]
        for item in itens:
            nome_item = item.nome_prato_snapshot
            if item.variacao_nome_snapshot:
                nome_item = f"{nome_item} - {item.variacao_nome_snapshot}"
            if item.classificacao_saida != ItemPedido.ClassificacaoSaida.VENDIDA:
                nome_item = f"{nome_item} ({item.get_classificacao_saida_display()})"
            item_line = f"- {item.quantidade}x {nome_item} | {_pedido_confirmacao_item_unit_value(item)} un."
            if item.quantidade > 1:
                item_line = f"{item_line} | {_money_line_value(item.subtotal)}"
            linhas.append(item_line)
            if item.observacao:
                linhas.append(f"  Obs: {item.observacao}")
        if not itens:
            linhas.append("- Sem itens")
        linhas.extend(
            [
                "",
                _pedido_confirmacao_entrega_line(pedido),
                "",
                f"Entrega: {_pedido_frete_line_value(pedido)}",
                f"Pagamento: {pedido.get_forma_pagamento_display()}",
            ]
        )
        linhas.extend(_pedido_confirmacao_desconto_lines(pedido))
        linhas.extend(
            [
                f"Total: {_money_line_value(pedido.total)}",
                _pedido_confirmacao_pagamento_line(pedido),
                "",
                _pedido_confirmacao_status_final(pedido),
            ]
        )
        return "\n".join(linhas)

    linhas = [
        saudacao,
        f"Recebemos seu pedido #{pedido.numero}.",
        "",
        _pedido_confirmacao_entrega_line(pedido),
        "",
        f"Pagamento: {pedido.get_forma_pagamento_display()}",
    ]
    linhas.extend(_pedido_confirmacao_desconto_lines(pedido))
    linhas.extend(
        [
            f"Total: {_money_line_value(pedido.total)}",
            _pedido_confirmacao_pagamento_line(pedido),
            "",
            _pedido_confirmacao_status_final(pedido),
        ]
    )
    return "\n".join(linhas)


def _normalize_whatsapp_number(value):
    return "".join(char for char in _safe_text(value) if char.isdigit())


def _configured_whatsapp_number(config=None):
    config = config or ConfiguracaoEntrega.objects.order_by("pk").first()
    numero = _normalize_whatsapp_number(getattr(config, "whatsapp_numero", ""))
    return numero or _normalize_whatsapp_number(getattr(settings, "RESTAURANT_WHATSAPP", ""))


def _build_whatsapp_order_url(pedido, config=None):
    numero = _configured_whatsapp_number(config)
    if not numero:
        return ""
    return f"https://wa.me/{numero}?text={quote(montar_mensagem_whatsapp(pedido))}"


def _access_session_key(request):
    if not request.session.session_key:
        request.session.create()
    return request.session.session_key or ""


def _safe_metric_decimal(value):
    try:
        return Decimal(str(value or "0").replace(",", ".")).quantize(Decimal("0.01"))
    except (InvalidOperation, TypeError, ValueError):
        return Decimal("0.00")


def _safe_metric_int(value):
    try:
        return max(0, int(value or 0))
    except (TypeError, ValueError):
        return 0


def _record_access_event(request, event_type, **extra):
    if event_type not in dict(AccessEvent.EventType.choices):
        return None
    event_path = _safe_text(extra.get("path") or request.path)[:160]
    session_key = _access_session_key(request)
    metadata = extra.get("metadata") if isinstance(extra.get("metadata"), dict) else {}
    allowed_metadata = {"origem", "variacao", "tipo_coleta", "forma_pagamento", "pedido_id", "page_open_id"}
    clean_metadata = {
        str(key)[:40]: str(value)[:160]
        for key, value in metadata.items()
        if key in allowed_metadata and value not in (None, "")
    }
    dedupe_metadata_key = _safe_text(extra.get("dedupe_metadata_key"))
    if dedupe_metadata_key and clean_metadata.get(dedupe_metadata_key):
        metadata_filter = {f"metadata__{dedupe_metadata_key}": clean_metadata[dedupe_metadata_key]}
        if AccessEvent.objects.filter(
            event_type=event_type,
            path=event_path,
            session_key=session_key,
            **metadata_filter,
        ).exists():
            return None
    dedupe_seconds = 0 if (dedupe_metadata_key and clean_metadata.get(dedupe_metadata_key)) else _safe_metric_int(extra.get("dedupe_seconds"))
    if dedupe_seconds:
        recent_since = timezone.now() - timedelta(seconds=dedupe_seconds)
        if AccessEvent.objects.filter(
            event_type=event_type,
            path=event_path,
            session_key=session_key,
            created_at__gte=recent_since,
        ).exists():
            return None
    return AccessEvent.objects.create(
        event_type=event_type,
        path=event_path,
        session_key=session_key,
        item_type=_safe_text(extra.get("item_type"))[:20],
        item_id=extra.get("item_id") if isinstance(extra.get("item_id"), int) and extra.get("item_id") > 0 else None,
        cart_items_count=_safe_metric_int(extra.get("cart_items_count")),
        cart_total=_safe_metric_decimal(extra.get("cart_total")),
        metadata=clean_metadata,
    )


@never_cache
def cardapio(request):
    config = ConfiguracaoEntrega.get_solo()
    current = timezone.localtime()
    whatsapp_numero = _configured_whatsapp_number(config)
    whatsapp_cardapio_url = (
        f"https://wa.me/{whatsapp_numero}?text={quote('Olá! Estou vendo o cardápio e queria tirar uma dúvida.')}"
        if whatsapp_numero
        else ""
    )
    regular_context = _resolve_cardapio_pratos(config=config, now=current)
    turno_context = resolve_turno_publico(current)
    cardapio_context = regular_context
    disponibilidade_itens = {}
    prato_pronto_contextual = turno_context["owns_menu_context"]
    if prato_pronto_contextual:
        itens_turno = (
            turno_context["itens_disponiveis"]
            if turno_context["accepting_orders"]
            else turno_context["itens"]
        )
        pratos = [item.prato for item in itens_turno]
        disponibilidade_itens = {item.prato_id: item for item in itens_turno}
        cardapio_context = {
            "pratos": pratos,
            "weekday_key": WEEKDAYS[current.weekday()],
            "is_today": True,
            "day_offset": 0,
            "target_date": current.date(),
            "title_lines": ["PRATO", "PRONTO"],
            "empty_label": "neste turno",
        }
    else:
        pratos = cardapio_context["pratos"]
    adicionais = (
        Adicional.objects.none()
        if prato_pronto_contextual
        else Adicional.objects.filter(ativo=True)
    )
    bebidas = Bebida.objects.filter(ativo=True)
    pratos_serializados = [
        serializar_prato(prato, disponibilidade_itens.get(prato.id))
        for prato in pratos
    ]
    adicionais_serializados = [serializar_adicional(adicional) for adicional in adicionais]
    bebidas_serializadas = [serializar_bebida(bebida) for bebida in bebidas]
    for item, payload in zip(pratos, pratos_serializados):
        item.imagem_cardapio_url = payload["imagem"]
        item.preco_cardapio = payload["preco"]
        item.preco_cardapio_formatado = payload["preco_formatado"]
        item.disponibilidade_turno_item_id = payload.get("disponibilidade_turno_item_id")
        item.quantidade_restante = payload.get("quantidade_restante")
    for item, payload in zip(adicionais, adicionais_serializados):
        item.imagem_cardapio_url = payload["imagem"]
        item.preco_cardapio = payload["preco"]
    for item, payload in zip(bebidas, bebidas_serializadas):
        item.imagem_cardapio_url = payload["imagem"]
        item.preco_cardapio = payload["preco"]
    turno_banner = None
    if prato_pronto_contextual:
        turno_banner = {
            "state": "open" if turno_context["is_open"] else "scheduled",
            "label": "Prato Pronto",
            "orientation": (
                "Marmitas seladas e refrigeradas, feitas no dia. "
                "Aqueça e sirva."
            ),
        }

    return render(
        request,
        "pedidos/cardapio.html",
        {
            "pratos": pratos,
            "adicionais": adicionais,
            "bebidas": bebidas,
            "pratos_json": json.dumps(pratos_serializados, ensure_ascii=False),
            "adicionais_json": json.dumps(adicionais_serializados, ensure_ascii=False),
            "bebidas_json": json.dumps(bebidas_serializadas, ensure_ascii=False),
              "cardapio_title_lines": cardapio_context["title_lines"],
              "cardapio_empty_label": cardapio_context["empty_label"],
              "cardapio_status_tag": _cardapio_status_tag(
                  config,
                  regular_context,
                  turno_context=turno_context,
                  now=current,
              ),
              "cardapio_closed_notice": _cardapio_closed_notice(config, regular_context, now=current),
              "horario_abertura": config.horario_abertura,
              "horario_fechamento": config.horario_fechamento,
              "whatsapp_cardapio_url": whatsapp_cardapio_url,
              "turno_prato_pronto_aberto": turno_context["accepting_orders"],
              "turno_prato_pronto_banner": turno_banner,
              "exibir_adicionais": not prato_pronto_contextual,
              "permite_promocao_marmita": (
                  turno_context["prato_pronto"].permite_promocao
                  if prato_pronto_contextual
                  else True
              ),
        },
    )


@never_cache
def checkout(request):
    config = ConfiguracaoEntrega.get_solo()
    turno_context = resolve_turno_publico()
    pratos_lookup = {f"prato:{prato.id}": {**serializar_prato(prato), "tipo": "prato"} for prato in Prato.objects.all()}
    adicionais_lookup = {
        f"adicional:{adicional.id}": {**serializar_adicional(adicional), "tipo": "adicional"}
        for adicional in Adicional.objects.all()
    }
    bebidas_lookup = {
        f"bebida:{bebida.id}": {**serializar_bebida(bebida), "tipo": "bebida"} for bebida in Bebida.objects.all()
    }
    itens_lookup = {**pratos_lookup, **adicionais_lookup, **bebidas_lookup}
    return render(
        request,
        "pedidos/checkout.html",
        {
            "pratos_lookup_json": itens_lookup,
            "pix_chave": config.pix_chave,
            "is_atendente": user_is_atendente(request.user),
            "bairros_sugestoes": RIO_VERDE_BAIRROS_OFICIAIS,
            "turno_prato_pronto_aberto": turno_context["accepting_orders"],
        },
    )


@never_cache
def carrinho(request):
    config = ConfiguracaoEntrega.get_solo()
    turno_context = resolve_turno_publico()
    pratos_lookup = {f"prato:{prato.id}": {**serializar_prato(prato), "tipo": "prato"} for prato in Prato.objects.filter(ativo=True)}
    return render(
        request,
        "pedidos/carrinho.html",
        {
            "cart_closed_notice": None if turno_context["accepting_orders"] else _cart_closed_notice(config),
            "pratos_lookup_json": pratos_lookup,
            "turno_prato_pronto_aberto": turno_context["accepting_orders"],
        },
    )


@require_POST
def api_metric_event(request):
    try:
        payload = json.loads(request.body.decode("utf-8") or "{}")
    except (json.JSONDecodeError, UnicodeDecodeError):
        payload = request.POST

    event_type = _safe_text(payload.get("event_type"))
    if event_type not in dict(AccessEvent.EventType.choices):
        return JsonResponse({"ok": False, "message": "Evento invalido."}, status=400)

    item_id = None
    try:
        raw_item_id = payload.get("item_id")
        item_id = int(raw_item_id) if raw_item_id not in (None, "") else None
    except (TypeError, ValueError):
        item_id = None

    metadata = payload.get("metadata") if isinstance(payload.get("metadata"), dict) else {}
    is_page_view_event = event_type in {
        AccessEvent.EventType.MENU_VIEW,
        AccessEvent.EventType.CART_VIEW,
        AccessEvent.EventType.CHECKOUT_VIEW,
    }
    event = _record_access_event(
        request,
        event_type,
        path=_safe_text(payload.get("path") or request.META.get("HTTP_REFERER") or request.path),
        item_type=_safe_text(payload.get("item_type")),
        item_id=item_id,
        cart_items_count=payload.get("cart_items_count"),
        cart_total=payload.get("cart_total"),
        metadata=metadata,
        dedupe_metadata_key="page_open_id" if is_page_view_event else "",
        dedupe_seconds=60 if is_page_view_event else 0,
    )
    return JsonResponse({"ok": True, "event_id": event.id if event else None, "duplicate": event is None})


def _safe_text(value):
    return str(value or "").strip()


def _safe_float(value):
    try:
        normalized = str(value or "").strip().replace(",", ".")
        if not normalized:
            return None
        return float(normalized)
    except (TypeError, ValueError):
        return None


def _parse_optional_time(value):
    raw = _safe_text(value)
    if not raw:
        return None
    parsed = parse_time(raw)
    if parsed is None:
        raise ValueError("Informe um horário válido no formato HH:MM.")
    return parsed


def _setting_float(name, default):
    try:
        value = float(getattr(settings, name, default))
        return value if value > 0 else float(default)
    except (TypeError, ValueError):
        return float(default)


def _setting_int(name, default):
    try:
        value = int(getattr(settings, name, default))
        return value if value >= 0 else int(default)
    except (TypeError, ValueError):
        return int(default)


DELIVERY_ETA_MULTIPLIER = _setting_float("DELIVERY_ETA_MULTIPLIER", 2.2)
DELIVERY_ETA_BUFFER_MINUTES = _setting_int("DELIVERY_ETA_BUFFER_MINUTES", 3)
DELIVERY_ETA_SHORT_TRIP_KM = _setting_float("DELIVERY_ETA_SHORT_TRIP_KM", 6.0)
DELIVERY_ETA_SHORT_TRIP_PENALTY_MINUTES = _setting_int("DELIVERY_ETA_SHORT_TRIP_PENALTY_MINUTES", 1)
DELIVERY_FRETE_PADRAO = Decimal("0.00")


def _build_order_address(rua="", numero="", bairro="", cidade="", estado="", lote_quadra=""):
    street_line = _safe_text(rua)
    numero = _safe_text(numero)
    bairro = _safe_text(bairro)
    cidade = _safe_text(cidade)
    estado = _safe_text(estado)
    lote_quadra = _safe_text(lote_quadra)

    if numero:
        street_line = f"{street_line}, {numero}" if street_line else numero
    if bairro:
        street_line = f"{street_line} - {bairro}" if street_line else bairro
    if lote_quadra:
        street_line = f"{street_line} - {lote_quadra}" if street_line else lote_quadra
    if cidade and estado and street_line:
        return f"{street_line}, {cidade} - {estado}"
    return street_line


def _load_bairros_coords_manual():
    path = Path(__file__).resolve().parent / "data" / "bairros_coords_manual.json"
    if not path.exists():
        return {}
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (ValueError, OSError):
        return {}
    if not isinstance(payload, dict):
        return {}
    data = {}
    for name, coords in payload.items():
        if not isinstance(coords, dict):
            continue
        lat = _safe_float(coords.get("lat"))
        lng = _safe_float(coords.get("lng"))
        if lat is None or lng is None:
            continue
        data[_normalize_key(name)] = {"bairro": _safe_text(name), "lat": lat, "lng": lng}
    return data


def _fetch_rio_verde_bairros():
    now = timezone.now()
    cache_updated_at = _RIO_VERDE_BAIRROS_CACHE.get("updated_at")
    if cache_updated_at and (now - cache_updated_at).total_seconds() < 21600:
        return _RIO_VERDE_BAIRROS_CACHE.get("data", [])

    min_lon, min_lat, max_lon, max_lat = [part.strip() for part in RIO_VERDE_BBOX.split(",")]
    query = f"""
[out:json][timeout:25];
(
  node["place"~"neighbourhood|suburb|quarter"]({min_lat},{min_lon},{max_lat},{max_lon});
  way["place"~"neighbourhood|suburb|quarter"]({min_lat},{min_lon},{max_lat},{max_lon});
  relation["place"~"neighbourhood|suburb|quarter"]({min_lat},{min_lon},{max_lat},{max_lon});
  relation["boundary"="administrative"]["admin_level"~"9|10"]({min_lat},{min_lon},{max_lat},{max_lon});
  way["boundary"="administrative"]["admin_level"~"9|10"]({min_lat},{min_lon},{max_lat},{max_lon});
);
out center;
"""
    payload = urlencode({"data": query})
    request = Request(
        f"{OVERPASS_API_URL}?{payload}",
        headers={"User-Agent": "PRATO-DELIVERY/1.0"},
    )
    bairros = []
    seen = set()
    try:
        with urlopen(request, timeout=12) as response:
            result = json.loads(response.read().decode("utf-8"))
    except (HTTPError, URLError, TimeoutError, ValueError, json.JSONDecodeError):
        return _RIO_VERDE_BAIRROS_CACHE.get("data", [])

    for item in result.get("elements", []):
        tags = item.get("tags") or {}
        nome = _safe_text(tags.get("name"))
        if not nome:
            continue
        # Ignore clearly non-neighborhood labels that may appear in admin boundaries.
        low = nome.lower()
        if low in {"rio verde", "goias", "goias", "brasil", "brazil"}:
            continue
        if len(nome) < 2:
            continue
        lat = item.get("lat")
        lng = item.get("lon")
        if lat is None or lng is None:
            center = item.get("center") or {}
            lat = center.get("lat")
            lng = center.get("lon")
        latf = _safe_float(lat)
        lngf = _safe_float(lng)
        if latf is None or lngf is None:
            continue
        key = nome.lower()
        if key in seen:
            continue
        seen.add(key)
        bairros.append({"bairro": nome, "lat": latf, "lng": lngf})

    bairros.sort(key=lambda row: row["bairro"].lower())
    by_norm = {_normalize_key(row["bairro"]): row for row in bairros if row.get("bairro")}
    manual = _load_bairros_coords_manual()

    final = []
    for bairro_nome in RIO_VERDE_BAIRROS_OFICIAIS:
        key = _normalize_key(bairro_nome)
        found = by_norm.get(key)
        if not found:
            found = manual.get(key)
        if found:
            final.append({"bairro": bairro_nome, "lat": found["lat"], "lng": found["lng"]})

    _RIO_VERDE_BAIRROS_CACHE["updated_at"] = now
    _RIO_VERDE_BAIRROS_CACHE["data"] = final
    return final


def _fetch_rio_verde_bairros_polygons():
    now = timezone.now()
    cache_updated_at = _RIO_VERDE_BAIRROS_POLYGONS_CACHE.get("updated_at")
    if cache_updated_at and (now - cache_updated_at).total_seconds() < 21600:
        return _RIO_VERDE_BAIRROS_POLYGONS_CACHE.get("data", {})

    min_lon, min_lat, max_lon, max_lat = [part.strip() for part in RIO_VERDE_BBOX.split(",")]
    # Fetch administrative/neighborhood boundaries from OSM (relations + ways).
    joined = "|".join(name.replace('"', '\\"') for name in RIO_VERDE_BAIRROS_OFICIAIS)
    query = f"""
[out:json][timeout:60];
(
  relation["name"~"^({joined})$",i]["boundary"="administrative"]({min_lat},{min_lon},{max_lat},{max_lon});
  relation["name"~"^({joined})$",i]["place"~"neighbourhood|suburb|quarter"]({min_lat},{min_lon},{max_lat},{max_lon});
  way["name"~"^({joined})$",i]["boundary"="administrative"]({min_lat},{min_lon},{max_lat},{max_lon});
  way["name"~"^({joined})$",i]["place"~"neighbourhood|suburb|quarter"]({min_lat},{min_lon},{max_lat},{max_lon});
  way["name"~"^({joined})$",i]["landuse"="residential"]({min_lat},{min_lon},{max_lat},{max_lon});
);
out geom;
"""
    body = f"data={query}".encode("utf-8")
    request = Request(
        OVERPASS_API_URL,
        data=body,
        headers={"User-Agent": "PRATO-DELIVERY/1.0", "Content-Type": "application/x-www-form-urlencoded"},
    )

    polygons = {}
    try:
        with urlopen(request, timeout=70) as response:
            payload = json.loads(response.read().decode("utf-8"))
    except (HTTPError, URLError, TimeoutError, ValueError, json.JSONDecodeError):
        return _RIO_VERDE_BAIRROS_POLYGONS_CACHE.get("data", {})

    official_by_norm = {_normalize_key(name): name for name in RIO_VERDE_BAIRROS_OFICIAIS}

    def _ring_from_geom(geom):
        ring = []
        for point in geom or []:
            lat = _safe_float(point.get("lat"))
            lng = _safe_float(point.get("lon"))
            if lat is None or lng is None:
                continue
            ring.append([lng, lat])
        if len(ring) < 4:
            return None
        if ring[0] != ring[-1]:
            ring.append(ring[0])
        return ring

    def _feature_from_way(item, official_name):
        ring = _ring_from_geom(item.get("geometry") or [])
        if not ring:
            return None
        return {
            "type": "Feature",
            "geometry": {"type": "Polygon", "coordinates": [ring]},
            "properties": {"bairro": official_name},
        }

    def _feature_from_relation(item, official_name):
        members = item.get("members") or []
        outers = []
        inners = []
        for member in members:
            geom = member.get("geometry") or []
            ring = _ring_from_geom(geom)
            if not ring:
                continue
            role = (member.get("role") or "").lower()
            if role == "inner":
                inners.append(ring)
            else:
                outers.append(ring)
        if not outers:
            return None
        polygons_coords = []
        for outer in outers:
            coords = [outer]
            coords.extend(inners)
            polygons_coords.append(coords)
        if len(polygons_coords) == 1:
            return {
                "type": "Feature",
                "geometry": {"type": "Polygon", "coordinates": polygons_coords[0]},
                "properties": {"bairro": official_name},
            }
        return {
            "type": "Feature",
            "geometry": {"type": "MultiPolygon", "coordinates": polygons_coords},
            "properties": {"bairro": official_name},
        }

    for item in payload.get("elements", []):
        tags = item.get("tags") or {}
        name = _safe_text(tags.get("name"))
        if not name:
            continue
        official_name = official_by_norm.get(_normalize_key(name))
        if not official_name:
            continue
        if official_name in polygons:
            continue
        f = None
        if item.get("type") == "relation":
            f = _feature_from_relation(item, official_name)
        elif item.get("type") == "way":
            f = _feature_from_way(item, official_name)
        if f:
            polygons[official_name] = f

    _RIO_VERDE_BAIRROS_POLYGONS_CACHE["updated_at"] = now
    _RIO_VERDE_BAIRROS_POLYGONS_CACHE["data"] = polygons
    return polygons


def _load_generated_bairros_polygons_geojson():
    path = Path(__file__).resolve().parent / "data" / "bairros_polygons_generated.geojson"
    if not path.exists():
        return None
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
    if not isinstance(payload, dict) or payload.get("type") != "FeatureCollection":
        return None
    features = payload.get("features")
    if not isinstance(features, list):
        return None
    return payload


def _saved_origin_snapshot(config=None):
    config = config or ConfiguracaoEntrega.objects.order_by("pk").first()
    return {
        "endereco": _safe_text(getattr(config, "origem_endereco", "")),
        "latitude": _safe_float(getattr(config, "origem_latitude", None)),
        "longitude": _safe_float(getattr(config, "origem_longitude", None)),
    }


def _resolve_saved_origin_result(config=None):
    origin = _saved_origin_snapshot(config)
    if origin["latitude"] is None or origin["longitude"] is None:
        return None
    return _origin_result_from_coordinates(
        origin["endereco"],
        origin["latitude"],
        origin["longitude"],
    )


def _blank_origin_result():
    return {
        "label": "",
        "street": "",
        "number": "",
        "district": "",
        "city": "Rio Verde",
        "state": "GO",
        "country": "Brasil",
        "countrycode": "BR",
        "lat": None,
        "lng": None,
        "type": "pending",
        "precision": "pending",
        "precision_label": _address_precision_label("pending"),
        "mode": "pending",
        "source": "pending",
    }


def _same_coordinate(left, right):
    left_float = _safe_float(left)
    right_float = _safe_float(right)
    if left_float is None or right_float is None:
        return left_float is None and right_float is None
    return abs(left_float - right_float) <= 0.0000001


def _origin_matches_saved_config(origin, config=None):
    saved_origin = _saved_origin_snapshot(config)
    return (
        _safe_text(origin.get("endereco")) == _safe_text(saved_origin["endereco"])
        and _same_coordinate(origin.get("latitude"), saved_origin["latitude"])
        and _same_coordinate(origin.get("longitude"), saved_origin["longitude"])
    )


def _origin_result_from_coordinates(address_text=None, latitude=None, longitude=None):
    lat = _safe_float(latitude)
    lng = _safe_float(longitude)
    if lat is None or lng is None:
        return None
    endereco = _safe_text(address_text) or "Ponto confirmado no mapa"
    return {
        "label": endereco,
        "street": endereco,
        "number": "",
        "district": "",
        "city": "Rio Verde",
        "state": "GO",
        "country": "Brasil",
        "countrycode": "BR",
        "lat": lat,
        "lng": lng,
        "type": "manual",
        "precision": "manual",
        "precision_label": _address_precision_label("manual"),
        "mode": "manual",
        "source": "manual",
    }


def _get_origin_coordinates():
    origin_result = _resolve_saved_origin_result()
    if not origin_result:
        return None, None
    return origin_result["lat"], origin_result["lng"]


def _fetch_route_summary(origin_lat, origin_lng, destination_lat, destination_lng):
    origin = f"{origin_lng:.7f},{origin_lat:.7f}"
    destination = f"{destination_lng:.7f},{destination_lat:.7f}"
    query = urlencode({"overview": "false", "alternatives": "false", "steps": "false"})
    request_url = f"{OSRM_ROUTE_BASE_URL}{origin};{destination}?{query}"
    external_request = Request(request_url, headers={"User-Agent": "PRATO-DELIVERY/1.0"})
    try:
        with urlopen(external_request, timeout=6) as response:
            payload = json.loads(response.read().decode("utf-8"))
    except (HTTPError, URLError, TimeoutError, ValueError, json.JSONDecodeError):
        return None, None

    if payload.get("code") != "Ok":
        return None, None

    routes = payload.get("routes") or []
    if not routes:
        return None, None

    route = routes[0]
    duration = _safe_float(route.get("duration"))
    distance = _safe_float(route.get("distance"))
    if duration is None or duration <= 0 or distance is None or distance <= 0:
        return None, None
    return duration, distance


def _to_decimal(value):
    try:
        if value is None:
            return None
        return Decimal(str(value).strip().replace(",", "."))
    except (InvalidOperation, TypeError, ValueError):
        return None


def _parse_percentual_0_100(value, label):
    percentual = _to_decimal(value)
    if percentual is None:
        raise ValueError(f"Informe {label} entre 0,00% e 100,00%.")
    percentual = percentual.quantize(Decimal("0.01"))
    if percentual < Decimal("0.00") or percentual > Decimal("100.00"):
        raise ValueError(f"Informe {label} entre 0,00% e 100,00%.")
    return percentual


def _calcular_frete_por_distancia(distance_km, faixas=None):
    distance_decimal = _to_decimal(distance_km)
    if distance_decimal is None or distance_decimal < 0:
        return DELIVERY_FRETE_PADRAO, None

    if faixas is None:
        faixas = list(FaixaFrete.objects.filter(ativo=True).order_by("ordem", "km_limite", "id"))
    else:
        normalized_faixas = []
        for faixa in faixas:
            if not getattr(faixa, "ativo", True):
                continue
            tipo = _safe_text(getattr(faixa, "tipo", FaixaFrete.Tipo.ATE)) or FaixaFrete.Tipo.ATE
            km_limite = _to_decimal(getattr(faixa, "km_limite", None))
            valor = _to_decimal(getattr(faixa, "valor", None))
            if km_limite is None or valor is None:
                continue
            ordem = getattr(faixa, "ordem", 0) or 0
            normalized_faixas.append(
                SimpleNamespace(
                    id=getattr(faixa, "id", None),
                    tipo=tipo,
                    km_limite=km_limite,
                    valor=valor,
                    ordem=int(ordem),
                    ativo=True,
                )
            )
        faixas = sorted(normalized_faixas, key=lambda faixa: (faixa.ordem, faixa.km_limite, faixa.id or 0))

    if not faixas:
        return DELIVERY_FRETE_PADRAO, None

    faixas_ate = [faixa for faixa in faixas if faixa.tipo == FaixaFrete.Tipo.ATE]
    faixas_acima = [faixa for faixa in faixas if faixa.tipo == FaixaFrete.Tipo.ACIMA]

    for faixa in faixas_ate:
        if distance_decimal <= faixa.km_limite:
            return faixa.valor, faixa

    if faixas_acima:
        candidatas = [faixa for faixa in faixas_acima if distance_decimal >= faixa.km_limite]
        if candidatas:
            candidatas.sort(key=lambda faixa: (faixa.km_limite, faixa.ordem, faixa.id), reverse=True)
            return candidatas[0].valor, candidatas[0]
        return faixas_acima[0].valor, faixas_acima[0]

    return faixas_ate[-1].valor, faixas_ate[-1]


def _format_faixa_label(faixa):
    if not faixa:
        return "Nenhuma faixa encontrada"
    prefixo = "Até" if faixa.tipo == FaixaFrete.Tipo.ATE else "Acima de"
    return f"{prefixo} {faixa.km_limite} km -> R$ {faixa.valor:.2f}".replace(".", ",")


def _masked_api_key(value):
    key = _safe_text(value)
    if len(key) <= 10:
        return key
    return f"{key[:6]}...{key[-4:]}"


def _google_maps_status():
    config = ConfiguracaoEntrega.objects.order_by("pk").first()
    api_key = config.google_maps_api_key_effective if config else getattr(settings, "GOOGLE_MAPS_API_KEY", "")
    enabled = bool(api_key)
    return {
        "enabled": enabled,
        "provider_label": "Google Maps" if enabled else "Google Maps nao configurado",
        "api_key_masked": _masked_api_key(api_key),
        "api_key_value": getattr(config, "google_maps_api_key", "") if config else "",
        "language": config.google_maps_language_effective if config else getattr(settings, "GOOGLE_MAPS_LANGUAGE", "pt-BR"),
        "region": config.google_maps_region_effective if config else getattr(settings, "GOOGLE_MAPS_REGION", "BR"),
        "language_value": getattr(config, "google_maps_language", "") if config else "",
        "region_value": getattr(config, "google_maps_region", "") if config else "",
        "required_apis": ["Maps JavaScript API", "Geocoding API", "Places API / Places API (New)"],
        "env_var_names": ["GOOGLE_MAPS_API_KEY", "GOOGLE_MAPS_LANGUAGE", "GOOGLE_MAPS_REGION"],
    }


def _next_faixa_ordem(faixas):
    if not faixas:
        return 10
    return max(int(getattr(faixa, "ordem", 0) or 0) for faixa in faixas) + 10


def _serialize_faixas_for_form(faixas=None, extra_rows=2):
    faixas = list(faixas if faixas is not None else FaixaFrete.objects.order_by("ordem", "km_limite", "id"))
    rows = []
    for faixa in faixas:
        rows.append(
            {
                "row_key": f"existing-{faixa.id}",
                "id": faixa.id,
                "tipo": faixa.tipo,
                "km_limite": f"{faixa.km_limite}",
                "valor": f"{faixa.valor}",
                "ativo": faixa.ativo,
                "ordem": faixa.ordem,
                "is_new": False,
            }
        )

    next_ordem = _next_faixa_ordem(faixas)
    for index in range(extra_rows):
        rows.append(
            {
                "row_key": f"new-{index + 1}",
                "id": "",
                "tipo": FaixaFrete.Tipo.ATE,
                "km_limite": "",
                "valor": "",
                "ativo": True,
                "ordem": next_ordem + (index * 10),
                "is_new": True,
            }
        )
    return rows


def _ensure_default_user_groups():
    Group.objects.get_or_create(name=ATENDENTE_GROUP_NAME)
    Group.objects.get_or_create(name=DIRETOR_GROUP_NAME)
    Group.objects.get_or_create(name=GERENTE_GROUP_NAME)


def _user_can_manage_order_payment(user):
    return bool(user.is_superuser or user.groups.filter(name=GERENTE_GROUP_NAME).exists())


def _serialize_users_for_admin():
    users = get_user_model().objects.prefetch_related("groups").order_by("username")
    return [
        {
            "id": user.id,
            "username": user.username,
            "first_name": user.first_name,
            "display_name": user.get_full_name() or user.username,
            "email": user.email,
            "is_active": user.is_active,
            "is_staff": user.is_staff,
            "is_superuser": user.is_superuser,
            "last_login": timezone.localtime(user.last_login).strftime("%d/%m/%Y %H:%M") if user.last_login else "",
            "date_joined": timezone.localtime(user.date_joined).strftime("%d/%m/%Y") if user.date_joined else "",
            "groups": list(user.groups.order_by("name")),
            "group_ids": {group.id for group in user.groups.all()},
            "group_ids_text": " ".join(str(group.id) for group in user.groups.all()),
        }
        for user in users
    ]


def _serialize_groups_for_admin():
    return Group.objects.annotate(user_count=Count("user")).order_by("name")


def _serialize_pedido_api_keys():
    return PedidoApiKey.objects.select_related("criado_por").order_by("-criado_em")


def _serialize_lista_impressao_admin():
    return PedidoListaImpressao.objects.select_related("pedido").order_by("-criado_em", "-id")[:100]


def _require_superuser_for_user_admin(request):
    if not request.user.is_superuser:
        raise PermissionError("Somente superusuarios podem administrar usuarios e classes.")


def _save_user_groups(user, group_ids):
    groups = Group.objects.filter(id__in=group_ids)
    user.groups.set(groups)


def _handle_user_admin_action(request):
    _require_superuser_for_user_admin(request)
    User = get_user_model()
    action = _safe_text(request.POST.get("action"))

    if action == "create_group":
        name = _safe_text(request.POST.get("group_name"))
        if not name:
            raise ValueError("Informe o nome da classe.")
        Group.objects.get_or_create(name=name)
        return "Classe criada."

    if action == "update_group":
        group = get_object_or_404(Group, id=request.POST.get("group_id"))
        name = _safe_text(request.POST.get("group_name"))
        if not name:
            raise ValueError("Informe o nome da classe.")
        if Group.objects.exclude(id=group.id).filter(name=name).exists():
            raise ValueError("Ja existe uma classe com esse nome.")
        group.name = name
        group.save()
        return "Classe atualizada."

    if action == "delete_group":
        group = get_object_or_404(Group, id=request.POST.get("group_id"))
        if group.user_set.exists():
            raise ValueError("Remova os usuarios dessa classe antes de excluir.")
        group.delete()
        return "Classe excluida."

    if action == "create_user":
        username = _safe_text(request.POST.get("username"))
        password = _safe_text(request.POST.get("password"))
        if not username:
            raise ValueError("Informe o usuario.")
        if not password:
            raise ValueError("Informe uma senha inicial.")
        if User.objects.filter(username=username).exists():
            raise ValueError("Ja existe um usuario com esse login.")
        user = User.objects.create_user(
            username=username,
            password=password,
            first_name=_safe_text(request.POST.get("first_name")),
            email=_safe_text(request.POST.get("email")),
        )
        user.is_active = request.POST.get("is_active") == "on"
        user.is_staff = request.POST.get("is_staff") == "on"
        user.is_superuser = request.POST.get("is_superuser") == "on"
        user.save()
        _save_user_groups(user, request.POST.getlist("groups"))
        return "Usuario criado."

    if action == "update_user":
        user = get_object_or_404(User, id=request.POST.get("user_id"))
        username = _safe_text(request.POST.get("username")) or user.username
        if User.objects.exclude(id=user.id).filter(username=username).exists():
            raise ValueError("Ja existe outro usuario com esse login.")
        user.username = username
        user.first_name = _safe_text(request.POST.get("first_name"))
        user.email = _safe_text(request.POST.get("email"))
        user.is_active = request.POST.get("is_active") == "on"
        user.is_staff = request.POST.get("is_staff") == "on"
        user.is_superuser = request.POST.get("is_superuser") == "on"
        if user.id == request.user.id:
            user.is_active = True
            user.is_staff = True
            user.is_superuser = True
        password = _safe_text(request.POST.get("password"))
        if password:
            user.set_password(password)
        user.save()
        _save_user_groups(user, request.POST.getlist("groups"))
        return "Usuario atualizado."

    raise ValueError("Acao de usuarios desconhecida.")


def _parse_faixa_rows(post_data):
    row_keys = post_data.getlist("faixa_row_key")
    faixa_ids = post_data.getlist("faixa_id")
    tipos = post_data.getlist("faixa_tipo")
    km_limites = post_data.getlist("faixa_km_limite")
    valores = post_data.getlist("faixa_valor")
    ordens = post_data.getlist("faixa_ordem")
    ativos = set(post_data.getlist("faixa_ativo"))
    deletes = set(post_data.getlist("faixa_delete"))

    rows = []
    total_rows = max(len(row_keys), len(tipos), len(km_limites), len(valores), len(ordens))
    for index in range(total_rows):
        row_key = _safe_text(row_keys[index] if index < len(row_keys) else f"row-{index}")
        faixa_id = _safe_text(faixa_ids[index] if index < len(faixa_ids) else "")
        tipo = _safe_text(tipos[index] if index < len(tipos) else FaixaFrete.Tipo.ATE) or FaixaFrete.Tipo.ATE
        km_limite = _safe_text(km_limites[index] if index < len(km_limites) else "")
        valor = _safe_text(valores[index] if index < len(valores) else "")
        ordem_raw = _safe_text(ordens[index] if index < len(ordens) else "0")
        try:
            ordem = int(ordem_raw or "0")
        except ValueError:
            ordem = 0

        rows.append(
            {
                "row_key": row_key,
                "id": faixa_id,
                "tipo": tipo if tipo in dict(FaixaFrete.Tipo.choices) else FaixaFrete.Tipo.ATE,
                "km_limite": km_limite,
                "valor": valor,
                "ativo": row_key in ativos,
                "delete": row_key in deletes,
                "ordem": ordem,
                "is_new": not faixa_id,
            }
        )
    return rows


def _preview_faixas_from_rows(rows):
    preview_rows = []
    for row in rows:
        if row.get("delete"):
            continue
        km_limite = _to_decimal(row.get("km_limite"))
        valor = _to_decimal(row.get("valor"))
        if km_limite is None or valor is None:
            continue
        preview_rows.append(
            SimpleNamespace(
                id=int(row["id"]) if _safe_text(row.get("id")).isdigit() else None,
                tipo=row.get("tipo") or FaixaFrete.Tipo.ATE,
                km_limite=km_limite,
                valor=valor,
                ordem=int(row.get("ordem") or 0),
                ativo=bool(row.get("ativo")),
            )
        )
    return preview_rows


def _save_faixa_rows(rows):
    existing_ids = set()
    for row in rows:
        faixa_id = _safe_text(row.get("id"))
        if faixa_id.isdigit():
            existing_ids.add(int(faixa_id))

    for faixa in FaixaFrete.objects.exclude(id__in=existing_ids):
        if any(row.get("delete") and _safe_text(row.get("id")) == str(faixa.id) for row in rows):
            faixa.delete()

    for row in rows:
        faixa_id = _safe_text(row.get("id"))
        km_limite = _to_decimal(row.get("km_limite"))
        valor = _to_decimal(row.get("valor"))
        should_delete = row.get("delete")

        if not faixa_id and km_limite is None and valor is None:
            continue

        if km_limite is None or valor is None:
            raise ValueError("Toda faixa precisa de km limite e valor para ser salva.")

        defaults = {
            "tipo": row.get("tipo") or FaixaFrete.Tipo.ATE,
            "km_limite": km_limite,
            "valor": valor,
            "ativo": bool(row.get("ativo")),
            "ordem": int(row.get("ordem") or 0),
        }

        if faixa_id.isdigit():
            faixa = FaixaFrete.objects.get(id=int(faixa_id))
            if should_delete:
                faixa.delete()
                continue
            for field, value in defaults.items():
                setattr(faixa, field, value)
            faixa.save()
            continue

        if should_delete:
            continue
        FaixaFrete.objects.create(**defaults)


def _build_destination_query_from_values(values):
    address = _safe_text(values.get("address") or values.get("destino_teste"))
    if address:
        return address

    parts = [
        _safe_text(values.get("bairro")),
        _safe_text(values.get("rua")),
        _safe_text(values.get("numero")),
        _safe_text(values.get("cidade")) or "Rio Verde",
        _safe_text(values.get("estado")) or "GO",
    ]
    return ", ".join([part for part in parts if part])


def _build_destination_query(request):
    return _build_destination_query_from_values(request.GET)


def _destination_result_from_values(values):
    destination_lat = _safe_float(values.get("lat") or values.get("latitude") or values.get("destino_teste_lat"))
    destination_lng = _safe_float(values.get("lng") or values.get("longitude") or values.get("destino_teste_lng"))
    resolved_label = _safe_text(values.get("resolved_label") or values.get("endereco_formatado") or values.get("destino_teste_label"))
    resolved_type = _safe_text(values.get("resolved_type") or values.get("geocode_tipo") or values.get("destino_teste_tipo"))
    resolved_precision = _safe_text(
        values.get("resolved_precision") or values.get("geocode_precision") or values.get("destino_teste_precision")
    )

    if destination_lat is not None and destination_lng is not None:
        precision = resolved_precision or _address_precision({"type": resolved_type, "number": ""})
        return {
            "label": resolved_label or _build_destination_query_from_values(values),
            "street": _safe_text(values.get("rua")) or resolved_label,
            "number": _safe_text(values.get("numero")),
            "district": _safe_text(values.get("bairro")),
            "city": _safe_text(values.get("cidade")) or "Rio Verde",
            "state": _safe_text(values.get("estado")) or "GO",
            "country": "Brasil",
            "countrycode": "BR",
            "lat": destination_lat,
            "lng": destination_lng,
            "type": resolved_type or "manual",
            "precision": precision or "manual",
            "precision_label": _address_precision_label(precision or "manual"),
            "source": "request",
        }

    return None


def _destination_result_from_request(request):
    return _destination_result_from_values(request.GET)


def _address_precision(item):
    if not item:
        return "pending"
    kind = _safe_text(item.get("type")).lower()
    number = _safe_text(item.get("number"))
    if kind == "manual":
        return "manual"
    if number:
        return "exact"
    if kind in {"house", "housenumber", "building", "entrance"}:
        return "exact"
    if kind in {"street", "road", "residential", "suburb", "district", "neighbourhood", "locality"}:
        return "approximate"
    return "approximate"


def _address_precision_label(precision):
    mapping = {
        "exact": "Endereço confirmado",
        "approximate": "Endereço aproximado",
        "manual": "Coordenadas manuais",
        "fallback": "Origem padrao",
        "pending": "A confirmar",
    }
    return mapping.get(precision, "A confirmar")


@require_GET
def api_address_delivery_time(request):
    destination_result = _destination_result_from_request(request)
    if not destination_result:
        return JsonResponse({"ok": False, "eta_minutes": None})

    origin_config = _saved_origin_snapshot()
    origin_result = _resolve_saved_origin_result()
    if not origin_result:
        return JsonResponse(
            {
                "ok": False,
                "eta_minutes": None,
                "error": "origin_not_configured",
            }
        )
    origin_lat = origin_result["lat"]
    origin_lng = origin_result["lng"]
    duration_seconds, distance_meters = _fetch_route_summary(
        origin_lat,
        origin_lng,
        destination_result["lat"],
        destination_result["lng"],
    )
    if duration_seconds is None or distance_meters is None:
        return JsonResponse({"ok": False, "eta_minutes": None})

    raw_minutes = duration_seconds / 60.0
    distance_km = max(distance_meters / 1000.0, 0.0)
    adjusted_minutes = (raw_minutes * DELIVERY_ETA_MULTIPLIER) + DELIVERY_ETA_BUFFER_MINUTES
    if distance_km <= DELIVERY_ETA_SHORT_TRIP_KM:
        adjusted_minutes += DELIVERY_ETA_SHORT_TRIP_PENALTY_MINUTES
    eta_minutes = max(1, int(math.ceil(adjusted_minutes)))
    frete_valor, faixa = _calcular_frete_por_distancia(distance_km)
    return JsonResponse(
        {
            "ok": True,
            "eta_minutes": eta_minutes,
            "distance_km": round(distance_km, 2),
            "shipping_fee": float(frete_valor),
            "shipping_fee_formatted": f"R$ {frete_valor:.2f}".replace(".", ","),
            "shipping_rule": {
                "id": faixa.id if faixa else None,
                "tipo": faixa.tipo if faixa else None,
                "km_limite": float(faixa.km_limite) if faixa else None,
                "valor": float(faixa.valor) if faixa else None,
            },
            "origin": origin_config["endereco"],
            "origin_lat": origin_lat,
            "origin_lng": origin_lng,
            "origin_mode": origin_result["mode"],
            "origin_label": origin_result["label"],
            "origin_type": origin_result["type"],
            "origin_precision": origin_result["precision"],
            "origin_precision_label": origin_result["precision_label"],
            "destination_lat": destination_result["lat"],
            "destination_lng": destination_result["lng"],
            "destination_label": destination_result["label"],
            "destination_type": destination_result["type"],
            "destination_precision": destination_result["precision"],
            "destination_precision_label": destination_result["precision_label"],
        }
    )


_money_decimal = money_decimal
_normalize_coupon_code = normalize_coupon_code

@require_POST
def api_validar_cupom(request):
    subtotal = _money_decimal(request.POST.get("subtotal"))
    frete = _money_decimal(request.POST.get("frete"))
    result = validar_cupom(request.POST.get("codigo"), subtotal, frete)
    return JsonResponse(
        {
            "ok": result["ok"],
            "message": result["message"],
            "codigo": _normalize_coupon_code(request.POST.get("codigo")),
            "desconto": f"{result['discount']:.2f}",
            "desconto_formatado": f"R$ {result['discount']:.2f}".replace(".", ","),
            "total": f"{(subtotal + frete - result['discount']):.2f}",
            "total_formatado": f"R$ {(subtotal + frete - result['discount']):.2f}".replace(".", ","),
        }
    )


def _tokens_from_history_payload(payload):
    tokens = []
    if not isinstance(payload, list):
        return tokens
    for item in payload:
        if isinstance(item, dict):
            token = str(item.get("token") or "").strip()
        else:
            token = str(item or "").strip()
        if token and token not in tokens:
            tokens.append(token)
        if len(tokens) >= 30:
            break
    return tokens


def _parse_order_history_tokens(raw):
    try:
        payload = json.loads(raw)
    except (TypeError, json.JSONDecodeError):
        return []
    return _tokens_from_history_payload(payload)


def _known_order_tokens_from_request(request):
    tokens = []
    for token in _parse_order_history_tokens(request.POST.get("known_order_tokens") or "[]"):
        if token not in tokens:
            tokens.append(token)
    for token in _parse_order_history_tokens(request.COOKIES.get(ORDER_HISTORY_COOKIE) or "[]"):
        if token not in tokens:
            tokens.append(token)
        if len(tokens) >= 30:
            break
    return tokens


def _checkout_key_from_request(request, prefix):
    raw_key = _safe_text(request.POST.get("checkout_key"))
    allowed = set("abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789_-")
    normalized = "".join(char for char in raw_key if char in allowed).strip()
    if len(normalized) < 16:
        return None
    return f"{prefix}:{normalized[:64]}"


def _request_expects_json(request):
    return (
        request.headers.get("X-Requested-With") == "XMLHttpRequest"
        or "application/json" in request.headers.get("Accept", "")
    )


def _order_history_for_response(request, pedido):
    history = [
        {
            "token": pedido.public_token,
            "numero": pedido.numero,
        }
    ]
    for token in _known_order_tokens_from_request(request):
        if token != pedido.public_token:
            history.append({"token": token})
        if len(history) >= 30:
            break
    return history


def _attach_order_history_cookie(response, request, pedido):
    response.set_cookie(
        ORDER_HISTORY_COOKIE,
        json.dumps(_order_history_for_response(request, pedido)),
        max_age=ORDER_HISTORY_COOKIE_MAX_AGE,
        samesite="Lax",
    )
    return response


def _order_created_response(request, pedido, reused=False):
    if _request_expects_json(request):
        response = JsonResponse(
            {
                "ok": True,
                "reused": reused,
                "pedido": _pedido_public_payload(pedido),
                "success_url": reverse("pedidos:sucesso", args=[pedido.public_token]),
            }
        )
        return _attach_order_history_cookie(response, request, pedido)
    return redirect("pedidos:sucesso", public_token=pedido.public_token)


def _record_order_created_metric(request, pedido):
    itens = list(pedido.itens.all())
    _record_access_event(
        request,
        AccessEvent.EventType.ORDER_CREATED,
        cart_items_count=sum(max(item.quantidade or 0, 0) for item in itens),
        cart_total=pedido.total,
        metadata={"pedido_id": pedido.id, "tipo_coleta": pedido.tipo_coleta},
    )


@require_POST
def criar_pedido(request):
    payload = request.POST.get("carrinho_payload")
    if not payload:
        return HttpResponseBadRequest("Carrinho vazio.")

    try:
        itens_payload = json.loads(payload)
    except json.JSONDecodeError:
        return HttpResponseBadRequest("Carrinho invalido.")

    if not itens_payload:
        return HttpResponseBadRequest("Carrinho vazio.")

    nome_cliente = request.POST.get("nome_cliente", "").strip() or "Cliente"
    telefone = normalize_phone(request.POST.get("telefone", ""))
    rua = request.POST.get("rua", "").strip()
    numero = request.POST.get("numero", "").strip()
    bairro = request.POST.get("bairro", "").strip()
    cidade = request.POST.get("cidade", "").strip()
    estado = request.POST.get("estado", "").strip()
    endereco_formatado = request.POST.get("endereco_formatado", "").strip()
    latitude_raw = request.POST.get("latitude", "").strip()
    longitude_raw = request.POST.get("longitude", "").strip()
    endereco_antigo = request.POST.get("endereco", "").strip()
    lote_quadra = request.POST.get("lote_quadra", "").strip()
    complemento = request.POST.get("complemento", "").strip()
    ponto_referencia = request.POST.get("ponto_referencia", "").strip()
    enviar_talheres_raw = request.POST.get("enviar_talheres", "nao").strip().lower()
    observacao_geral = request.POST.get("observacao_geral", "").strip()
    valor_frete_raw = request.POST.get("valor_frete", "").strip()
    distancia_km_raw = request.POST.get("distancia_km", "").strip()
    forma_pagamento = _safe_text(request.POST.get("forma_pagamento"))
    cupom_codigo = _normalize_coupon_code(request.POST.get("cupom_codigo"))
    checkout_key = _checkout_key_from_request(request, "entrega")
    if checkout_key:
        existing_pedido = Pedido.objects.filter(checkout_key=checkout_key).first()
        if existing_pedido:
            return _order_created_response(request, existing_pedido, reused=True)
    config_entrega = ConfiguracaoEntrega.get_solo()
    if not _configured_whatsapp_number(config_entrega):
        return HttpResponseBadRequest("Configure o número do WhatsApp antes de finalizar pedidos.")

    try:
        turno_pedido, disponibilidade_turno = validar_turno_para_pedido(Pedido.TipoColeta.ENTREGA)
    except ValueError as exc:
        return HttpResponseBadRequest(str(exc))
    if disponibilidade_turno is None and any(
        item.get("disponibilidade_turno_item_id") not in (None, "")
        for item in itens_payload
    ):
        return HttpResponseBadRequest(
            "O turno Prato Pronto encerrou. Atualize o cardapio e monte o carrinho novamente."
        )

    endereco = _build_order_address(rua, numero, bairro, cidade, estado, lote_quadra) or endereco_formatado or endereco_antigo

    if not all([cidade, estado, endereco]) or not (rua or endereco_formatado):
        return HttpResponseBadRequest("Preencha os campos obrigatorios.")
    formas_pagamento_checkout = {
        Pedido.FormaPagamento.PIX,
        Pedido.FormaPagamento.DINHEIRO,
        Pedido.FormaPagamento.CARTAO,
    }
    if forma_pagamento not in formas_pagamento_checkout:
        return HttpResponseBadRequest("Selecione uma forma de pagamento.")
    if forma_pagamento == Pedido.FormaPagamento.PIX and not _safe_text(config_entrega.pix_chave):
        return HttpResponseBadRequest("Configure a chave Pix antes de aceitar pagamento online.")
    enviar_talheres = enviar_talheres_raw != "nao"

    latitude = None
    longitude = None
    valor_frete = Decimal("0.00")
    distancia_km = Decimal("0.00")
    destination_result = _destination_result_from_values(request.POST)
    if not destination_result:
        return HttpResponseBadRequest("Confirme o ponto de entrega no mapa.")

    origin_result = _resolve_saved_origin_result()
    if not origin_result:
        return HttpResponseBadRequest("Configure e salve a origem de entrega antes de receber pedidos.")

    try:
        if latitude_raw:
            latitude = Decimal(latitude_raw)
        if longitude_raw:
            longitude = Decimal(longitude_raw)
    except (InvalidOperation, TypeError):
        latitude = None
        longitude = None
    try:
        if valor_frete_raw:
            valor_frete = Decimal(valor_frete_raw)
        if distancia_km_raw:
            distancia_km = Decimal(distancia_km_raw)
    except (InvalidOperation, TypeError):
        valor_frete = Decimal("0.00")
        distancia_km = Decimal("0.00")

    if destination_result.get("label"):
        endereco_formatado = destination_result["label"]
    latitude = _to_decimal(destination_result.get("lat"))
    longitude = _to_decimal(destination_result.get("lng"))
    duration_seconds, distance_meters = _fetch_route_summary(
        origin_result["lat"],
        origin_result["lng"],
        destination_result["lat"],
        destination_result["lng"],
    )
    if duration_seconds is None or distance_meters is None:
        return HttpResponseBadRequest("Não foi possível calcular a rota para o ponto de entrega.")
    distancia_km = Decimal(str(round(max(distance_meters / 1000.0, 0.0), 2)))
    valor_frete, _ = _calcular_frete_por_distancia(distancia_km)

    if distancia_km > 0 and valor_frete == Decimal("0.00"):
        valor_frete, _ = _calcular_frete_por_distancia(distancia_km)
    if valor_frete < 0:
        valor_frete = Decimal("0.00")

    try:
        with transaction.atomic():
            pedido = Pedido.objects.create(
                nome_cliente=nome_cliente,
                telefone=telefone,
                rua=rua,
                numero_endereco=numero,
                bairro=bairro,
                cidade=cidade,
                estado=estado,
                endereco_formatado=endereco_formatado,
                latitude=latitude,
                longitude=longitude,
                endereco=endereco,
                lote_quadra=lote_quadra,
                complemento=complemento,
                ponto_referencia=ponto_referencia,
                tipo_coleta=Pedido.TipoColeta.ENTREGA,
                forma_pagamento=forma_pagamento,
                enviar_talheres=enviar_talheres,
                canal=Pedido.Canal.SITE,
                ifood=False,
                turno=turno_pedido,
                disponibilidade_turno=disponibilidade_turno,
                orientacao_turno_snapshot=(
                    turno_pedido.orientacao_cliente if disponibilidade_turno else ""
                ),
                observacao_geral=observacao_geral,
                status=Pedido.Status.AGUARDANDO_APROVACAO,
                valor_frete=valor_frete,
                distancia_km=distancia_km,
                checkout_key=checkout_key,
            )
            create_order_items_from_payload(
                pedido,
                itens_payload,
                disponibilidade_turno=disponibilidade_turno,
                permitir_promocao=turno_pedido.permite_promocao,
            )
            recalculate_order_totals(pedido, cupom_codigo=cupom_codigo)
            sync_movimentacao_caixa_pedido(pedido)
            inherit_customer_from_known_tokens(pedido, _known_order_tokens_from_request(request))
    except ValueError as exc:
        return HttpResponseBadRequest(str(exc))
    _record_order_created_metric(request, pedido)
    return _order_created_response(request, pedido)


@require_POST
def criar_retirada(request):
    payload = request.POST.get("carrinho_payload")
    if not payload:
        return HttpResponseBadRequest("Carrinho vazio.")

    try:
        itens_payload = json.loads(payload)
    except json.JSONDecodeError:
        return HttpResponseBadRequest("Carrinho invalido.")

    if not itens_payload:
        return HttpResponseBadRequest("Carrinho vazio.")

    config_entrega = ConfiguracaoEntrega.get_solo()
    if not _configured_whatsapp_number(config_entrega):
        return HttpResponseBadRequest("Configure o número do WhatsApp antes de finalizar pedidos.")

    try:
        turno_pedido, disponibilidade_turno = validar_turno_para_pedido(Pedido.TipoColeta.RETIRADA)
    except ValueError as exc:
        return HttpResponseBadRequest(str(exc))
    if disponibilidade_turno is None and any(
        item.get("disponibilidade_turno_item_id") not in (None, "")
        for item in itens_payload
    ):
        return HttpResponseBadRequest(
            "O turno Prato Pronto encerrou. Atualize o cardapio e monte o carrinho novamente."
        )

    nome_cliente = request.POST.get("nome_cliente", "").strip() or "Cliente"
    observacao_geral = request.POST.get("observacao_geral", "").strip()
    enviar_talheres_raw = request.POST.get("enviar_talheres", "nao").strip().lower()
    cupom_codigo = _normalize_coupon_code(request.POST.get("cupom_codigo"))
    checkout_key = _checkout_key_from_request(request, "retirada")
    if checkout_key:
        existing_pedido = Pedido.objects.filter(checkout_key=checkout_key).first()
        if existing_pedido:
            return _order_created_response(request, existing_pedido, reused=True)

    try:
        with transaction.atomic():
            pedido = Pedido.objects.create(
                nome_cliente=nome_cliente,
                telefone="",
                rua="",
                numero_endereco="",
                bairro="",
                cidade="Rio Verde",
                estado="GO",
                endereco_formatado="Retirada no local",
                endereco="Retirada no local",
                tipo_coleta=Pedido.TipoColeta.RETIRADA,
                forma_pagamento=Pedido.FormaPagamento.DINHEIRO,
                enviar_talheres=enviar_talheres_raw != "nao",
                canal=Pedido.Canal.SITE,
                ifood=False,
                turno=turno_pedido,
                disponibilidade_turno=disponibilidade_turno,
                orientacao_turno_snapshot=(
                    turno_pedido.orientacao_cliente if disponibilidade_turno else ""
                ),
                observacao_geral=observacao_geral,
                status=Pedido.Status.AGUARDANDO_APROVACAO,
                valor_frete=Decimal("0.00"),
                distancia_km=Decimal("0.00"),
                checkout_key=checkout_key,
            )
            create_order_items_from_payload(
                pedido,
                itens_payload,
                disponibilidade_turno=disponibilidade_turno,
                permitir_promocao=turno_pedido.permite_promocao,
            )
            recalculate_order_totals(pedido, cupom_codigo=cupom_codigo)
            sync_movimentacao_caixa_pedido(pedido)
            inherit_customer_from_known_tokens(pedido, _known_order_tokens_from_request(request))
    except ValueError as exc:
        return HttpResponseBadRequest(str(exc))
    _record_order_created_metric(request, pedido)
    return _order_created_response(request, pedido)


@never_cache
def sucesso(request, public_token):
    pedido = get_object_or_404(Pedido.objects.prefetch_related("itens"), public_token=public_token)
    mensagem = montar_mensagem_whatsapp(pedido)
    whatsapp_url = _build_whatsapp_order_url(pedido)
    response = render(
        request,
        "pedidos/sucesso.html",
        {
            "pedido": pedido,
            "whatsapp_url": whatsapp_url,
        },
    )
    return _attach_order_history_cookie(response, request, pedido)


def _pedido_public_payload(pedido, include_items=False):
    status_step = {
        Pedido.Status.AGUARDANDO_APROVACAO: 2,
        Pedido.Status.NOVO: 3,
        Pedido.Status.EM_PREPARO: 3,
        Pedido.Status.PAGAMENTO_RECEBIDO: 3,
        Pedido.Status.AGUARDANDO_ENTREGADOR: 3,
        Pedido.Status.SAIU_ENTREGA: 4,
        Pedido.Status.FINALIZADO: 4,
        Pedido.Status.CANCELADO: 0,
    }
    payload = {
        "token": pedido.public_token,
        "numero": pedido.numero,
        "status": pedido.status,
        "status_label": pedido.status_label_contextual,
        "status_step": status_step.get(pedido.status, 1),
        "criado_em": timezone.localtime(pedido.criado_em).isoformat(),
        "horario": timezone.localtime(pedido.criado_em).strftime("%d/%m %H:%M"),
        "itens_count": pedido.itens.count(),
        "total": f"R$ {pedido.total:.2f}".replace(".", ","),
        "promocao_desconto": f"R$ {pedido.promocao_desconto:.2f}".replace(".", ",") if pedido.promocao_desconto else "",
        "promocao_descricao": pedido.promocao_descricao,
        "cupom_desconto": f"R$ {pedido.cupom_desconto:.2f}".replace(".", ",") if pedido.cupom_desconto else "",
        "cupom_codigo": pedido.cupom_codigo,
        "pagamento": pedido.get_forma_pagamento_display(),
        "turno": pedido.turno.codigo if pedido.turno_id else TurnoAtendimento.Codigo.PRINCIPAL,
        "turno_label": turno_label_pedido(pedido),
        "orientacao_turno": pedido.orientacao_turno_snapshot,
        "acompanhamento_url": reverse("pedidos:acompanhar_pedido", args=[pedido.public_token]),
    }
    if include_items:
        payload["endereco"] = pedido.endereco
        payload["valor_frete"] = f"R$ {pedido.valor_frete:.2f}".replace(".", ",")
        payload["itens"] = [
            {
                "nome": item.nome_prato_snapshot,
                "variacao": item.variacao_nome_snapshot,
                "quantidade": item.quantidade,
                "observacao": item.observacao,
                "classificacao_saida": item.classificacao_saida,
                "classificacao_saida_label": item.get_classificacao_saida_display(),
                "subtotal": f"R$ {item.subtotal:.2f}".replace(".", ","),
            }
            for item in pedido.itens.all()
        ]
    return payload


@never_cache
def meus_pedidos(request):
    return render(request, "pedidos/meus_pedidos.html")


@never_cache
def acompanhar_pedido(request, public_token):
    pedido = get_object_or_404(Pedido.objects.prefetch_related("itens"), public_token=public_token)
    whatsapp_url = _build_whatsapp_order_url(pedido)
    return render(
        request,
        "pedidos/acompanhar_pedido.html",
        {
            "pedido": pedido,
            "pedido_publico": _pedido_public_payload(pedido, include_items=True),
            "whatsapp_url": whatsapp_url,
        },
    )


@never_cache
@require_POST
def api_meus_pedidos(request):
    try:
        payload = json.loads(request.body.decode("utf-8") or "{}")
    except (json.JSONDecodeError, UnicodeDecodeError):
        return HttpResponseBadRequest("Payload invalido.")

    raw_tokens = payload.get("tokens", [])
    if not isinstance(raw_tokens, list):
        return HttpResponseBadRequest("Lista de tokens invalida.")

    tokens = []
    seen = set()
    for raw_token in raw_tokens:
        if not isinstance(raw_token, str):
            continue
        token = raw_token.strip()
        if not token or len(token) > 64 or token in seen:
            continue
        seen.add(token)
        tokens.append(token)
        if len(tokens) >= 30:
            break

    pedidos = Pedido.objects.prefetch_related("itens").filter(public_token__in=tokens).order_by("-criado_em")
    return JsonResponse({"pedidos": [_pedido_public_payload(pedido) for pedido in pedidos]})


def _pedidos_api_queryset(params, *, include_items=False):
    pedidos = Pedido.objects.select_related("cupom").all()
    if include_items:
        pedidos = pedidos.prefetch_related("itens")

    status = _safe_text(params.get("status"))
    if status:
        pedidos = pedidos.filter(status=status)

    tipo_coleta = _safe_text(params.get("tipo_coleta"))
    if tipo_coleta:
        pedidos = pedidos.filter(tipo_coleta=tipo_coleta)

    criado_em = _safe_text(params.get("criado_em"))
    if criado_em:
        parsed_date = parse_date(criado_em)
        pedidos = pedidos.filter(criado_em__date=parsed_date) if parsed_date else pedidos.none()

    updated_after = _safe_text(params.get("updated_after") or params.get("atualizado_apos"))
    if updated_after:
        updated_after = updated_after.replace(" ", "+")
        parsed_datetime = parse_datetime(updated_after)
        if parsed_datetime and timezone.is_naive(parsed_datetime):
            parsed_datetime = timezone.make_aware(parsed_datetime, timezone.get_current_timezone())
        pedidos = pedidos.filter(atualizado_em__gt=parsed_datetime) if parsed_datetime else pedidos.none()

    numero = _safe_text(params.get("numero"))
    if numero:
        try:
            pedidos = pedidos.filter(numero=int(numero))
        except (TypeError, ValueError):
            pedidos = pedidos.none()

    telefone = _safe_text(params.get("telefone"))
    if telefone:
        pedidos = pedidos.filter(telefone__icontains=telefone)

    return pedidos.order_by("-criado_em", "-id")


def _safe_api_limit(params, default=None, max_limit=100):
    raw_limit = _safe_text(params.get("limit"))
    if not raw_limit:
        return default
    try:
        return min(max(int(raw_limit), 1), max_limit)
    except (TypeError, ValueError):
        return default


PEDIDOS_API_DEFAULT_LIMIT = 50
PEDIDOS_API_MAX_LIMIT = 100
PEDIDOS_API_FULL_DEFAULT_LIMIT = 10
PEDIDOS_API_FULL_MAX_LIMIT = 25


def _safe_api_offset(params):
    try:
        return max(int(params.get("offset", 0)), 0)
    except (TypeError, ValueError):
        return 0


def _api_key_from_request(request):
    authorization = _safe_text(request.headers.get("Authorization"))
    if authorization.lower().startswith("bearer "):
        return authorization[7:].strip()
    return _safe_text(request.headers.get("X-API-Key"))


def _require_pedidos_api_key(request):
    api_key = PedidoApiKey.authenticate(_api_key_from_request(request))
    if not api_key:
        return None
    usage_cache_key = f"pedido-api-key-usage:{api_key.pk}:{api_key.prefixo}"
    if cache.add(usage_cache_key, "1", 60):
        PedidoApiKey.objects.filter(pk=api_key.pk).update(ultimo_uso_em=timezone.now())
    return api_key


def _invalid_api_key_response():
    return JsonResponse({"ok": False, "error": "invalid_api_key"}, status=401)


def _serialize_print_queue_item(item):
    return {
        "id": item.id,
        "nome_cliente": item.nome_cliente,
        "public_token": item.public_token,
        "criado_em": item.criado_em,
    }


def _api_rate_limit_response(retry_after):
    response = JsonResponse(
        {
            "ok": False,
            "error": "rate_limited",
            "message": "Aguarde antes de consultar novamente.",
            "retry_after": retry_after,
        },
        status=429,
    )
    response["Retry-After"] = str(retry_after)
    return response


def _check_api_rate_limit(api_key, bucket, limit, window_seconds):
    cache_key = f"pedido-api-rate:{api_key.pk}:{api_key.prefixo}:{bucket}"
    current = cache.get(cache_key, 0)
    if current >= limit:
        return _api_rate_limit_response(window_seconds)
    if current:
        cache.incr(cache_key)
    else:
        cache.set(cache_key, 1, window_seconds)
    return None


@require_GET
def api_pedidos(request):
    api_key = _require_pedidos_api_key(request)
    if not api_key:
        return _invalid_api_key_response()
    limited = _check_api_rate_limit(api_key, "pedidos-list", 12, 60)
    if limited:
        return limited

    fields = _safe_text(request.GET.get("fields")).lower()
    include_items = fields in {"full", "complete", "detalhado"}
    pedidos = _pedidos_api_queryset(request.GET, include_items=include_items)
    count = pedidos.count()
    limit = _safe_api_limit(
        request.GET,
        default=PEDIDOS_API_FULL_DEFAULT_LIMIT if include_items else PEDIDOS_API_DEFAULT_LIMIT,
        max_limit=PEDIDOS_API_FULL_MAX_LIMIT if include_items else PEDIDOS_API_MAX_LIMIT,
    )
    offset = _safe_api_offset(request.GET)
    has_more = False
    if limit is not None:
        page = list(pedidos[offset : offset + limit + 1])
        has_more = len(page) > limit
        pedidos_payload = page[:limit]
    else:
        pedidos_payload = pedidos
    serializer = serialize_pedido_api if include_items else serialize_pedido_summary_api
    return JsonResponse(
        {
            "count": count,
            "limit": limit,
            "offset": offset,
            "has_more": has_more,
            "next_offset": offset + len(pedidos_payload) if limit is not None and has_more else None,
            "pedidos": [serializer(pedido) for pedido in pedidos_payload],
        }
    )


@require_GET
def api_pedido_detalhe(request, pedido_id):
    api_key = _require_pedidos_api_key(request)
    if not api_key:
        return _invalid_api_key_response()
    limited = _check_api_rate_limit(api_key, "pedido-detail", 180, 60)
    if limited:
        return limited
    pedido = get_object_or_404(
        Pedido.objects.select_related("cupom").prefetch_related("itens"),
        id=pedido_id,
    )
    return JsonResponse({"pedido": serialize_pedido_api(pedido)})


@require_GET
def api_pedido_detalhe_token(request, public_token):
    api_key = _require_pedidos_api_key(request)
    if not api_key:
        return _invalid_api_key_response()
    limited = _check_api_rate_limit(api_key, "pedido-token-detail", 180, 60)
    if limited:
        return limited
    pedido = get_object_or_404(
        Pedido.objects.select_related("cupom").prefetch_related("itens"),
        public_token=public_token,
    )
    return JsonResponse({"pedido": serialize_pedido_api(pedido)})


@require_GET
def api_lista_impressao(request):
    api_key = _require_pedidos_api_key(request)
    if not api_key:
        return _invalid_api_key_response()
    limited = _check_api_rate_limit(api_key, "lista-impressao", 60, 60)
    if limited:
        return limited
    itens = PedidoListaImpressao.objects.all()

    desde_id = _safe_text(request.GET.get("desde_id"))
    if desde_id:
        try:
            itens = itens.filter(id__gt=int(desde_id))
        except (TypeError, ValueError):
            itens = itens.none()

    try:
        limit = min(max(int(request.GET.get("limit", 100)), 1), 500)
    except (TypeError, ValueError):
        limit = 100

    page = list(itens.order_by("criado_em", "id")[: limit + 1])
    has_more = len(page) > limit
    data = [_serialize_print_queue_item(item) for item in page[:limit]]
    next_desde_id = data[-1]["id"] if data else (int(desde_id) if desde_id.isdigit() else None)
    return JsonResponse(
        {
            "count": len(data),
            "itens": data,
            "has_more": has_more,
            "next_desde_id": next_desde_id,
        }
    )


def _dashboard_periodo(period):
    hoje = timezone.localdate()
    if period == "month":
        inicio = hoje.replace(day=1)
        return {
            "key": "month",
            "label": "Este mês",
            "inicio": inicio,
            "fim": hoje,
        }
    inicio = hoje - timedelta(days=6)
    return {
        "key": "7d",
        "label": "Últimos 7 dias",
        "inicio": inicio,
        "fim": hoje,
    }


def _dashboard_days(inicio, fim):
    days = []
    cursor = inicio
    while cursor <= fim:
        days.append(cursor)
        cursor += timedelta(days=1)
    return days


def _dashboard_percentual_vs_ontem(vendas_hoje, vendas_ontem):
    if vendas_ontem <= 0:
        if vendas_hoje <= 0:
            return 0
        return 100
    variacao = ((vendas_hoje - vendas_ontem) / vendas_ontem) * 100
    return int(round(variacao))


def _dashboard_insight(percentual):
    if percentual <= -1:
        return f"As vendas cairam {abs(percentual)}% em relacao a ontem."
    if percentual >= 1:
        return f"As vendas cresceram {percentual}% em relacao a ontem."
    return "As vendas ficaram estaveis em relacao a ontem."


def _heatmap_period_window(period_key):
    hoje = timezone.localdate()
    normalized = (period_key or "7d").strip().lower()

    if normalized == "today":
        return hoje, hoje
    if normalized == "30d":
        return hoje - timedelta(days=29), hoje
    if normalized == "all":
        return None, None
    return hoje - timedelta(days=6), hoje


def _line_path(points):
    if not points:
        return ""
    first = points[0]
    return " ".join(
        [f"M {first['x']:.2f} {first['y']:.2f}"]
        + [f"L {point['x']:.2f} {point['y']:.2f}" for point in points[1:]]
    )


def _build_sales_chart_context(labels, values, average):
    width = 960
    height = 320
    pad = {"top": 28, "right": 26, "bottom": 42, "left": 52}
    chart_w = width - pad["left"] - pad["right"]
    chart_h = height - pad["top"] - pad["bottom"]
    max_y = max([1, *values, average])
    step_x = chart_w / (len(values) - 1) if len(values) > 1 else chart_w

    points = []
    for idx, value in enumerate(values):
        x = pad["left"] + step_x * idx
        y = pad["top"] + chart_h - ((value / max_y) * chart_h)
        points.append({"x": x, "y": y, "v": value})

    line_d = _line_path(points)
    area_d = ""
    if points:
        area_d = (
            f"{line_d} L {points[-1]['x']:.2f} {pad['top'] + chart_h:.2f} "
            f"L {points[0]['x']:.2f} {pad['top'] + chart_h:.2f} Z"
        )
    avg_y = pad["top"] + chart_h - ((average / max_y) * chart_h)
    label_points = []
    for idx, label in enumerate(labels):
        label_points.append(
            {
                "x": pad["left"] + step_x * idx,
                "label": label,
            }
        )

    grid_lines = [pad["top"] + (chart_h / 4) * i for i in range(5)]
    return {
        "width": width,
        "height": height,
        "left": pad["left"],
        "right": width - pad["right"],
        "avg_y": avg_y,
        "line_d": line_d,
        "area_d": area_d,
        "points": points,
        "labels": label_points,
        "grid_lines": grid_lines,
    }


def _build_hour_chart_context(labels, values):
    width = 960
    height = 260
    pad = {"top": 20, "right": 20, "bottom": 34, "left": 40}
    chart_w = width - pad["left"] - pad["right"]
    chart_h = height - pad["top"] - pad["bottom"]
    max_y = max([1, *values])
    step_x = chart_w / (len(values) - 1) if len(values) > 1 else chart_w

    points = []
    for idx, value in enumerate(values):
        x = pad["left"] + step_x * idx
        y = pad["top"] + chart_h - ((value / max_y) * chart_h)
        points.append({"x": x, "y": y, "v": value})

    line_d = _line_path(points)
    label_points = []
    for idx, label in enumerate(labels):
        if idx % 2 == 0 or idx == len(labels) - 1:
            label_points.append({"x": pad["left"] + step_x * idx, "label": label})

    grid_lines = [pad["top"] + (chart_h / 4) * i for i in range(5)]
    return {
        "width": width,
        "height": height,
        "left": pad["left"],
        "right": width - pad["right"],
        "baseline": pad["top"] + chart_h,
        "line_d": line_d,
        "labels": label_points,
        "grid_lines": grid_lines,
    }


def _metrics_period(period):
    hoje = timezone.localdate()
    normalized = _safe_text(period).lower()
    if normalized == "today":
        return {"key": "today", "label": "Hoje", "inicio": hoje, "fim": hoje}
    if normalized == "30d":
        return {"key": "30d", "label": "Ultimos 30 dias", "inicio": hoje - timedelta(days=29), "fim": hoje}
    if normalized == "month":
        return {"key": "month", "label": "Este mes", "inicio": hoje.replace(day=1), "fim": hoje}
    return {"key": "7d", "label": "Ultimos 7 dias", "inicio": hoje - timedelta(days=6), "fim": hoje}


def _percent(part, total):
    if not total:
        return 0
    return int(round((part / total) * 100))


def _rate_label(part, total):
    return f"{_percent(part, total)}%"


def _access_page_label(path):
    normalized = _safe_text(path)
    if normalized in {"/", ""}:
        return "Cardapio"
    if normalized.startswith("/carrinho"):
        return "Carrinho"
    if normalized.startswith("/checkout"):
        return "Caixa"
    if normalized.startswith("/meus-pedidos"):
        return "Meus pedidos"
    return normalized[:80] or "-"


def _access_history_rows(events):
    return [
        {
            "time": timezone.localtime(event.created_at).strftime("%H:%M:%S"),
            "date": timezone.localtime(event.created_at).strftime("%d/%m/%Y"),
            "label": event.get_event_type_display(),
            "page": _access_page_label(event.path),
            "session": event.session_key[-6:] if event.session_key else "-",
            "detail": event.metadata.get("origem", "") if isinstance(event.metadata, dict) else "",
        }
        for event in events
    ]


def _build_access_metrics_context(period):
    periodo = _metrics_period(period)
    inicio = periodo["inicio"]
    fim = periodo["fim"]
    days = _dashboard_days(inicio, fim)
    events = AccessEvent.objects.filter(created_at__date__gte=inicio, created_at__date__lte=fim)
    visible_events = events.exclude(event_type=AccessEvent.EventType.PAGE_ACTIVE)

    event_counts_raw = visible_events.values("event_type").annotate(total=Count("id"))
    event_counts = {row["event_type"]: row["total"] for row in event_counts_raw}
    event_sessions = {
        row["event_type"]: row["total"]
        for row in visible_events.values("event_type").annotate(total=Count("session_key", distinct=True))
    }

    menu_sessions = event_sessions.get(AccessEvent.EventType.MENU_VIEW, 0)
    cart_sessions = event_sessions.get(AccessEvent.EventType.CART_VIEW, 0)
    pickup_sessions = event_sessions.get(AccessEvent.EventType.PICKUP_SUBMIT, 0)
    checkout_sessions = event_sessions.get(AccessEvent.EventType.CHECKOUT_VIEW, 0)
    order_sessions = event_sessions.get(AccessEvent.EventType.ORDER_CREATED, 0)
    closing_sessions = pickup_sessions + checkout_sessions

    funnel_steps = [
        {
            "label": "Cardapio",
            "count": menu_sessions,
            "rate": "100%" if menu_sessions else "0%",
            "width": 100 if menu_sessions else 0,
        },
        {
            "label": "Carrinho",
            "count": cart_sessions,
            "rate": _rate_label(cart_sessions, menu_sessions),
            "width": _percent(cart_sessions, menu_sessions),
        },
        {
            "label": "Retirada",
            "count": pickup_sessions,
            "rate": _rate_label(pickup_sessions, cart_sessions),
            "width": _percent(pickup_sessions, menu_sessions),
        },
        {
            "label": "Caixa",
            "count": checkout_sessions,
            "rate": _rate_label(checkout_sessions, cart_sessions),
            "width": _percent(checkout_sessions, menu_sessions),
        },
        {
            "label": "Pedido criado",
            "count": order_sessions,
            "rate": _rate_label(order_sessions, closing_sessions or cart_sessions),
            "width": _percent(order_sessions, menu_sessions),
        },
    ]

    events_by_day_raw = (
        visible_events.filter(event_type=AccessEvent.EventType.MENU_VIEW)
        .annotate(day=TruncDate("created_at"))
        .values("day")
        .annotate(total=Count("id"))
        .order_by("day")
    )
    events_by_day_map = {row["day"]: row["total"] for row in events_by_day_raw}
    access_series = [events_by_day_map.get(day, 0) for day in days]
    access_labels = [day.strftime("%d/%m") for day in days]
    max_access = max(access_series) if access_series else 0
    access_bars = [
        {
            "label": label,
            "value": value,
            "height": max(4, _percent(value, max_access)) if max_access else 4,
        }
        for label, value in zip(access_labels, access_series)
    ]

    events_by_hour_raw = (
        visible_events.annotate(hour=ExtractHour("created_at"))
        .values("hour")
        .annotate(total=Count("id"))
        .order_by("hour")
    )
    events_by_hour_map = {int(row["hour"] or 0): row["total"] for row in events_by_hour_raw}
    hour_series = [events_by_hour_map.get(hour, 0) for hour in range(24)]
    peak_hour_value = max(hour_series) if hour_series else 0
    peak_hour = hour_series.index(peak_hour_value) if peak_hour_value else 0

    top_items = list(
        visible_events.filter(event_type=AccessEvent.EventType.ADD_TO_CART)
        .exclude(item_type="")
        .exclude(item_id__isnull=True)
        .values("item_type", "item_id")
        .annotate(total=Count("id"), unidades=Sum("cart_items_count"))
        .order_by("-total", "item_type", "item_id")[:8]
    )

    item_names = {
        "prato": {item.id: item.nome for item in Prato.objects.filter(id__in=[row["item_id"] for row in top_items if row["item_type"] == "prato"])},
        "adicional": {item.id: item.nome for item in Adicional.objects.filter(id__in=[row["item_id"] for row in top_items if row["item_type"] == "adicional"])},
        "bebida": {item.id: item.nome for item in Bebida.objects.filter(id__in=[row["item_id"] for row in top_items if row["item_type"] == "bebida"])},
    }
    for row in top_items:
        row["nome"] = item_names.get(row["item_type"], {}).get(row["item_id"], f"{row['item_type']} #{row['item_id']}")

    metrics_payload = {
        "access_labels": access_labels,
        "access_series": access_series,
        "hour_labels": [f"{hour:02d}h" for hour in range(24)],
        "hour_series": hour_series,
    }

    pedidos_abertos = Pedido.objects.exclude(
        status__in=[Pedido.Status.RASCUNHO, Pedido.Status.AGUARDANDO_APROVACAO, Pedido.Status.FINALIZADO, Pedido.Status.CANCELADO]
    ).count()
    aprovacao_count = Pedido.objects.filter(status=Pedido.Status.AGUARDANDO_APROVACAO).count()
    pedidos_metricas = Pedido.objects.filter(criado_em__date__gte=inicio, criado_em__date__lte=fim).exclude(status=Pedido.Status.RASCUNHO)
    envio_orders_total = pedidos_metricas.filter(tipo_coleta=Pedido.TipoColeta.ENTREGA).count()
    retirada_orders_total = pedidos_metricas.filter(tipo_coleta=Pedido.TipoColeta.RETIRADA).count()
    total_pedidos_metricas = pedidos_metricas.count()
    pedidos_por_canal_raw = pedidos_metricas.values("canal").annotate(total=Count("id"))
    pedidos_por_canal = {row["canal"]: row["total"] for row in pedidos_por_canal_raw}
    pedidos_balcao_total = pedidos_por_canal.get(Pedido.Canal.BALCAO, 0)
    pedidos_site_total = pedidos_por_canal.get(Pedido.Canal.SITE, 0)
    pedidos_ifood_total = pedidos_por_canal.get(Pedido.Canal.IFOOD, 0)
    envio_orders_share = _rate_label(envio_orders_total, total_pedidos_metricas)
    retirada_orders_share = _rate_label(retirada_orders_total, total_pedidos_metricas)
    active_since = timezone.now() - timedelta(seconds=90)
    active_users_count = AccessEvent.objects.filter(
        event_type=AccessEvent.EventType.PAGE_ACTIVE,
        created_at__gte=active_since,
    ).values("session_key").distinct().count()
    access_history = _access_history_rows(visible_events.order_by("-created_at", "-id")[:40])

    return {
        "periodo_key": periodo["key"],
        "periodo_label": periodo["label"],
        "total_cardapio": event_counts.get(AccessEvent.EventType.MENU_VIEW, 0),
        "total_carrinho": event_counts.get(AccessEvent.EventType.CART_VIEW, 0),
        "total_retirada": event_counts.get(AccessEvent.EventType.PICKUP_SUBMIT, 0),
        "total_checkout": event_counts.get(AccessEvent.EventType.CHECKOUT_VIEW, 0),
        "total_pedidos_metricas": total_pedidos_metricas,
        "pedidos_balcao_total": pedidos_balcao_total,
        "pedidos_site_total": pedidos_site_total,
        "pedidos_ifood_total": pedidos_ifood_total,
        "envio_orders_total": envio_orders_total,
        "retirada_orders_total": retirada_orders_total,
        "envio_orders_share": envio_orders_share,
        "retirada_orders_share": retirada_orders_share,
        "checkout_submit_total": event_counts.get(AccessEvent.EventType.CHECKOUT_SUBMIT, 0),
        "pickup_submit_total": event_counts.get(AccessEvent.EventType.PICKUP_SUBMIT, 0),
        "cart_conversion": _rate_label(cart_sessions, menu_sessions),
        "pickup_conversion": _rate_label(pickup_sessions, cart_sessions),
        "checkout_conversion": _rate_label(checkout_sessions, cart_sessions),
        "order_conversion": _rate_label(order_sessions, closing_sessions or cart_sessions),
        "funnel_steps": funnel_steps,
        "top_items": top_items,
        "peak_hour": peak_hour,
        "peak_hour_value": peak_hour_value,
        "metrics_payload": metrics_payload,
        "access_bars": access_bars,
        "active_users_count": active_users_count,
        "access_history": access_history,
        "pedidos_badge": pedidos_abertos,
        "aprovacao_count": aprovacao_count,
    }


def _access_metrics_json(context):
    return {
        "periodo_key": context["periodo_key"],
        "periodo_label": context["periodo_label"],
        "kpis": {
            "total_cardapio": context["total_cardapio"],
            "total_carrinho": context["total_carrinho"],
            "total_retirada": context["total_retirada"],
            "total_checkout": context["total_checkout"],
            "total_pedidos_metricas": context["total_pedidos_metricas"],
            "pedidos_balcao_total": context["pedidos_balcao_total"],
            "pedidos_site_total": context["pedidos_site_total"],
            "pedidos_ifood_total": context["pedidos_ifood_total"],
            "envio_orders_total": context["envio_orders_total"],
            "retirada_orders_total": context["retirada_orders_total"],
            "envio_orders_share": context["envio_orders_share"],
            "retirada_orders_share": context["retirada_orders_share"],
            "cart_conversion": context["cart_conversion"],
            "pickup_conversion": context["pickup_conversion"],
            "checkout_conversion": context["checkout_conversion"],
            "order_conversion": context["order_conversion"],
        },
        "funnel_steps": context["funnel_steps"],
        "access_bars": context["access_bars"],
        "top_items": [
            {
                "nome": row["nome"],
                "total": int(row["total"] or 0),
                "item_type": row["item_type"],
                "item_id": row["item_id"],
            }
            for row in context["top_items"]
        ],
        "peak": {
            "hour": context["peak_hour"],
            "value": context["peak_hour_value"],
            "label": f"{context['peak_hour']:02d}h concentrou {context['peak_hour_value']} eventos.",
        },
        "active_users_count": context["active_users_count"],
        "access_history": context["access_history"],
        "updated_at": timezone.localtime().strftime("%H:%M:%S"),
        "pedidos_badge": context["pedidos_badge"],
        "aprovacao_count": context["aprovacao_count"],
    }


@staff_member_required(login_url="/admin/login/")
def metricas_acesso(request):
    context = _build_access_metrics_context(request.GET.get("period", "today"))
    return render(request, "pedidos/metricas_acesso.html", context)


@staff_member_required(login_url="/admin/login/")
@require_GET
def api_metricas_acesso(request):
    context = _build_access_metrics_context(request.GET.get("period", "today"))
    return JsonResponse(_access_metrics_json(context))


def _dashboard_nav_context(data_selecionada, hoje):
    week_labels = ["Seg", "Ter", "Qua", "Qui", "Sex", "Sab", "Dom"]
    inicio = data_selecionada - timedelta(days=3)
    fim = data_selecionada + timedelta(days=3)
    dates = [inicio + timedelta(days=index) for index in range(7)]

    pedidos_por_dia = {
        row["day"]: int(row["total"] or 0)
        for row in (
            Pedido.objects.filter(status=Pedido.Status.FINALIZADO, criado_em__date__gte=inicio, criado_em__date__lte=fim)
            .annotate(day=TruncDate("criado_em"))
            .values("day")
            .annotate(total=Count("id"))
        )
    }
    marmitas_por_dia = {
        row["day"]: int(row["total"] or 0)
        for row in (
            ItemPedido.objects.filter(
                pedido__status=Pedido.Status.FINALIZADO,
                pedido__criado_em__date__gte=inicio,
                pedido__criado_em__date__lte=fim,
            )
            .filter(
                Q(prato__isnull=False)
                | Q(
                    pedido__observacao_geral__startswith="[IMPORTADO DO SISTEMA ANTIGO]",
                    prato__isnull=True,
                    bebida__isnull=True,
                    adicional__isnull=True,
                    observacao__icontains="Tipo legado: dish",
                )
            )
            .filter(classificacao_saida=ItemPedido.ClassificacaoSaida.VENDIDA)
            .annotate(day=TruncDate("pedido__criado_em"))
            .values("day")
            .annotate(total=Sum("quantidade"))
        )
    }

    nav_days = []
    for day in dates:
        delta = (day - hoje).days
        if delta == 0:
            relative_label = "Hoje"
        elif delta == -1:
            relative_label = "Ontem"
        elif delta == 1:
            relative_label = "Amanha"
        else:
            relative_label = week_labels[day.weekday()]
        nav_days.append(
            {
                "date": day,
                "weekday": week_labels[day.weekday()],
                "relative_label": relative_label,
                "pedidos": pedidos_por_dia.get(day, 0),
                "marmitas": marmitas_por_dia.get(day, 0),
                "is_selected": day == data_selecionada,
                "is_today": day == hoje,
                "has_data": bool(pedidos_por_dia.get(day, 0) or marmitas_por_dia.get(day, 0)),
            }
        )

    def day_summary(day):
        delta = (day - hoje).days
        return {
            "date": day,
            "relative_label": "Hoje" if delta == 0 else "Ontem" if delta == -1 else "Amanha" if delta == 1 else week_labels[day.weekday()],
            "pedidos": pedidos_por_dia.get(day, 0),
            "marmitas": marmitas_por_dia.get(day, 0),
        }

    return {
        "days": nav_days,
        "previous_day": day_summary(data_selecionada - timedelta(days=1)),
        "next_day": day_summary(data_selecionada + timedelta(days=1)),
        "previous_week": data_selecionada - timedelta(days=7),
        "next_week": data_selecionada + timedelta(days=7),
    }


@staff_member_required(login_url="/admin/login/")
def cozinha(request):
    if not user_is_diretor(request.user):
        return redirect("pedidos:cozinha_operacao")

    hoje = timezone.localdate()
    data_selecionada = parse_date(_safe_text(request.GET.get("data"))) or hoje
    dias_semana = ["Segunda-feira", "Terca-feira", "Quarta-feira", "Quinta-feira", "Sexta-feira", "Sabado", "Domingo"]
    if data_selecionada == hoje:
        rotulo_data_dashboard = "Hoje"
    elif data_selecionada == hoje - timedelta(days=1):
        rotulo_data_dashboard = "Ontem"
    elif data_selecionada == hoje + timedelta(days=1):
        rotulo_data_dashboard = "Amanha"
    else:
        rotulo_data_dashboard = "Dia selecionado"

    if request.method == "POST":
        try:
            marmitas_produzidas = max(int(request.POST.get("marmitas_produzidas") or 0), 0)
            consumo_interno = max(int(request.POST.get("consumo_interno") or 0), 0)
            custo_insumos = max(money_decimal(request.POST.get("custo_insumos")), Decimal("0.00"))
        except (TypeError, ValueError):
            return HttpResponseBadRequest("Quantidade de marmitas invalida.")
        resumo, _created = ResumoOperacionalDia.objects.get_or_create(data=data_selecionada)
        resumo.marmitas_produzidas = marmitas_produzidas
        resumo.consumo_interno = consumo_interno
        resumo.custo_insumos = custo_insumos
        resumo.save(update_fields=["marmitas_produzidas", "consumo_interno", "custo_insumos", "atualizado_em"])
        return redirect(f"{reverse('pedidos:cozinha')}?data={data_selecionada.isoformat()}")

    dashboard = get_dashboard_diaria(data_selecionada)
    pedidos_novos = Pedido.objects.filter(status=Pedido.Status.NOVO).count()

    return render(
        request,
        "pedidos/cozinha_dashboard.html",
        {
            "dashboard": dashboard,
            "data_selecionada": data_selecionada,
            "dia_semana_dashboard": dias_semana[data_selecionada.weekday()],
            "rotulo_data_dashboard": rotulo_data_dashboard,
            "dashboard_nav": _dashboard_nav_context(data_selecionada, hoje),
            "hoje": hoje,
            "pedidos_novos": pedidos_novos,
        },
    )


@staff_member_required(login_url="/admin/login/")
@require_GET
def exportar_pedidos_xlsx(request):
    if not user_is_diretor(request.user):
        return redirect("pedidos:cozinha_operacao")

    inicio = parse_date(_safe_text(request.GET.get("inicio")))
    fim = parse_date(_safe_text(request.GET.get("fim")))
    if not inicio or not fim:
        return HttpResponseBadRequest("Informe as datas inicial e final.")
    if inicio > fim:
        return HttpResponseBadRequest("A data inicial nao pode ser posterior a data final.")

    pedidos = _marcar_clientes_recorrentes(
        Pedido.objects.filter(criado_em__date__range=(inicio, fim))
        .exclude(status=Pedido.Status.RASCUNHO)
        .prefetch_related("itens")
        .order_by("criado_em", "id")
    )

    workbook = Workbook()
    pedidos_sheet = workbook.active
    pedidos_sheet.title = "Pedidos"
    pedidos_headers = [
        "Numero",
        "Data",
        "Hora",
        "Cliente",
        "Telefone",
        "Relacionamento do cliente",
        "Canal",
        "Status",
        "Tipo de coleta",
        "Forma de pagamento",
        "Pagamento recebido",
        "Endereco",
        "Bairro",
        "Itens",
        "Quantidade de itens",
        "Subtotal dos itens",
        "Frete",
        "Desconto promocao",
        "Desconto cupom",
        "Total",
        "Observacao",
        "Turno",
        "Orientacao do produto",
    ]
    pedidos_sheet.append(pedidos_headers)

    def csv_safe(value):
        text = _safe_text(value)
        if text.startswith(("=", "+", "-", "@")):
            return f"'{text}"
        return text

    for pedido in pedidos:
        criado_em = timezone.localtime(pedido.criado_em)
        itens = list(pedido.itens.all())
        itens_texto = " | ".join(
            (
                f"{item.quantidade}x {item.nome_prato_snapshot}"
                f"{f' - {item.variacao_nome_snapshot}' if item.variacao_nome_snapshot else ''}"
            )
            for item in itens
        )
        quantidade_itens = sum(item.quantidade for item in itens)
        subtotal_itens = sum((item.subtotal for item in itens), Decimal("0.00"))

        pedidos_sheet.append(
            [
                pedido.numero or pedido.id,
                criado_em.date(),
                criado_em.time().replace(second=0, microsecond=0),
                csv_safe(pedido.nome_cliente),
                csv_safe(pedido.telefone),
                csv_safe(pedido.cliente_recorrente_label if pedido.cliente_recorrente else ""),
                csv_safe(pedido.get_canal_display()),
                csv_safe(pedido.status_label_contextual),
                csv_safe(pedido.get_tipo_coleta_display()),
                csv_safe(pedido.get_forma_pagamento_display()),
                "Sim" if pedido.pagamento_recebido else "Nao",
                csv_safe(pedido.endereco_formatado or pedido.endereco),
                csv_safe(pedido.bairro),
                csv_safe(itens_texto),
                quantidade_itens,
                subtotal_itens,
                pedido.valor_frete,
                pedido.promocao_desconto,
                pedido.cupom_desconto,
                pedido.total,
                csv_safe(pedido.observacao_geral),
                csv_safe(turno_label_pedido(pedido)),
                csv_safe(pedido.orientacao_turno_snapshot),
            ]
        )

    dashboard_sheet = workbook.create_sheet("Dashboard diaria")
    dashboard_headers = [
        "Data",
        "Pedidos finalizados",
        "Pedidos recorrentes",
        "Pedidos balcao",
        "Pedidos site",
        "Pedidos iFood",
        "Faturamento total",
        "Faturamento iFood",
        "Pedidos Prato Pronto",
        "Faturamento Prato Pronto",
        "Pratos Prontos vendidos",
        "Estoque Prato Pronto inicial",
        "Estoque Prato Pronto restante",
        "Taxa iFood (%)",
        "Taxa iFood (R$)",
        "Custo entrega",
        "Marmitas produzidas",
        "Marmitas saida",
        "Marmitas vendidas",
        "Marmitas cortesia",
        "Marmitas promocao",
        "Consumo interno",
        "Marmitas excedentes",
        "Custo insumos",
        "Custo por marmita produzida",
        "Custo por marmita vendida",
        "Custo total producao",
        "Custo total operacional",
        "Resultado operacional",
    ]
    dashboard_sheet.append(dashboard_headers)

    data_atual = inicio
    while data_atual <= fim:
        dashboard = get_dashboard_diaria(data_atual)
        canais = {canal["key"]: canal["total"] for canal in dashboard["canais"]}
        dashboard_sheet.append(
            [
                data_atual,
                dashboard["total_pedidos"],
                dashboard["pedidos_recorrentes"],
                canais.get(Pedido.Canal.BALCAO, 0),
                canais.get(Pedido.Canal.SITE, 0),
                canais.get(Pedido.Canal.IFOOD, 0),
                dashboard["faturamento_total"],
                dashboard["faturamento_ifood"],
                dashboard["pedidos_prato_pronto"],
                dashboard["faturamento_prato_pronto"],
                dashboard["pratos_prontos_vendidos"],
                dashboard["estoque_prato_pronto_inicial"],
                dashboard["estoque_prato_pronto_restante"],
                dashboard["taxa_ifood_percentual"],
                dashboard["taxa_ifood_valor"],
                dashboard["custo_entrega"],
                dashboard["marmitas_produzidas"],
                dashboard["marmitas_saida"],
                dashboard["marmitas_vendidas"],
                dashboard["marmitas_cortesia"],
                dashboard["marmitas_promocao"],
                dashboard["consumo_interno"],
                dashboard["marmitas_excedentes"],
                dashboard["custo_insumos"],
                dashboard["custo_unitario_marmita"],
                dashboard["custo_unitario_marmita_vendida"],
                dashboard["custo_total_producao"],
                dashboard["custo_total_operacional"],
                dashboard["resultado_operacional"],
            ]
        )
        data_atual += timedelta(days=1)

    header_fill = PatternFill("solid", fgColor="4E342E")
    header_font = Font(color="FFFFFF", bold=True)
    currency_format = 'R$ #,##0.00'

    for sheet in (pedidos_sheet, dashboard_sheet):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 34

    pedidos_sheet.column_dimensions["A"].width = 11
    pedidos_sheet.column_dimensions["B"].width = 12
    pedidos_sheet.column_dimensions["C"].width = 9
    pedidos_sheet.column_dimensions["D"].width = 26
    pedidos_sheet.column_dimensions["E"].width = 17
    pedidos_sheet.column_dimensions["F"].width = 24
    pedidos_sheet.column_dimensions["G"].width = 14
    pedidos_sheet.column_dimensions["H"].width = 22
    pedidos_sheet.column_dimensions["I"].width = 17
    pedidos_sheet.column_dimensions["J"].width = 22
    pedidos_sheet.column_dimensions["K"].width = 19
    pedidos_sheet.column_dimensions["L"].width = 38
    pedidos_sheet.column_dimensions["M"].width = 20
    pedidos_sheet.column_dimensions["N"].width = 48
    pedidos_sheet.column_dimensions["O"].width = 18
    for column in ("P", "Q", "R", "S", "T"):
        pedidos_sheet.column_dimensions[column].width = 18
    pedidos_sheet.column_dimensions["U"].width = 38

    for row in pedidos_sheet.iter_rows(min_row=2):
        row[1].number_format = "dd/mm/yyyy"
        row[2].number_format = "hh:mm"
        for index in range(15, 20):
            row[index].number_format = currency_format
        for index in (3, 11, 13, 20):
            row[index].alignment = Alignment(vertical="top", wrap_text=True)

    for column_index in range(1, len(dashboard_headers) + 1):
        dashboard_sheet.column_dimensions[get_column_letter(column_index)].width = 20
    dashboard_sheet.column_dimensions["A"].width = 13
    for row in dashboard_sheet.iter_rows(min_row=2):
        row[0].number_format = "dd/mm/yyyy"
        for index in (6, 7, 9, 10, 18, 19, 20, 21, 22, 23):
            row[index].number_format = currency_format
        row[8].number_format = "0.00"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        f'attachment; filename="pedidos_dashboard_{inicio.isoformat()}_a_{fim.isoformat()}.xlsx"'
    )
    return response


def _contabil_query(data_selecionada, extra=None):
    params = {"data": data_selecionada.isoformat()}
    if extra:
        params.update(extra)
    return urlencode(params)


def _redirect_contabil_caixa(data_selecionada, extra=None):
    return redirect(f"{reverse('pedidos:contabil_caixa')}?{_contabil_query(data_selecionada, extra)}")


def _redirect_contabil_conta(data_selecionada, extra=None):
    return redirect(f"{reverse('pedidos:contabil_conta')}?{_contabil_query(data_selecionada, extra)}")


def _pedidos_contabil_context(data_selecionada, filtros):
    statuses_ignorados = [Pedido.Status.RASCUNHO, Pedido.Status.CANCELADO]
    base = Pedido.objects.exclude(status__in=statuses_ignorados).filter(**filtros)
    pedidos_dia = (
        base.filter(criado_em__date=data_selecionada)
        .select_related("terminal")
        .order_by("terminal__ordem", "criado_em", "id")
    )
    total_dia = pedidos_dia.aggregate(total=Sum("total")).get("total") or Decimal("0.00")
    total_anterior = (
        base.filter(criado_em__date__lt=data_selecionada).aggregate(total=Sum("total")).get("total") or Decimal("0.00")
    )
    total_atual = total_anterior + total_dia
    return {
        "pedidos": list(pedidos_dia),
        "total_anterior": total_anterior,
        "total_dia": total_dia,
        "total_atual": total_atual,
        "total_anterior_label": money_label(total_anterior),
        "total_dia_label": money_label(total_dia),
        "total_dia_signed_label": signed_money_label(total_dia, positive_prefix=True),
        "total_atual_label": money_label(total_atual),
    }


def _caixa_summary_items(caixa):
    return [
        {"label": "Saldo ontem", "value": caixa["saldo_anterior_label"]},
        {"label": "Saidas", "value": caixa["saidas_dia_label"]},
        {"label": "Entradas", "value": caixa["entradas_dia_label"]},
        {"label": "Caixa atual", "value": caixa["saldo_atual_label"]},
    ]


def _conta_summary_items(conta):
    return [
        {"label": "Saldo ontem", "value": conta["saldo_anterior_label"]},
        {"label": "Saidas", "value": conta["saidas_dia_label"]},
        {"label": "Entradas", "value": conta["entradas_dia_label"]},
        {"label": "Conta atual", "value": conta["saldo_atual_label"]},
    ]


def _pedidos_summary_items(pedidos_contabil):
    return [
        {"label": "Total ontem", "value": pedidos_contabil["total_anterior_label"]},
        {"label": "Saidas", "value": "- R$ 0,00"},
        {"label": "Entradas", "value": pedidos_contabil["total_dia_signed_label"]},
        {"label": "Total atual", "value": pedidos_contabil["total_atual_label"]},
    ]


def _render_contabil_pedidos(request, *, aba, rota, titulo, kicker, filtros, mensagem_vazia):
    if not user_is_diretor(request.user):
        return redirect("pedidos:cozinha_operacao")

    hoje = timezone.localdate()
    data_selecionada = parse_date(_safe_text(request.GET.get("data"))) or hoje
    dias_semana = ["Segunda-feira", "Terca-feira", "Quarta-feira", "Quinta-feira", "Sexta-feira", "Sabado", "Domingo"]
    pedidos_contabil = _pedidos_contabil_context(data_selecionada, filtros)
    return render(
        request,
        "pedidos/contabil_pedidos_lista.html",
        {
            "active": "contabil",
            "contabil_aba": aba,
            "rota_atual": rota,
            "titulo": titulo,
            "kicker": kicker,
            "mensagem_vazia": mensagem_vazia,
            "data_selecionada": data_selecionada,
            "dia_semana_dashboard": dias_semana[data_selecionada.weekday()],
            "dashboard_nav": _dashboard_nav_context(data_selecionada, hoje),
            "hoje": hoje,
            "pedidos_contabil": pedidos_contabil,
            "contabil_summary_items": _pedidos_summary_items(pedidos_contabil),
            "pedidos_novos": Pedido.objects.filter(status=Pedido.Status.NOVO).count(),
        },
    )


@staff_member_required(login_url="/admin/login/")
def contabil(request):
    return redirect("pedidos:contabil_caixa")


@staff_member_required(login_url="/admin/login/")
def contabil_caixa(request):
    if not user_is_diretor(request.user):
        return redirect("pedidos:cozinha_operacao")

    hoje = timezone.localdate()
    data_selecionada = parse_date(_safe_text(request.GET.get("data"))) or hoje
    feedback = None
    feedback_kind = "success"

    if request.method == "POST":
        action = _safe_text(request.POST.get("action"))
        try:
            if action == "create_movimentacao":
                criar_movimentacao_caixa(
                    data_movimento=data_selecionada,
                    post=request.POST,
                    user=request.user,
                )
                return _redirect_contabil_caixa(data_selecionada)

            if action == "update_movimentacao":
                movimento = get_object_or_404(
                    MovimentacaoCaixa,
                    id=request.POST.get("movimentacao_id"),
                    excluido_em__isnull=True,
                )
                atualizar_movimentacao_caixa(movimento, request.POST, request.user)
                return _redirect_contabil_caixa(data_selecionada)

            if action == "duplicate_movimentacao":
                movimento = get_object_or_404(
                    MovimentacaoCaixa,
                    id=request.POST.get("movimentacao_id"),
                    excluido_em__isnull=True,
                )
                duplicada = duplicar_movimentacao_caixa(movimento, request.user)
                return _redirect_contabil_caixa(duplicada.data_movimento, {"movimentacao": duplicada.id})

            if action == "delete_movimentacao":
                movimento = get_object_or_404(
                    MovimentacaoCaixa,
                    id=request.POST.get("movimentacao_id"),
                    excluido_em__isnull=True,
                )
                excluir_movimentacao_caixa(movimento, request.user)
                return _redirect_contabil_caixa(data_selecionada)

            if action == "create_terminal":
                nome = _safe_text(request.POST.get("terminal_nome"))
                if not nome:
                    raise ValueError("Informe o nome do terminal.")
                codigo_base = unicodedata.normalize("NFKD", nome).encode("ascii", "ignore").decode("ascii").lower()
                codigo_base = "-".join(part for part in codigo_base.replace("_", "-").split() if part)[:24] or "terminal"
                codigo = codigo_base
                suffix = 2
                while TerminalCaixa.objects.filter(codigo=codigo).exists():
                    codigo = f"{codigo_base}-{suffix}"[:30]
                    suffix += 1
                novo_terminal = TerminalCaixa.objects.create(
                    nome=nome,
                    codigo=codigo,
                    ordem=(TerminalCaixa.objects.aggregate(max_ordem=Max("ordem")).get("max_ordem") or 0) + 10,
                )
                return _redirect_contabil_caixa(data_selecionada)
        except ValueError as exc:
            feedback = str(exc)
            feedback_kind = "error"

    caixa = caixa_diario_context(data_selecionada)
    movimentacao_edicao = None
    movimentacao_id = request.GET.get("movimentacao")
    if movimentacao_id:
        movimentacao_edicao = (
            MovimentacaoCaixa.objects.filter(id=movimentacao_id, excluido_em__isnull=True)
            .select_related("terminal", "categoria")
            .first()
        )

    dias_semana = ["Segunda-feira", "Terca-feira", "Quarta-feira", "Quinta-feira", "Sexta-feira", "Sabado", "Domingo"]
    terminais = list(TerminalCaixa.objects.filter(ativo=True).order_by("ordem", "nome"))
    categorias = list(CategoriaMovimentacaoCaixa.objects.filter(ativo=True).order_by("nome"))
    return render(
        request,
        "pedidos/contabil_caixa.html",
        {
            "active": "contabil",
            "contabil_aba": "caixa",
            "data_selecionada": data_selecionada,
            "dia_semana_dashboard": dias_semana[data_selecionada.weekday()],
            "dashboard_nav": _dashboard_nav_context(data_selecionada, hoje),
            "hoje": hoje,
            "terminais": terminais,
            "terminal_padrao": terminal_padrao(),
            "categorias": categorias,
            "categorias_datalist_id": "contabil-categorias",
            "caixa": caixa,
            "contabil_summary_items": _caixa_summary_items(caixa),
            "contabil_action_modal": "#contabil-movimento-modal",
            "feedback": feedback,
            "feedback_kind": feedback_kind,
            "movimentacao_edicao": movimentacao_edicao,
            "money_label": money_label,
            "signed_money_label": signed_money_label,
            "pedidos_novos": Pedido.objects.filter(status=Pedido.Status.NOVO).count(),
        },
    )


@staff_member_required(login_url="/admin/login/")
def contabil_ifood(request):
    return _render_contabil_pedidos(
        request,
        aba="ifood",
        rota="pedidos:contabil_ifood",
        titulo="Ifood",
        kicker="Pedidos iFood",
        filtros={"canal": Pedido.Canal.IFOOD},
        mensagem_vazia="Nenhum pedido do iFood registrado neste dia.",
    )


@staff_member_required(login_url="/admin/login/")
def contabil_conta(request):
    if not user_is_diretor(request.user):
        return redirect("pedidos:cozinha_operacao")

    hoje = timezone.localdate()
    data_selecionada = parse_date(_safe_text(request.GET.get("data"))) or hoje
    feedback = None
    feedback_kind = "success"

    if request.method == "POST":
        action = _safe_text(request.POST.get("action"))
        try:
            if action == "create_movimentacao":
                criar_movimentacao_conta(
                    data_movimento=data_selecionada,
                    post=request.POST,
                    user=request.user,
                )
                return _redirect_contabil_conta(data_selecionada)

            if action == "update_movimentacao":
                movimento = get_object_or_404(
                    MovimentacaoConta,
                    id=request.POST.get("movimentacao_id"),
                    excluido_em__isnull=True,
                )
                atualizar_movimentacao_conta(movimento, request.POST, request.user)
                return _redirect_contabil_conta(data_selecionada)

            if action == "duplicate_movimentacao":
                movimento = get_object_or_404(
                    MovimentacaoConta,
                    id=request.POST.get("movimentacao_id"),
                    excluido_em__isnull=True,
                )
                duplicada = duplicar_movimentacao_conta(movimento, request.user)
                return _redirect_contabil_conta(duplicada.data_movimento, {"movimentacao": duplicada.id})

            if action == "delete_movimentacao":
                movimento = get_object_or_404(
                    MovimentacaoConta,
                    id=request.POST.get("movimentacao_id"),
                    excluido_em__isnull=True,
                )
                excluir_movimentacao_conta(movimento, request.user)
                return _redirect_contabil_conta(data_selecionada)
        except ValueError as exc:
            feedback = str(exc)
            feedback_kind = "error"

    conta = conta_diario_context(data_selecionada)
    movimentacao_edicao = None
    movimentacao_id = request.GET.get("movimentacao")
    if movimentacao_id:
        movimentacao_edicao = (
            MovimentacaoConta.objects.filter(id=movimentacao_id, excluido_em__isnull=True)
            .select_related("banco", "categoria")
            .first()
        )

    dias_semana = ["Segunda-feira", "Terca-feira", "Quarta-feira", "Quinta-feira", "Sexta-feira", "Sabado", "Domingo"]
    bancos = list(BancoConta.objects.filter(ativo=True).order_by("ordem", "nome"))
    categorias = list(CategoriaMovimentacaoConta.objects.filter(ativo=True).order_by("nome"))
    return render(
        request,
        "pedidos/contabil_conta.html",
        {
            "active": "contabil",
            "contabil_aba": "conta",
            "data_selecionada": data_selecionada,
            "dia_semana_dashboard": dias_semana[data_selecionada.weekday()],
            "dashboard_nav": _dashboard_nav_context(data_selecionada, hoje),
            "hoje": hoje,
            "bancos": bancos,
            "banco_padrao": banco_padrao(),
            "categorias": categorias,
            "categorias_datalist_id": "contabil-conta-categorias",
            "conta": conta,
            "contabil_summary_items": _conta_summary_items(conta),
            "contabil_action_modal": "#contabil-conta-movimento-modal",
            "feedback": feedback,
            "feedback_kind": feedback_kind,
            "movimentacao_edicao": movimentacao_edicao,
            "pedidos_novos": Pedido.objects.filter(status=Pedido.Status.NOVO).count(),
        },
    )


@staff_member_required(login_url="/admin/login/")
def contabil_ajustes(request):
    return redirect(f"{reverse('pedidos:ajustes_admin')}?aba=contabil")


@staff_member_required(login_url="/admin/login/")
@require_GET
def api_order_heatmap(request):
    period = request.GET.get("period", "7d")
    start_date, end_date = _heatmap_period_window(period)

    queryset = Pedido.objects.exclude(status__in=[Pedido.Status.CANCELADO, Pedido.Status.RASCUNHO]).filter(
        latitude__isnull=False,
        longitude__isnull=False,
    )

    if start_date and end_date:
        queryset = queryset.filter(criado_em__date__gte=start_date, criado_em__date__lte=end_date)

    points_raw = queryset.order_by("-criado_em").values("latitude", "longitude", "bairro")[:5000]

    points = []
    for row in points_raw:
        latf = _safe_float(row.get("latitude"))
        lngf = _safe_float(row.get("longitude"))
        if latf is None or lngf is None:
            continue
        if not (-90 <= latf <= 90 and -180 <= lngf <= 180):
            continue
        points.append(
            {
                "lat": latf,
                "lng": lngf,
                "weight": 1,
                "bairro": (row.get("bairro") or "").strip(),
            }
        )

    return JsonResponse(points, safe=False)


@require_GET
def api_bairros_rio_verde(request):
    if request.GET.get("refresh") == "1":
        _RIO_VERDE_BAIRROS_CACHE["updated_at"] = None
        _RIO_VERDE_BAIRROS_CACHE["data"] = []
    bairros = _fetch_rio_verde_bairros()
    return JsonResponse(bairros, safe=False)


@require_GET
def api_bairros_polygons(request):
    if request.GET.get("refresh") == "1":
        _RIO_VERDE_BAIRROS_POLYGONS_CACHE["updated_at"] = None
        _RIO_VERDE_BAIRROS_POLYGONS_CACHE["data"] = {}
    polygons = _fetch_rio_verde_bairros_polygons()
    features = list(polygons.values())
    if len(features) >= 20:
        return JsonResponse({"type": "FeatureCollection", "features": features})

    # Fallback 1: generated local GeoJSON (street-based approximation).
    generated = _load_generated_bairros_polygons_geojson()
    if generated and isinstance(generated.get("features"), list) and generated.get("features"):
        return JsonResponse(generated)

    # Fallback 2: no polygons available.
    return JsonResponse({"type": "FeatureCollection", "features": []})


def _weekday_label_pt(date_value):
    names = [
        "SEGUNDA-FEIRA",
        "TERCA-FEIRA",
        "QUARTA-FEIRA",
        "QUINTA-FEIRA",
        "SEXTA-FEIRA",
        "SABADO",
        "DOMINGO",
    ]
    return names[date_value.weekday()]


def _cozinha_operacao_payload():
    today = timezone.localdate()
    now = timezone.localtime()
    entregues_hoje = (
        ItemPedido.objects.filter(
            pedido__status=Pedido.Status.FINALIZADO,
            pedido__criado_em__date=today,
            prato_id__isnull=False,
        ).aggregate(total=Sum("quantidade")).get("total")
        or 0
    )

    pedidos_em_producao_qs = Pedido.objects.filter(status=Pedido.Status.EM_PREPARO)
    pedidos_em_producao = (
        pedidos_em_producao_qs
        .prefetch_related("itens")
        .order_by("-criado_em", "-id")[:12]
    )
    pratos_em_producao = (
        ItemPedido.objects.filter(pedido__status=Pedido.Status.EM_PREPARO, prato_id__isnull=False)
        .values("nome_prato_snapshot")
        .annotate(total=Sum("quantidade"))
        .order_by("-total", "nome_prato_snapshot")
    )
    total_para_producao = sum(int(row["total"] or 0) for row in pratos_em_producao)

    pratos = []
    for row in pratos_em_producao:
        pratos.append(
            {
                "nome": row["nome_prato_snapshot"],
                "quantidade": int(row["total"] or 0),
            }
        )

    pedidos_cards = []
    for pedido in pedidos_em_producao:
        pratos_total = sum(int(item.quantidade or 0) for item in pedido.itens.all())
        tempo_base = pedido.producao_iniciada_em or pedido.criado_em
        elapsed_min = max(0, int((now - timezone.localtime(tempo_base)).total_seconds() // 60))
        pedidos_cards.append(
            {
                "id": pedido.id,
                "pedido_numero": pedido.numero,
                "cliente": pedido.nome_cliente,
                "criado_em": _format_local_datetime(pedido.criado_em, "%d/%m, %H:%M"),
                "icone_url": pedido.icone_pedido_url,
                "item_type_counts": pedido.item_type_counts,
                "item_lines": _pedido_item_lines(pedido),
                "tempo_producao": _tempo_producao_pedido(pedido),
                "pratos_total": pratos_total,
                "elapsed_min": elapsed_min,
                "turno_label": turno_label_pedido(pedido),
            }
        )

    return {
        "entregues_hoje": entregues_hoje,
        "total_para_producao": total_para_producao,
        "pratos_em_producao": pratos,
        "pedidos_cards": pedidos_cards,
        "pedidos_em_producao": pedidos_em_producao_qs.count(),
        "weekday_label": _weekday_label_pt(today),
        "date_label": today.strftime("%d/%m"),
    }


@staff_member_required(login_url="/admin/login/")
def cozinha_pedidos(request):
    payload = _cozinha_operacao_payload()
    return render(
        request,
        "pedidos/cozinha_operacao.html",
        {
            **payload,
        },
    )


def _prato_pronto_operacao_legado(request):
    _principal, turno = ensure_turnos_padrao()
    today = timezone.localdate()
    feedback = ""
    feedback_kind = "success"

    disponibilidade, _created = DisponibilidadeTurnoDia.objects.get_or_create(
        turno=turno,
        data=today,
    )
    if not disponibilidade.personalizado:
        disponibilidade = sincronizar_disponibilidade_automatica(disponibilidade)

    if request.method == "POST":
        action = _safe_text(request.POST.get("action"))
        try:
            if action == "save_turno":
                turno.ativo = request.POST.get("ativo") == "on"
                selected_days = [
                    day for day in TURNO_WEEKDAYS
                    if day in set(request.POST.getlist("dias_semana"))
                ]
                if not selected_days:
                    raise ValueError("Selecione pelo menos um dia da semana.")
                turno.dias_semana = ",".join(selected_days)
                turno.horario_inicio = _parse_optional_time(request.POST.get("horario_inicio"))
                turno.horario_fim = _parse_optional_time(request.POST.get("horario_fim"))
                turno.horario_limite_entrega = _parse_optional_time(request.POST.get("horario_limite_entrega"))
                if not turno.horario_inicio or not turno.horario_fim:
                    raise ValueError("Informe o inicio e o fim do turno.")
                if turno.horario_inicio >= turno.horario_fim:
                    raise ValueError("O fim do turno deve ser posterior ao inicio.")
                if (
                    turno.horario_limite_entrega
                    and not (turno.horario_inicio < turno.horario_limite_entrega <= turno.horario_fim)
                ):
                    raise ValueError("O limite de entrega deve ficar dentro do turno.")
                turno.permite_entrega = request.POST.get("permite_entrega") == "on"
                turno.permite_retirada = request.POST.get("permite_retirada") == "on"
                if not turno.permite_entrega and not turno.permite_retirada:
                    raise ValueError("Habilite entrega, retirada ou ambas.")
                modo_cardapio = _safe_text(request.POST.get("modo_cardapio"))
                if modo_cardapio not in dict(TurnoAtendimento.ModoCardapio.choices):
                    raise ValueError("Selecione uma origem valida para o cardapio.")
                turno.modo_cardapio = modo_cardapio
                selected_pratos = {
                    int(value)
                    for value in request.POST.getlist("pratos_padrao")
                    if str(value).isdigit()
                }
                if (
                    turno.modo_cardapio == TurnoAtendimento.ModoCardapio.FIXO
                    and not selected_pratos
                ):
                    raise ValueError("Selecione ao menos um prato fixo para esse modo.")
                turno.preco_padrao = (
                    money_decimal(request.POST.get("preco_padrao"))
                    if _safe_text(request.POST.get("preco_padrao"))
                    else None
                )
                turno.desconto_padrao = money_decimal(request.POST.get("desconto_padrao"))
                if turno.preco_padrao is not None and turno.preco_padrao <= 0:
                    raise ValueError("O preco padrao deve ser maior que zero.")
                if turno.desconto_padrao < 0:
                    raise ValueError("O desconto nao pode ser negativo.")
                turno.permite_promocao = request.POST.get("permite_promocao") == "on"
                turno.permite_cupom = request.POST.get("permite_cupom") == "on"
                turno.mensagem = _safe_text(request.POST.get("mensagem"))[:180]
                turno.orientacao_cliente = _safe_text(
                    request.POST.get("orientacao_cliente")
                )[:220]
                if not turno.orientacao_cliente:
                    raise ValueError(
                        "Informe como o Prato Pronto deve ser conservado e consumido."
                    )
                turno.save()

                TurnoPratoPadrao.objects.filter(turno=turno).exclude(prato_id__in=selected_pratos).delete()
                for ordem, prato in enumerate(
                    Prato.objects.filter(id__in=selected_pratos, ativo=True).order_by("nome"),
                    start=1,
                ):
                    TurnoPratoPadrao.objects.update_or_create(
                        turno=turno,
                        prato=prato,
                        defaults={"ativo": True, "ordem": ordem * 10},
                    )
                if not disponibilidade.personalizado:
                    sincronizar_disponibilidade_automatica(disponibilidade)
                return redirect(f"{request.path}?saved=turno")

            if action == "save_today":
                disponibilidade.personalizado = True
                disponibilidade.pausado = request.POST.get("pausado") == "on"
                disponibilidade.horario_inicio = _parse_optional_time(request.POST.get("horario_inicio_dia"))
                disponibilidade.horario_fim = _parse_optional_time(request.POST.get("horario_fim_dia"))
                disponibilidade.horario_limite_entrega = _parse_optional_time(
                    request.POST.get("horario_limite_entrega_dia")
                )
                inicio_dia = disponibilidade.horario_inicio or turno.horario_inicio
                fim_dia = disponibilidade.horario_fim or turno.horario_fim
                limite_entrega_dia = (
                    disponibilidade.horario_limite_entrega
                    or turno.horario_limite_entrega
                    or fim_dia
                )
                if not inicio_dia or not fim_dia or inicio_dia >= fim_dia:
                    raise ValueError("Confira os horarios da operacao de hoje.")
                if limite_entrega_dia and not (inicio_dia < limite_entrega_dia <= fim_dia):
                    raise ValueError("O limite de entrega de hoje deve ficar dentro do turno.")
                disponibilidade.permite_entrega = request.POST.get("permite_entrega_dia") == "on"
                disponibilidade.permite_retirada = request.POST.get("permite_retirada_dia") == "on"
                if not disponibilidade.permite_entrega and not disponibilidade.permite_retirada:
                    raise ValueError("Habilite entrega, retirada ou ambas para hoje.")
                disponibilidade.mensagem = _safe_text(request.POST.get("mensagem_dia"))[:180]

                selected_ids = {
                    int(value)
                    for value in request.POST.getlist("itens_hoje")
                    if str(value).isdigit()
                }
                if not selected_ids and not disponibilidade.pausado:
                    raise ValueError("Selecione ao menos um Prato Pronto ou pause o turno de hoje.")
                selected_pratos = list(
                    Prato.objects.filter(id__in=selected_ids, ativo=True).order_by("nome")
                )
                if selected_ids and not selected_pratos:
                    raise ValueError("Os pratos selecionados nao estao mais ativos.")

                with transaction.atomic():
                    itens_existentes = {
                        item.prato_id: item
                        for item in DisponibilidadeTurnoItem.objects.select_for_update().filter(
                            disponibilidade=disponibilidade
                        )
                    }
                    itens_preparados = []
                    for ordem, prato in enumerate(selected_pratos, start=1):
                        preco_raw = _safe_text(request.POST.get(f"preco_item_{prato.id}"))
                        preco = (
                            money_decimal(preco_raw)
                            if preco_raw
                            else preco_turno_prato(turno, prato)
                        )
                        if preco <= 0:
                            raise ValueError(
                                f"Informe um preco maior que zero para {prato.nome}."
                            )
                        estoque_raw = _safe_text(request.POST.get(f"estoque_item_{prato.id}"))
                        estoque = int(estoque_raw) if estoque_raw else None
                        if estoque is not None and estoque < 0:
                            raise ValueError(
                                f"O estoque de {prato.nome} nao pode ser negativo."
                            )
                        item = itens_existentes.get(prato.id)
                        if (
                            item
                            and estoque is not None
                            and estoque < item.quantidade_reservada
                        ):
                            raise ValueError(
                                f"O estoque de {prato.nome} nao pode ser menor que "
                                f"{item.quantidade_reservada}, ja reservado."
                            )
                        itens_preparados.append((ordem, prato, preco, estoque, item))

                    disponibilidade.save()
                    ids_ativos = {prato.id for _ordem, prato, *_rest in itens_preparados}
                    for ordem, prato, preco, estoque, item in itens_preparados:
                        if item is None:
                            item = DisponibilidadeTurnoItem(
                                disponibilidade=disponibilidade,
                                prato=prato,
                            )
                        item.preco = preco
                        item.quantidade_disponivel = estoque
                        item.ativo = True
                        item.ordem = ordem * 10
                        item.save()

                    for item in disponibilidade.itens.exclude(prato_id__in=ids_ativos):
                        if item.quantidade_reservada:
                            item.ativo = False
                            item.save(update_fields=["ativo", "atualizado_em"])
                        else:
                            item.delete()
                return redirect(f"{request.path}?saved=hoje")

            if action == "toggle_pause":
                disponibilidade.pausado = not disponibilidade.pausado
                disponibilidade.save(update_fields=["pausado", "atualizado_em"])
                return redirect(f"{request.path}?saved=pausa")

            if action == "reset_today":
                disponibilidade.personalizado = False
                disponibilidade.pausado = False
                disponibilidade.horario_inicio = None
                disponibilidade.horario_fim = None
                disponibilidade.horario_limite_entrega = None
                disponibilidade.permite_entrega = None
                disponibilidade.permite_retirada = None
                disponibilidade.mensagem = ""
                disponibilidade.save()
                sincronizar_disponibilidade_automatica(disponibilidade)
                return redirect(f"{request.path}?saved=padrao")
        except (TypeError, ValueError) as exc:
            feedback = str(exc)
            feedback_kind = "error"
            turno.refresh_from_db()
            disponibilidade.refresh_from_db()

    saved = _safe_text(request.GET.get("saved"))
    if saved and not feedback:
        feedback = {
            "turno": "Configuracao permanente salva.",
            "hoje": "Disponibilidade de hoje salva.",
            "pausa": "Situacao de hoje atualizada.",
            "padrao": "Hoje voltou a seguir a configuracao automatica.",
        }.get(saved, "Alteracoes salvas.")

    disponibilidade.refresh_from_db()
    itens_map = {
        item.prato_id: item
        for item in disponibilidade.itens.select_related("prato").all()
    }
    pratos_rows = []
    pratos_padrao_ids = set(
        TurnoPratoPadrao.objects.filter(turno=turno, ativo=True).values_list("prato_id", flat=True)
    )
    for prato in Prato.objects.filter(ativo=True).order_by("nome"):
        item = itens_map.get(prato.id)
        pratos_rows.append(
            {
                "prato": prato,
                "selected_today": bool(item and item.ativo),
                "selected_default": prato.id in pratos_padrao_ids,
                "preco_today": (
                    f"{item.preco:.2f}" if item else f"{preco_turno_prato(turno, prato):.2f}"
                ),
                "estoque_today": (
                    item.quantidade_disponivel
                    if item and item.quantidade_disponivel is not None
                    else ""
                ),
                "reservado": item.quantidade_reservada if item else 0,
                "restante": item.quantidade_restante if item else None,
            }
        )

    turno_context = resolve_turno_publico()
    weekday_choices = [
        ("seg", "Seg"),
        ("ter", "Ter"),
        ("qua", "Qua"),
        ("qui", "Qui"),
        ("sex", "Sex"),
        ("sab", "Sab"),
        ("dom", "Dom"),
    ]
    return render(
        request,
        "pedidos/prato_pronto_operacao.html",
        {
            "active": "prato_pronto",
            "turno": turno,
            "turno_days": turno.dias_semana_set,
            "modo_cardapio_choices": TurnoAtendimento.ModoCardapio.choices,
            "weekday_choices": weekday_choices,
            "disponibilidade": disponibilidade,
            "pratos_rows": pratos_rows,
            "feedback": feedback,
            "feedback_kind": feedback_kind,
            "turno_context": turno_context,
            "today": today,
        },
    )


def _duplicar_imagem_prato(source, target):
    if not source or not source.imagem:
        return False
    try:
        source.imagem.open("rb")
        image_content = ContentFile(source.imagem.read())
        source.imagem.close()
        target.imagem.save(
            Path(source.imagem.name).name,
            image_content,
            save=False,
        )
        return True
    except (FileNotFoundError, OSError, ValueError):
        target.imagem = None
        return False


@staff_member_required(login_url="/admin/login/")
def prato_pronto_operacao(request):
    _principal, turno = ensure_turnos_padrao()
    today = timezone.localdate()
    feedback = ""
    feedback_kind = "success"
    open_dish_modal = request.GET.get("new") == "1"
    open_general_modal = False
    edit_row = None

    turno.modo_cardapio = TurnoAtendimento.ModoCardapio.FIXO
    disponibilidade, _created = DisponibilidadeTurnoDia.objects.get_or_create(
        turno=turno,
        data=today,
    )
    if not disponibilidade.personalizado:
        sincronizar_disponibilidade_automatica(disponibilidade)

    edit_id = request.GET.get("edit")
    if edit_id and str(edit_id).isdigit():
        edit_row = get_object_or_404(
            TurnoPratoPadrao.objects.select_related("prato"),
            pk=edit_id,
            turno=turno,
        )
        open_dish_modal = True

    dish_form = PratoProntoForm(
        instance=edit_row.prato if edit_row else None,
        turno_prato=edit_row,
    )

    if request.method == "POST":
        action_values = request.POST.getlist("action")
        action = _safe_text(action_values[-1] if action_values else "")
        try:
            if action == "save_general":
                open_general_modal = True
                selected_days = [
                    day for day in TURNO_WEEKDAYS
                    if day in set(request.POST.getlist("dias_semana"))
                ]
                if not selected_days:
                    raise ValueError("Selecione pelo menos um dia de atendimento.")
                inicio = _parse_optional_time(request.POST.get("horario_inicio"))
                fim = _parse_optional_time(request.POST.get("horario_fim"))
                if not inicio or not fim:
                    raise ValueError("Informe o início e o fim do atendimento.")
                if inicio >= fim:
                    raise ValueError("O fim do atendimento deve ser posterior ao início.")

                turno.ativo = request.POST.get("ativo") == "on"
                turno.dias_semana = ",".join(selected_days)
                turno.horario_inicio = inicio
                turno.horario_fim = fim
                turno.horario_limite_entrega = fim
                turno.permite_entrega = True
                turno.permite_retirada = True
                turno.modo_cardapio = TurnoAtendimento.ModoCardapio.FIXO
                turno.preco_padrao = None
                turno.desconto_padrao = Decimal("0.00")
                turno.permite_promocao = False
                turno.permite_cupom = False
                turno.save()

                disponibilidade.personalizado = False
                disponibilidade.pausado = False
                disponibilidade.horario_inicio = None
                disponibilidade.horario_fim = None
                disponibilidade.horario_limite_entrega = None
                disponibilidade.permite_entrega = None
                disponibilidade.permite_retirada = None
                disponibilidade.mensagem = ""
                disponibilidade.save()
                sincronizar_disponibilidade_automatica(disponibilidade)
                return redirect(f"{request.path}?saved=general")

            if action == "close_now":
                live_context = resolve_turno_publico()
                if not live_context["is_open"]:
                    raise ValueError("O Prato Pronto não está aberto agora.")
                current_time = timezone.localtime().time().replace(microsecond=0)
                disponibilidade.personalizado = True
                disponibilidade.pausado = False
                disponibilidade.horario_fim = current_time
                disponibilidade.horario_limite_entrega = current_time
                disponibilidade.save()
                return redirect(f"{request.path}?saved=closed_now")

            if action == "postpone_today":
                if not turno_agendado_no_dia(turno, today):
                    raise ValueError("Não há abertura do Prato Pronto programada para hoje.")
                disponibilidade.personalizado = True
                disponibilidade.pausado = True
                disponibilidade.save(update_fields=["personalizado", "pausado", "atualizado_em"])
                return redirect(f"{request.path}?saved=postponed")

            if action == "resume_today":
                disponibilidade.personalizado = False
                disponibilidade.pausado = False
                disponibilidade.horario_inicio = None
                disponibilidade.horario_fim = None
                disponibilidade.horario_limite_entrega = None
                disponibilidade.permite_entrega = None
                disponibilidade.permite_retirada = None
                disponibilidade.mensagem = ""
                disponibilidade.save()
                sincronizar_disponibilidade_automatica(disponibilidade)
                return redirect(f"{request.path}?saved=resumed")

            if action == "save_dish":
                row_id = _safe_text(request.POST.get("turno_prato_id"))
                edit_row = (
                    get_object_or_404(
                        TurnoPratoPadrao.objects.select_related("prato"),
                        pk=row_id,
                        turno=turno,
                    )
                    if row_id.isdigit()
                    else None
                )
                source_id = _safe_text(request.POST.get("copiar_de"))
                source = (
                    Prato.objects.filter(
                        pk=source_id,
                        exclusivo_prato_pronto=False,
                    ).first()
                    if source_id.isdigit()
                    else None
                )
                form_data = request.POST.copy()
                if source and not edit_row:
                    defaults = {
                        "nome": source.nome,
                        "descricao": source.descricao,
                        "variacoes": source.variacoes,
                        "preco_prato_pronto": source.preco_site_resolvido or source.preco or "22.00",
                    }
                    for field_name, default_value in defaults.items():
                        if not _safe_text(form_data.get(field_name)):
                            form_data[field_name] = default_value

                form_instance = edit_row.prato if edit_row else None
                image_source = source if source and not edit_row else None
                if edit_row and not edit_row.prato.exclusivo_prato_pronto:
                    original = edit_row.prato
                    image_source = original
                    form_instance = Prato(
                        nome=original.nome,
                        descricao=original.descricao,
                        variacoes=original.variacoes,
                        imagem=original.imagem,
                        preco=original.preco,
                        preco_balcao=original.preco_balcao,
                        preco_site=original.preco_site,
                        preco_ifood=original.preco_ifood,
                        ativo=True,
                        exclusivo_prato_pronto=True,
                    )

                dish_form = PratoProntoForm(
                    form_data,
                    request.FILES,
                    instance=form_instance,
                    turno_prato=edit_row,
                )
                open_dish_modal = True
                if dish_form.is_valid():
                    with transaction.atomic():
                        prato = dish_form.save(commit=False)
                        if image_source and not request.FILES.get("imagem"):
                            _duplicar_imagem_prato(image_source, prato)
                        preco = dish_form.cleaned_data["preco_prato_pronto"]
                        prato.preco = preco
                        prato.preco_site = preco
                        prato.ativo = True
                        prato.exclusivo_prato_pronto = True
                        prato.save()
                        if edit_row:
                            edit_row.prato = prato
                            edit_row.preco = preco
                            edit_row.dias_semana = dish_form.cleaned_data["recorrencia"]
                            edit_row.ativo = dish_form.cleaned_data["ativo_prato_pronto"]
                            edit_row.save()
                        else:
                            last_order = (
                                TurnoPratoPadrao.objects.filter(turno=turno)
                                .aggregate(max_order=Max("ordem"))["max_order"]
                                or 0
                            )
                            edit_row = TurnoPratoPadrao.objects.create(
                                turno=turno,
                                prato=prato,
                                preco=preco,
                                dias_semana=dish_form.cleaned_data["recorrencia"],
                                ativo=dish_form.cleaned_data["ativo_prato_pronto"],
                                ordem=last_order + 10,
                            )
                        disponibilidade.personalizado = False
                        disponibilidade.save(update_fields=["personalizado", "atualizado_em"])
                        sincronizar_disponibilidade_automatica(disponibilidade)
                    return redirect(f"{request.path}?saved=dish")
                feedback = "Confira os campos destacados."
                feedback_kind = "error"

            if action == "toggle_dish":
                row = get_object_or_404(
                    TurnoPratoPadrao,
                    pk=request.POST.get("turno_prato_id"),
                    turno=turno,
                )
                row.ativo = not row.ativo
                row.save(update_fields=["ativo", "atualizado_em"])
                disponibilidade.personalizado = False
                disponibilidade.save(update_fields=["personalizado", "atualizado_em"])
                sincronizar_disponibilidade_automatica(disponibilidade)
                return redirect(f"{request.path}?saved=status")

            if action == "delete_dish":
                row = get_object_or_404(
                    TurnoPratoPadrao.objects.select_related("prato"),
                    pk=request.POST.get("turno_prato_id"),
                    turno=turno,
                )
                prato = row.prato
                with transaction.atomic():
                    row.delete()
                    if prato.exclusivo_prato_pronto and prato.ativo:
                        prato.ativo = False
                        prato.save(update_fields=["ativo"])
                    disponibilidade.personalizado = False
                    disponibilidade.save(update_fields=["personalizado", "atualizado_em"])
                    sincronizar_disponibilidade_automatica(disponibilidade)
                return redirect(f"{request.path}?saved=deleted")

            if action == "update_stock":
                row = get_object_or_404(
                    TurnoPratoPadrao.objects.select_related("prato"),
                    pk=request.POST.get("turno_prato_id"),
                    turno=turno,
                )
                if not row.ativo:
                    raise ValueError("Ative o prato antes de informar o estoque.")
                if row.dias_semana_set and TURNO_WEEKDAYS[today.weekday()] not in row.dias_semana_set:
                    raise ValueError("Este prato não está programado para hoje.")
                disponibilidade.personalizado = False
                disponibilidade.save(update_fields=["personalizado", "atualizado_em"])
                sincronizar_disponibilidade_automatica(disponibilidade)
                item = get_object_or_404(
                    DisponibilidadeTurnoItem,
                    disponibilidade=disponibilidade,
                    prato=row.prato,
                )
                stock_raw = _safe_text(request.POST.get("quantidade_disponivel"))
                stock = int(stock_raw) if stock_raw else None
                if stock is not None and stock < 0:
                    raise ValueError("O estoque não pode ser negativo.")
                if stock is not None and stock < item.quantidade_reservada:
                    raise ValueError(
                        f"O estoque não pode ser menor que {item.quantidade_reservada}, já reservado."
                    )
                item.quantidade_disponivel = stock
                item.save(update_fields=["quantidade_disponivel", "atualizado_em"])
                return redirect(f"{request.path}?saved=stock")
        except (TypeError, ValueError) as exc:
            feedback = str(exc)
            feedback_kind = "error"

    saved = _safe_text(request.GET.get("saved"))
    if saved and not feedback:
        feedback = {
            "general": "Configurações gerais salvas.",
            "closed_now": "Atendimento de hoje encerrado.",
            "postponed": "A abertura de hoje foi adiada.",
            "resumed": "A abertura de hoje foi retomada.",
            "dish": "Prato Pronto salvo.",
            "status": "Situação do prato atualizada.",
            "deleted": "Prato removido do Prato Pronto.",
            "stock": "Estoque de hoje atualizado.",
        }.get(saved, "Alterações salvas.")

    disponibilidade.refresh_from_db()
    itens_map = {
        item.prato_id: item
        for item in disponibilidade.itens.select_related("prato").all()
    }
    weekday_labels = dict(
        [
            ("seg", "Seg"),
            ("ter", "Ter"),
            ("qua", "Qua"),
            ("qui", "Qui"),
            ("sex", "Sex"),
            ("sab", "Sáb"),
            ("dom", "Dom"),
        ]
    )
    pratos_rows = []
    for row in (
        TurnoPratoPadrao.objects.select_related("prato")
        .filter(turno=turno)
        .order_by("ordem", "prato__nome")
    ):
        item = itens_map.get(row.prato_id)
        recurrence = [
            weekday_labels[day]
            for day in TURNO_WEEKDAYS
            if not row.dias_semana_set or day in row.dias_semana_set
        ]
        pratos_rows.append(
            {
                "config": row,
                "prato": row.prato,
                "preco_display": preco_turno_prato(turno, row.prato, row.preco),
                "recurrence_label": ", ".join(recurrence),
                "scheduled_today": (
                    not row.dias_semana_set
                    or TURNO_WEEKDAYS[today.weekday()] in row.dias_semana_set
                ),
                "estoque_today": (
                    item.quantidade_disponivel
                    if item and item.quantidade_disponivel is not None
                    else ""
                ),
                "reservado": item.quantidade_reservada if item else 0,
                "restante": item.quantidade_restante if item else None,
            }
        )

    copy_catalog = []
    for prato in Prato.objects.filter(exclusivo_prato_pronto=False).order_by("nome"):
        imagem_url = ""
        if prato.imagem:
            try:
                if prato.imagem.storage.exists(prato.imagem.name):
                    imagem_url = prato.imagem.url
            except (OSError, ValueError):
                imagem_url = ""
        copy_catalog.append(
            {
                "id": prato.id,
                "nome": prato.nome,
                "descricao": prato.descricao,
                "variacoes": prato.variacoes,
                "preco": prato.preco,
                "preco_site": prato.preco_site,
                "imagem_url": imagem_url,
            }
        )
    return render(
        request,
        "pedidos/prato_pronto_operacao.html",
        {
            "active": "prato_pronto",
            "turno": turno,
            "turno_days": turno.dias_semana_set,
            "weekday_choices": weekday_labels.items(),
            "disponibilidade": disponibilidade,
            "pratos_rows": pratos_rows,
            "feedback": feedback,
            "feedback_kind": feedback_kind,
            "turno_context": resolve_turno_publico(),
            "today": today,
            "dish_form": dish_form,
            "edit_row": edit_row,
            "open_dish_modal": open_dish_modal,
            "open_general_modal": open_general_modal,
            "copy_catalog": copy_catalog,
        },
    )


@staff_member_required(login_url="/admin/login/")
@require_GET
def api_cozinha_operacao(request):
    return JsonResponse(_cozinha_operacao_payload())


@staff_member_required(login_url="/admin/login/")
def pedidos_admin(request):
    base = Pedido.objects.prefetch_related("itens")
    hoje = timezone.localdate()
    pedidos_ativos = _marcar_clientes_recorrentes(base.exclude(
        status__in=[Pedido.Status.RASCUNHO, Pedido.Status.AGUARDANDO_APROVACAO, Pedido.Status.FINALIZADO, Pedido.Status.CANCELADO]
    ).order_by("-criado_em", "-id")[:20])
    return render(
        request,
        "pedidos/pedidos_admin.html",
        {
            "pedidos_ativos": pedidos_ativos,
            "bairros_sugestoes": RIO_VERDE_BAIRROS_OFICIAIS,
            "aprovacao_count": base.filter(status=Pedido.Status.AGUARDANDO_APROVACAO).count(),
            "concluidos_count": base.filter(status=Pedido.Status.FINALIZADO, criado_em__date=hoje).count(),
            "pedidos_badge": base.exclude(
                status__in=[Pedido.Status.RASCUNHO, Pedido.Status.AGUARDANDO_APROVACAO, Pedido.Status.FINALIZADO, Pedido.Status.CANCELADO]
            ).count(),
        },
    )


@staff_member_required(login_url="/admin/login/")
def pedidos_aprovacao_admin(request):
    base = Pedido.objects.prefetch_related("itens")
    hoje = timezone.localdate()
    pedidos_aprovacao = _marcar_clientes_recorrentes(
        base.filter(status=Pedido.Status.AGUARDANDO_APROVACAO).order_by("-criado_em", "-id")[:20]
    )
    return render(
        request,
        "pedidos/pedidos_aprovacao_admin.html",
        {
            "pedidos_aprovacao": pedidos_aprovacao,
            "bairros_sugestoes": RIO_VERDE_BAIRROS_OFICIAIS,
            "aprovacao_count": base.filter(status=Pedido.Status.AGUARDANDO_APROVACAO).count(),
            "concluidos_count": base.filter(status=Pedido.Status.FINALIZADO, criado_em__date=hoje).count(),
            "pedidos_badge": base.exclude(
                status__in=[Pedido.Status.RASCUNHO, Pedido.Status.AGUARDANDO_APROVACAO, Pedido.Status.FINALIZADO, Pedido.Status.CANCELADO]
            ).count(),
        },
    )


@staff_member_required(login_url="/admin/login/")
def pedidos_concluidos_admin(request):
    data_selecionada = parse_date(_safe_text(request.GET.get("data"))) or timezone.localdate()
    base = Pedido.objects.prefetch_related("itens")
    concluidos = base.filter(status=Pedido.Status.FINALIZADO, criado_em__date=data_selecionada)
    cancelados = base.filter(status=Pedido.Status.CANCELADO, criado_em__date=data_selecionada)
    total_concluidos_geral = base.filter(status=Pedido.Status.FINALIZADO).count()
    hoje = timezone.localdate()
    nav_inicio = data_selecionada - timedelta(days=3)
    nav_fim = data_selecionada + timedelta(days=3)
    nav_counts = {}
    nav_rows = (
        base.filter(
            status__in=[Pedido.Status.FINALIZADO, Pedido.Status.CANCELADO],
            criado_em__date__gte=nav_inicio,
            criado_em__date__lte=nav_fim,
        )
        .annotate(dia=TruncDate("criado_em"))
        .values("dia", "status")
        .annotate(total=Count("id"))
    )
    for row in nav_rows:
        nav_counts.setdefault(row["dia"], {})[row["status"]] = row["total"]

    date_nav_days = []
    for offset in range(7):
        dia = nav_inicio + timedelta(days=offset)
        counts = nav_counts.get(dia, {})
        date_nav_days.append(
            {
                "date": dia,
                "weekday": _weekday_label_pt(dia)[:3],
                "is_selected": dia == data_selecionada,
                "is_today": dia == hoje,
                "concluidos": counts.get(Pedido.Status.FINALIZADO, 0),
                "cancelados": counts.get(Pedido.Status.CANCELADO, 0),
            }
        )

    delta_dias = (data_selecionada - hoje).days
    if delta_dias == 0:
        rotulo_data_relativa = "Hoje"
    elif delta_dias == -1:
        rotulo_data_relativa = "Ontem"
    elif delta_dias == 1:
        rotulo_data_relativa = "Amanha"
    elif delta_dias < 0:
        rotulo_data_relativa = f"Ha {abs(delta_dias)} dias"
    else:
        rotulo_data_relativa = f"Em {delta_dias} dias"

    pedidos_concluidos = _marcar_clientes_recorrentes(concluidos.order_by("-criado_em", "-id"))
    pedidos_cancelados = _marcar_clientes_recorrentes(cancelados.order_by("-criado_em", "-id"))
    resumo_dia = concluidos.aggregate(total=Sum("total"), ultimo=Max("criado_em"))
    return render(
        request,
        "pedidos/pedidos_concluidos_admin.html",
        {
            "pedidos_concluidos": pedidos_concluidos,
            "pedidos_cancelados": pedidos_cancelados,
            "bairros_sugestoes": RIO_VERDE_BAIRROS_OFICIAIS,
            "aprovacao_count": base.filter(status=Pedido.Status.AGUARDANDO_APROVACAO).count(),
            "concluidos_count": concluidos.count(),
            "total_concluidos_geral": total_concluidos_geral,
            "cancelados_count": cancelados.count(),
            "data_selecionada": data_selecionada,
            "data_anterior": data_selecionada - timedelta(days=1),
            "data_proxima": data_selecionada + timedelta(days=1),
            "data_nav_days": date_nav_days,
            "rotulo_data_relativa": rotulo_data_relativa,
            "dia_semana_selecionado": _weekday_label_pt(data_selecionada),
            "total_concluidos_dia": resumo_dia.get("total") or Decimal("0.00"),
            "ultimo_concluido_em": resumo_dia.get("ultimo"),
            "hoje": hoje,
            "pedidos_badge": base.exclude(
                status__in=[Pedido.Status.RASCUNHO, Pedido.Status.AGUARDANDO_APROVACAO, Pedido.Status.FINALIZADO, Pedido.Status.CANCELADO]
            ).count(),
        },
    )


@staff_member_required(login_url="/admin/login/")
def clientes_admin(request):
    clientes = (
        Cliente.objects.annotate(pedidos_count=Count("pedidos"))
        .filter(pedidos_count__gt=0)
        .order_by("-ultimo_pedido_em", "nome")
    )
    return render(
        request,
        "pedidos/clientes_admin.html",
        {
            "active": "clientes",
            "clientes": clientes,
            "conflitos_abertos_count": ClienteTokenConflito.objects.filter(status=ClienteTokenConflito.Status.ABERTO).count(),
            **_pedidos_base_counts(Pedido.objects.all()),
        },
    )


@staff_member_required(login_url="/admin/login/")
def clientes_conflitos_admin(request):
    conflitos = (
        ClienteTokenConflito.objects.select_related("pedido")
        .prefetch_related("clientes")
        .filter(status=ClienteTokenConflito.Status.ABERTO)
        .order_by("-criado_em")
    )
    return render(
        request,
        "pedidos/clientes_conflitos_admin.html",
        {
            "active": "clientes",
            "conflitos": conflitos,
            **_pedidos_base_counts(Pedido.objects.all()),
        },
    )


@staff_member_required(login_url="/admin/login/")
@transaction.atomic
def cliente_detalhe_admin(request, cliente_id):
    cliente = get_object_or_404(Cliente.objects.prefetch_related("enderecos"), id=cliente_id)
    feedback = None
    feedback_kind = "success"
    if request.method == "POST":
        if not _user_can_manage_order_payment(request.user):
            feedback = "Usuario sem permissao para editar cliente."
            feedback_kind = "error"
        else:
            nome = _safe_text(request.POST.get("nome"))
            if nome:
                cliente.nome = nome
                cliente.nome_editado_manualmente = True
                cliente.save(update_fields=["nome", "nome_editado_manualmente", "atualizado_em"])
                feedback = "Cliente atualizado."
            else:
                feedback = "Informe um nome para o cliente."
                feedback_kind = "error"

    pedidos = cliente.pedidos.prefetch_related("itens").exclude(status=Pedido.Status.RASCUNHO).order_by("-criado_em", "-id")
    tokens = [pedido.public_token for pedido in pedidos if pedido.public_token]
    return render(
        request,
        "pedidos/cliente_detalhe_admin.html",
        {
            "active": "clientes",
            "cliente": cliente,
            "pedidos": pedidos,
            "tokens": tokens,
            "feedback": feedback,
            "feedback_kind": feedback_kind,
            "can_edit_client": _user_can_manage_order_payment(request.user),
            **_pedidos_base_counts(Pedido.objects.all()),
        },
    )


def _tempo_producao_pedido(pedido):
    if not pedido.producao_iniciada_em:
        return "--"
    return timesince(pedido.producao_iniciada_em, timezone.now())


def _format_local_datetime(value, fmt):
    return timezone.localtime(value).strftime(fmt)


def _pedido_primeiro_item_line(pedido):
    primeiro_item = next(iter(pedido.itens.all()), None)
    if not primeiro_item:
        return "Sem itens"
    item_line = f"{primeiro_item.quantidade}x {primeiro_item.nome_prato_snapshot}"
    if primeiro_item.variacao_nome_snapshot:
        item_line = f"{item_line} - {primeiro_item.variacao_nome_snapshot}"
    return item_line


def _pedido_item_lines(pedido):
    lines = []
    for item in pedido.itens.all():
        item_line = f"{item.quantidade}x {item.nome_prato_snapshot}"
        if item.variacao_nome_snapshot:
            item_line = f"{item_line} - {item.variacao_nome_snapshot}"
        lines.append(item_line)
    return lines or ["Sem itens"]


def _pedido_item_rows(pedido):
    rows = [
        {
            "nome": f"{item.quantidade}x {item.nome_prato_snapshot}",
            "variacao": item.variacao_nome_snapshot,
        }
        for item in pedido.itens.all()
    ]
    return rows or [{"nome": "Sem itens", "variacao": ""}]


def _marcar_clientes_recorrentes(pedidos):
    pedidos = list(pedidos)
    if not pedidos:
        return pedidos

    pedido_ids = [pedido.id for pedido in pedidos]
    max_criado_em = max(pedido.criado_em for pedido in pedidos if pedido.criado_em)
    cliente_ids = {pedido.cliente_id for pedido in pedidos if pedido.cliente_id}
    telefones = {normalize_phone(pedido.telefone) for pedido in pedidos}
    telefones.discard("")
    clientes_historico = {}
    telefones_historico = {}
    for pedido in pedidos:
        if pedido.cliente_id:
            clientes_historico.setdefault(pedido.cliente_id, []).append(pedido.criado_em)
        telefone_normalizado = normalize_phone(pedido.telefone)
        if telefone_normalizado:
            telefones_historico.setdefault(telefone_normalizado, []).append(pedido.criado_em)

    if cliente_ids:
        clientes_anteriores = (
            Pedido.objects.exclude(status__in=RECURRENT_CUSTOMER_EXCLUDED_STATUSES)
            .exclude(id__in=pedido_ids)
            .filter(cliente_id__in=cliente_ids, criado_em__lt=max_criado_em)
            .values_list("cliente_id", "criado_em")
        )
        for cliente_id, criado_em in clientes_anteriores:
            clientes_historico.setdefault(cliente_id, []).append(criado_em)

    if telefones:
        pedidos_anteriores = (
            Pedido.objects.exclude(status__in=RECURRENT_CUSTOMER_EXCLUDED_STATUSES)
            .exclude(id__in=pedido_ids)
            .filter(criado_em__lt=max_criado_em)
            .values_list("telefone", "criado_em")
        )
        for telefone, criado_em in pedidos_anteriores:
            telefone_normalizado = normalize_phone(telefone)
            if telefone_normalizado in telefones:
                telefones_historico.setdefault(telefone_normalizado, []).append(criado_em)

    for pedido in pedidos:
        telefone_normalizado = normalize_phone(pedido.telefone)
        datas_anteriores = [
            criado_em
            for criado_em in (
                clientes_historico.get(pedido.cliente_id, [])
                + telefones_historico.get(telefone_normalizado, [])
            )
            if criado_em < pedido.criado_em
        ]
        pedido_anterior_em = max(datas_anteriores, default=None)
        pedido.cliente_recorrente = pedido_anterior_em is not None
        pedido.cliente_recorrente_recente = bool(
            pedido_anterior_em
            and pedido.criado_em - pedido_anterior_em <= timedelta(hours=72)
        )
        pedido.cliente_recorrente_label = (
            "Pediu recentemente"
            if pedido.cliente_recorrente_recente
            else "Cliente recorrente"
        )
    return pedidos


def _pedido_admin_summary(pedido):
    return {
        "id": pedido.id,
        "numero": pedido.numero,
        "cliente": pedido.nome_cliente,
        "sem_telefone": not bool(normalize_phone(pedido.telefone)),
        "criado_em": _format_local_datetime(pedido.criado_em, "%d/%m, %H:%M"),
        "canal": pedido.canal,
        "canal_label": pedido.get_canal_display(),
        "cliente_recorrente": bool(getattr(pedido, "cliente_recorrente", False)),
        "cliente_recorrente_recente": bool(getattr(pedido, "cliente_recorrente_recente", False)),
        "cliente_recorrente_label": getattr(pedido, "cliente_recorrente_label", "Cliente recorrente"),
        "item_line": _pedido_primeiro_item_line(pedido),
        "item_lines": _pedido_item_lines(pedido),
        "item_rows": _pedido_item_rows(pedido),
        "status": pedido.status,
        "status_label": pedido.status_label_contextual,
        "pagamento_recebido": pedido.pagamento_recebido,
        "pagamento_recebido_em": _format_local_datetime(pedido.pagamento_recebido_em, "%H:%M") if pedido.pagamento_recebido_em else "",
        "tipo_coleta": pedido.tipo_coleta,
        "turno": pedido.turno.codigo if pedido.turno_id else TurnoAtendimento.Codigo.PRINCIPAL,
        "turno_label": turno_label_pedido(pedido),
        "orientacao_turno": pedido.orientacao_turno_snapshot,
        "stage_labels": pedido.stage_labels,
        "icone_url": pedido.icone_pedido_url,
        "tempo_producao": _tempo_producao_pedido(pedido),
        "entregador_solicitado": pedido.entregador_solicitado,
        "total": f"R$ {pedido.total:.2f}".replace(".", ","),
        "detail_url": reverse("pedidos:pedido_detalhe_admin", args=[pedido.id]),
        "copy_url": reverse("pedidos:api_pedido_copias", args=[pedido.id]),
    }


def _pedidos_base_counts(base):
    hoje = timezone.localdate()
    return {
        "aprovacao_count": base.filter(status=Pedido.Status.AGUARDANDO_APROVACAO).count(),
        "aprovacao_ids": list(
            base.filter(status=Pedido.Status.AGUARDANDO_APROVACAO).order_by("-criado_em", "-id").values_list("id", flat=True)[:20]
        ),
        "pedidos_badge": base.exclude(
            status__in=[Pedido.Status.RASCUNHO, Pedido.Status.AGUARDANDO_APROVACAO, Pedido.Status.FINALIZADO, Pedido.Status.CANCELADO]
        ).count(),
        "concluidos_count": base.filter(status=Pedido.Status.FINALIZADO, criado_em__date=hoje).count(),
    }


def _pedidos_admin_payload():
    base = Pedido.objects.prefetch_related("itens")
    pedidos = _marcar_clientes_recorrentes(base.exclude(
        status__in=[Pedido.Status.RASCUNHO, Pedido.Status.AGUARDANDO_APROVACAO, Pedido.Status.FINALIZADO, Pedido.Status.CANCELADO]
    ).order_by("-criado_em", "-id")[:20])
    return {
        "pedidos": [_pedido_admin_summary(pedido) for pedido in pedidos],
        **_pedidos_base_counts(base),
    }


@staff_member_required(login_url="/admin/login/")
@require_GET
def api_pedidos_admin(request):
    return JsonResponse(_pedidos_admin_payload())


def _pedidos_aprovacao_payload():
    base = Pedido.objects.prefetch_related("itens")
    pedidos = _marcar_clientes_recorrentes(
        base.filter(status=Pedido.Status.AGUARDANDO_APROVACAO).order_by("-criado_em", "-id")[:20]
    )
    return {
        "pedidos": [_pedido_admin_summary(pedido) for pedido in pedidos],
        **_pedidos_base_counts(base),
    }


@staff_member_required(login_url="/admin/login/")
@require_GET
def api_pedidos_aprovacao_admin(request):
    return JsonResponse(_pedidos_aprovacao_payload())


def _pedidos_concluidos_payload(data_selecionada=None):
    data_selecionada = data_selecionada or timezone.localdate()
    base = Pedido.objects.prefetch_related("itens")
    concluidos = base.filter(status=Pedido.Status.FINALIZADO, criado_em__date=data_selecionada).order_by("-criado_em", "-id")
    cancelados = base.filter(status=Pedido.Status.CANCELADO, criado_em__date=data_selecionada).order_by("-criado_em", "-id")
    concluidos_marcados = _marcar_clientes_recorrentes(concluidos)
    cancelados_marcados = _marcar_clientes_recorrentes(cancelados)
    resumo_dia = concluidos.aggregate(total=Sum("total"), ultimo=Max("criado_em"))
    ultimo_concluido = resumo_dia.get("ultimo")
    return {
        **_pedidos_base_counts(base),
        "pedidos_concluidos": [_pedido_admin_summary(pedido) for pedido in concluidos_marcados],
        "pedidos_cancelados": [_pedido_admin_summary(pedido) for pedido in cancelados_marcados],
        "concluidos_count": concluidos.count(),
        "total_concluidos_geral": base.filter(status=Pedido.Status.FINALIZADO).count(),
        "cancelados_count": cancelados.count(),
        "data_label": data_selecionada.strftime("%d/%m/%Y"),
        "total_concluidos_dia": f"R$ {(resumo_dia.get('total') or Decimal('0.00')):.2f}".replace(".", ","),
        "ultimo_concluido_label": _format_local_datetime(ultimo_concluido, "%H:%M") if ultimo_concluido else "-",
    }


@staff_member_required(login_url="/admin/login/")
@require_GET
def api_pedidos_concluidos_admin(request):
    data_selecionada = parse_date(_safe_text(request.GET.get("data"))) or timezone.localdate()
    return JsonResponse(_pedidos_concluidos_payload(data_selecionada))


@staff_member_required(login_url="/admin/login/")
@require_GET
def api_pedido_copias(request, pedido_id):
    pedido = get_object_or_404(Pedido.objects.prefetch_related("itens"), id=pedido_id)
    return JsonResponse(
        {
            "cliente": montar_mensagem_whatsapp(pedido),
            "entregador": montar_mensagem_entregador(pedido),
            "confirmacao": montar_mensagem_confirmacao_pedido(pedido),
        }
    )


def _serialize_cliente_endereco(endereco):
    return {
        "id": endereco.id,
        "endereco": endereco.endereco,
        "endereco_formatado": endereco.endereco_formatado,
        "rua": endereco.rua,
        "numero_endereco": endereco.numero_endereco,
        "bairro": endereco.bairro,
        "cidade": endereco.cidade,
        "estado": endereco.estado,
        "complemento": endereco.complemento,
        "lote_quadra": endereco.lote_quadra,
        "ponto_referencia": endereco.ponto_referencia,
        "latitude": str(endereco.latitude) if endereco.latitude is not None else "",
        "longitude": str(endereco.longitude) if endereco.longitude is not None else "",
        "ultimo_uso_em": endereco.ultimo_uso_em.isoformat() if endereco.ultimo_uso_em else "",
    }


@staff_member_required(login_url="/admin/login/")
@require_GET
def api_cliente_enderecos_por_telefone(request):
    telefone_normalizado = normalize_phone(request.GET.get("telefone"))
    if not telefone_normalizado:
        return JsonResponse({"ok": True, "cliente": None, "enderecos": []})

    cliente = Cliente.objects.filter(telefone_normalizado=telefone_normalizado).first()
    if not cliente:
        return JsonResponse({"ok": True, "cliente": None, "enderecos": []})

    enderecos = EnderecoCliente.objects.filter(cliente=cliente).order_by("-ultimo_uso_em", "-id")[:5]
    return JsonResponse(
        {
            "ok": True,
            "cliente": {
                "id": cliente.id,
                "nome": cliente.nome,
                "telefone": cliente.telefone,
            },
            "enderecos": [_serialize_cliente_endereco(endereco) for endereco in enderecos],
        }
    )


@staff_member_required(login_url="/admin/login/")
@require_POST
def registrar_pedido_lista_impressao(request, pedido_id):
    pedido = get_object_or_404(Pedido, id=pedido_id)
    item = PedidoListaImpressao.objects.create(
        pedido=pedido,
        numero=pedido.numero,
        nome_cliente=pedido.nome_cliente,
        public_token=pedido.public_token,
    )
    return JsonResponse({"ok": True, "id": item.id})


@staff_member_required(login_url="/admin/login/")
def pedido_detalhe_admin(request, pedido_id):
    pedido = get_object_or_404(Pedido.objects.prefetch_related("itens"), id=pedido_id)
    context = _pedido_detail_context(request, pedido, is_new_order=pedido.status == Pedido.Status.RASCUNHO)
    if request.headers.get("x-requested-with") == "XMLHttpRequest":
        return render(request, "pedidos/_pedido_detail_modal_content.html", context)
    return render(request, "pedidos/pedido_detalhe_admin.html", context)


def _pedido_detail_context(request, pedido, *, is_new_order=False):
    itens_subtotal = pedido.itens.aggregate(total=Sum("subtotal")).get("total") or Decimal("0.00")
    frete_esperado, faixa_frete_atual = _calcular_frete_por_distancia(pedido.distancia_km)
    total_recalculado = itens_subtotal + pedido.valor_frete - pedido.promocao_desconto - pedido.cupom_desconto
    diferenca_frete = Decimal("0.00") if pedido.frete_gratis else pedido.valor_frete - frete_esperado

    context = {
        "active": "pedidos",
        "pedidos_badge": Pedido.objects.exclude(
            status__in=[Pedido.Status.RASCUNHO, Pedido.Status.FINALIZADO, Pedido.Status.CANCELADO]
        ).count(),
        "pedido": pedido,
        "itens_subtotal": itens_subtotal,
        "frete_esperado": frete_esperado,
        "faixa_frete_atual": faixa_frete_atual,
        "diferenca_frete": diferenca_frete,
        "frete_confere": diferenca_frete == Decimal("0.00"),
        "total_recalculado": total_recalculado,
        "total_confere": total_recalculado == pedido.total,
        "can_edit_payment": _user_can_manage_order_payment(request.user),
        "payment_choices": Pedido.FormaPagamento.choices,
        "canal_choices": Pedido.Canal.choices,
        "terminal_choices": TerminalCaixa.objects.filter(ativo=True).order_by("ordem", "nome"),
        "bairros_sugestoes": RIO_VERDE_BAIRROS_OFICIAIS,
        "is_new_order": is_new_order,
    }
    return context


def _pedido_modal_payload(pedido):
    pedido = Pedido.objects.prefetch_related("itens").get(pk=pedido.pk)
    itens_subtotal = pedido.itens.aggregate(total=Sum("subtotal")).get("total") or Decimal("0.00")
    frete_esperado, _faixa_frete_atual = _calcular_frete_por_distancia(pedido.distancia_km)
    diferenca_frete = Decimal("0.00") if pedido.frete_gratis else pedido.valor_frete - frete_esperado
    total_recalculado = itens_subtotal + pedido.valor_frete - pedido.promocao_desconto - pedido.cupom_desconto
    return {
        "ok": True,
        "pedido": {
            "id": pedido.id,
            "numero": pedido.numero,
            "nome_cliente": pedido.nome_cliente,
            "telefone": pedido.telefone,
            "cliente_nome": pedido.cliente.nome if pedido.cliente_id and pedido.cliente else "",
            "tipo_coleta": pedido.tipo_coleta,
            "tipo_coleta_label": pedido.get_tipo_coleta_display(),
            "forma_pagamento": pedido.forma_pagamento,
            "forma_pagamento_label": pedido.get_forma_pagamento_display(),
            "enviar_talheres": "sim" if pedido.enviar_talheres else "nao",
            "enviar_talheres_label": "Enviar" if pedido.enviar_talheres else "Nao enviar",
            "canal": pedido.canal,
            "canal_label": pedido.get_canal_display(),
            "terminal": str(pedido.terminal_id or ""),
            "terminal_label": pedido.terminal.nome if pedido.terminal_id and pedido.terminal else "",
            "ifood": "sim" if pedido.ifood else "nao",
            "ifood_label": "iFood" if pedido.ifood else "Balcao/PD",
            "observacao_geral": pedido.observacao_geral,
            "status": pedido.status,
            "status_label": pedido.status_label_contextual,
            "pagamento_recebido": pedido.pagamento_recebido,
            "pagamento_recebido_em": _format_local_datetime(pedido.pagamento_recebido_em, "%H:%M") if pedido.pagamento_recebido_em else "",
            "endereco": pedido.endereco,
            "google_maps_route_url": pedido.google_maps_route_url,
            "valor_frete": f"R$ {pedido.valor_frete:.2f}".replace(".", ","),
            "valor_frete_label": _pedido_frete_display_value(pedido),
            "frete_gratis": pedido.frete_gratis,
            "distancia_km": f"{pedido.distancia_km:.2f}".replace(".", ","),
            "distancia_label": _pedido_distancia_display_value(pedido),
            "itens_subtotal": f"R$ {itens_subtotal:.2f}".replace(".", ","),
            "total": f"R$ {pedido.total:.2f}".replace(".", ","),
            "cupom_codigo": pedido.cupom_codigo,
            "cupom_desconto": f"R$ {pedido.cupom_desconto:.2f}".replace(".", ","),
            "promocao_descricao": pedido.promocao_descricao,
            "promocao_desconto": f"R$ {pedido.promocao_desconto:.2f}".replace(".", ","),
            "frete_confere": diferenca_frete == Decimal("0.00"),
            "total_confere": total_recalculado == pedido.total,
        },
        "itens": [
            {
                "tipo": (
                    "prato_pronto"
                    if item.disponibilidade_turno_item_id
                    else "prato"
                    if item.prato_id
                    else "bebida"
                    if item.bebida_id
                    else "adicional"
                    if item.adicional_id
                    else ""
                ),
                "item_id": item.prato_id or item.bebida_id or item.adicional_id or "",
                "disponibilidade_turno_item_id": item.disponibilidade_turno_item_id or "",
                "nome": item.nome_prato_snapshot,
                "variacao": item.variacao_nome_snapshot,
                "quantidade": item.quantidade,
                "observacao": item.observacao,
                "classificacao_saida": item.classificacao_saida,
                "classificacao_saida_label": item.get_classificacao_saida_display(),
                "subtotal": f"R$ {item.subtotal:.2f}".replace(".", ","),
            }
            for item in pedido.itens.all()
        ],
    }


@staff_member_required(login_url="/admin/login/")
@require_GET
def pedido_novo_admin(request):
    canal = _safe_text(request.GET.get("canal")) or Pedido.Canal.BALCAO
    if canal not in dict(Pedido.Canal.choices):
        canal = Pedido.Canal.BALCAO
    pedido = Pedido.objects.create(
        nome_cliente="Cliente",
        telefone="",
        rua="",
        numero_endereco="",
        bairro="",
        cidade="Rio Verde",
        estado="GO",
        endereco_formatado="Retirada no local",
        endereco="Retirada no local",
        tipo_coleta=Pedido.TipoColeta.RETIRADA,
        forma_pagamento=Pedido.FormaPagamento.IFOOD if canal == Pedido.Canal.IFOOD else Pedido.FormaPagamento.DINHEIRO,
        enviar_talheres=False,
        canal=canal,
        ifood=canal == Pedido.Canal.IFOOD,
        status=Pedido.Status.RASCUNHO,
        valor_frete=Decimal("0.00"),
        distancia_km=Decimal("0.00"),
    )
    context = _pedido_detail_context(request, pedido, is_new_order=True)
    return render(request, "pedidos/_pedido_detail_modal_content.html", context)


@staff_member_required(login_url="/admin/login/")
@require_POST
@transaction.atomic
def finalizar_pedido_novo_admin(request, pedido_id):
    if not _user_can_manage_order_payment(request.user):
        return HttpResponseBadRequest("Usuario sem permissao para criar pedido.")
    pedido = get_object_or_404(Pedido.objects.prefetch_related("itens"), id=pedido_id, status=Pedido.Status.RASCUNHO)
    if not pedido.itens.exists():
        return HttpResponseBadRequest("Adicione pelo menos um item.")
    if pedido.tipo_coleta == Pedido.TipoColeta.ENTREGA and not _safe_text(pedido.endereco):
        return HttpResponseBadRequest("Informe o endereco de entrega.")
    try:
        recalculate_order_totals(pedido, cupom_codigo=pedido.cupom_codigo)
    except ValueError as exc:
        transaction.set_rollback(True)
        return HttpResponseBadRequest(str(exc))
    pedido.status = Pedido.Status.EM_PREPARO
    pedido.save(update_fields=["status"])
    sync_movimentacao_caixa_pedido(pedido, request.user)
    sync_customer_from_order(pedido)
    return JsonResponse(
        {
            "ok": True,
            "id": pedido.id,
            "detail_url": reverse("pedidos:pedido_detalhe_admin", args=[pedido.id]),
        }
    )


def _clone_order_as_draft(pedido):
    clone = Pedido.objects.create(
        nome_cliente=pedido.nome_cliente,
        telefone=pedido.telefone,
        cliente=pedido.cliente,
        rua=pedido.rua,
        numero_endereco=pedido.numero_endereco,
        bairro=pedido.bairro,
        cidade=pedido.cidade,
        estado=pedido.estado,
        endereco_formatado=pedido.endereco_formatado,
        latitude=pedido.latitude,
        longitude=pedido.longitude,
        endereco=pedido.endereco,
        complemento=pedido.complemento,
        lote_quadra=pedido.lote_quadra,
        ponto_referencia=pedido.ponto_referencia,
        tipo_coleta=pedido.tipo_coleta,
        forma_pagamento=pedido.forma_pagamento,
        terminal=pedido.terminal,
        enviar_talheres=pedido.enviar_talheres,
        canal=pedido.canal,
        ifood=pedido.ifood,
        observacao_geral=pedido.observacao_geral,
        status=Pedido.Status.RASCUNHO,
        distancia_km=pedido.distancia_km,
        valor_frete=pedido.valor_frete,
        frete_gratis=pedido.frete_gratis,
        cupom=pedido.cupom,
        cupom_codigo=pedido.cupom_codigo,
    )
    for item in pedido.itens.all():
        ItemPedido.objects.create(
            pedido=clone,
            prato=item.prato,
            bebida=item.bebida,
            adicional=item.adicional,
            nome_prato_snapshot=item.nome_prato_snapshot,
            variacao_nome_snapshot=item.variacao_nome_snapshot,
            preco_snapshot=item.preco_snapshot,
            quantidade=item.quantidade,
            observacao=item.observacao,
            classificacao_saida=item.classificacao_saida,
        )
    recalculate_order_totals(clone, cupom_codigo=clone.cupom_codigo)
    return clone


@staff_member_required(login_url="/admin/login/")
@require_POST
@transaction.atomic
def duplicar_pedido_admin(request, pedido_id):
    if not _user_can_manage_order_payment(request.user):
        return HttpResponseBadRequest("Usuario sem permissao para duplicar pedido.")
    pedido = get_object_or_404(Pedido.objects.prefetch_related("itens"), id=pedido_id)
    clone = _clone_order_as_draft(pedido)
    detail_url = reverse("pedidos:pedido_detalhe_admin", args=[clone.id])
    if request.headers.get("x-requested-with") == "XMLHttpRequest":
        return JsonResponse({"ok": True, "id": clone.id, "detail_url": detail_url})
    return redirect("pedidos:pedido_detalhe_admin", pedido_id=clone.id)


@staff_member_required(login_url="/admin/login/")
@require_POST
@transaction.atomic
def excluir_pedido_admin(request, pedido_id):
    if not _user_can_manage_order_payment(request.user):
        return HttpResponseBadRequest("Usuario sem permissao para excluir pedido.")
    pedido = get_object_or_404(Pedido, id=pedido_id)
    if pedido.disponibilidade_turno_id and not pedido.estoque_turno_liberado:
        liberar_estoque_pedido(pedido)
    excluir_movimentacao_caixa_pedido(pedido, request.user)
    pedido.delete()
    if request.headers.get("x-requested-with") == "XMLHttpRequest":
        return JsonResponse({"ok": True})
    return redirect("pedidos:cozinha_pedidos")


@staff_member_required(login_url="/admin/login/")
@require_POST
def atualizar_pagamento_pedido(request, pedido_id):
    if not _user_can_manage_order_payment(request.user):
        return HttpResponseBadRequest("Usuario sem permissao para alterar pagamento.")
    pedido = get_object_or_404(Pedido, id=pedido_id)
    forma_pagamento = request.POST.get("forma_pagamento")
    if forma_pagamento not in dict(Pedido.FormaPagamento.choices):
        return HttpResponseBadRequest("Forma de pagamento invalida.")
    pedido.forma_pagamento = forma_pagamento
    update_fields = ["forma_pagamento"]
    if forma_pagamento == Pedido.FormaPagamento.IFOOD:
        pedido.canal = Pedido.Canal.IFOOD
        pedido.ifood = True
        update_fields.extend(["canal", "ifood"])
    elif pedido.canal == Pedido.Canal.IFOOD:
        pedido.canal = Pedido.Canal.BALCAO
        pedido.ifood = False
        update_fields.extend(["canal", "ifood"])
    pedido.save(update_fields=update_fields)
    sync_movimentacao_caixa_pedido(pedido, request.user)
    return JsonResponse(_pedido_modal_payload(pedido))


@staff_member_required(login_url="/admin/login/")
@require_POST
@transaction.atomic
def atualizar_cupom_pedido(request, pedido_id):
    if not _user_can_manage_order_payment(request.user):
        return HttpResponseBadRequest("Usuario sem permissao para editar pedido.")
    pedido = get_object_or_404(Pedido, id=pedido_id)
    cupom_codigo = normalize_coupon_code(request.POST.get("cupom_codigo"))
    try:
        recalculate_order_totals(pedido, cupom_codigo=cupom_codigo)
        sync_movimentacao_caixa_pedido(pedido, request.user)
    except ValueError as exc:
        transaction.set_rollback(True)
        return HttpResponseBadRequest(str(exc))
    pedido.refresh_from_db()
    return JsonResponse(_pedido_modal_payload(pedido))


@staff_member_required(login_url="/admin/login/")
@require_GET
def api_catalogo_editor_pedido(request):
    _principal, turno = ensure_turnos_padrao()
    disponibilidade, _created = DisponibilidadeTurnoDia.objects.get_or_create(
        turno=turno,
        data=timezone.localdate(),
    )
    if not disponibilidade.personalizado:
        disponibilidade = sincronizar_disponibilidade_automatica(
            disponibilidade
        )
    turno_context = resolve_turno_publico()
    return JsonResponse(
        serialize_editor_catalog(
            disponibilidade_prato_pronto=disponibilidade,
            prato_pronto_open=turno_context["is_open"],
        )
    )


@staff_member_required(login_url="/admin/login/")
@require_POST
@transaction.atomic
def atualizar_itens_pedido(request, pedido_id):
    if not _user_can_manage_order_payment(request.user):
        return HttpResponseBadRequest("Usuario sem permissao para editar pedido.")
    pedido = get_object_or_404(Pedido.objects.prefetch_related("itens"), id=pedido_id)
    try:
        payload = json.loads(request.POST.get("itens_payload") or "[]")
    except json.JSONDecodeError:
        return HttpResponseBadRequest("Itens invalidos.")
    if not payload:
        return HttpResponseBadRequest("O pedido precisa ter pelo menos um item.")
    try:
        special_items = [
            item for item in payload
            if _safe_text(item.get("tipo")) == "prato_pronto"
        ]
        regular_dishes = [
            item for item in payload
            if _safe_text(item.get("tipo")) == "prato"
        ]
        if special_items and regular_dishes:
            raise ValueError(
                "Separe os Pratos Prontos dos pratos do atendimento principal."
            )

        if special_items:
            turno_context = resolve_turno_publico()
            turno = turno_context["prato_pronto"]
            disponibilidade, _created = DisponibilidadeTurnoDia.objects.get_or_create(
                turno=turno,
                data=timezone.localdate(),
            )
            if not disponibilidade.personalizado:
                disponibilidade = sincronizar_disponibilidade_automatica(
                    disponibilidade
                )
            available_by_prato = {
                item.prato_id: item
                for item in disponibilidade.itens.filter(
                    ativo=True,
                    prato__ativo=True,
                )
            }
            for item in special_items:
                try:
                    prato_id = int(item.get("item_id"))
                except (TypeError, ValueError):
                    raise ValueError("Um dos Pratos Prontos é inválido.")
                disponibilidade_item = available_by_prato.get(prato_id)
                if not disponibilidade_item:
                    raise ValueError(
                        "Um dos Pratos Prontos não está disponível hoje."
                    )
                item["disponibilidade_turno_item_id"] = disponibilidade_item.id
            pedido.turno = turno
            pedido.disponibilidade_turno = disponibilidade
            pedido.orientacao_turno_snapshot = turno.orientacao_cliente
            pedido.estoque_turno_liberado = False
            pedido.save(
                update_fields=[
                    "turno",
                    "disponibilidade_turno",
                    "orientacao_turno_snapshot",
                    "estoque_turno_liberado",
                ]
            )
        elif pedido.turno_id:
            pedido.turno = None
            pedido.disponibilidade_turno = None
            pedido.orientacao_turno_snapshot = ""
            pedido.estoque_turno_liberado = False
            pedido.save(
                update_fields=[
                    "turno",
                    "disponibilidade_turno",
                    "orientacao_turno_snapshot",
                    "estoque_turno_liberado",
                ]
            )

        replace_order_items(pedido, payload)
        sync_movimentacao_caixa_pedido(pedido, request.user)
    except ValueError as exc:
        transaction.set_rollback(True)
        return HttpResponseBadRequest(str(exc))
    return JsonResponse(_pedido_modal_payload(pedido))


@staff_member_required(login_url="/admin/login/")
@require_POST
@transaction.atomic
def atualizar_dados_pedido(request, pedido_id):
    if not _user_can_manage_order_payment(request.user):
        return HttpResponseBadRequest("Usuario sem permissao para editar pedido.")
    pedido = get_object_or_404(Pedido, id=pedido_id)
    field = _safe_text(request.POST.get("field"))
    if field == "nome_cliente":
        pedido.nome_cliente = _safe_text(request.POST.get("value")) or pedido.nome_cliente
        pedido.save(update_fields=["nome_cliente"])
        sync_movimentacao_caixa_pedido(pedido, request.user)
        sync_customer_from_order(pedido)
    elif field == "telefone":
        pedido.telefone = normalize_phone(request.POST.get("value"))
        update_fields = ["telefone"]
        if pedido.status == Pedido.Status.RASCUNHO:
            cliente = Cliente.objects.filter(telefone_normalizado=normalize_phone(pedido.telefone)).first()
            if cliente and pedido.nome_cliente.strip().casefold() in {"", "cliente"}:
                pedido.nome_cliente = cliente.nome
                update_fields.append("nome_cliente")
        pedido.save(update_fields=update_fields)
        sync_movimentacao_caixa_pedido(pedido, request.user)
        sync_customer_from_order(pedido)
    elif field == "enviar_talheres":
        pedido.enviar_talheres = request.POST.get("value") == "sim"
        pedido.save(update_fields=["enviar_talheres"])
    elif field == "ifood":
        pedido.ifood = request.POST.get("value") == "sim"
        pedido.canal = Pedido.Canal.IFOOD if pedido.ifood else Pedido.Canal.BALCAO
        update_fields = ["ifood", "canal"]
        if pedido.ifood:
            pedido.forma_pagamento = Pedido.FormaPagamento.IFOOD
            update_fields.append("forma_pagamento")
        elif pedido.forma_pagamento == Pedido.FormaPagamento.IFOOD:
            pedido.forma_pagamento = Pedido.FormaPagamento.DINHEIRO
            update_fields.append("forma_pagamento")
        pedido.save(update_fields=update_fields)
        try:
            reprice_order_items_from_catalog(pedido)
            recalculate_order_totals(pedido)
            sync_movimentacao_caixa_pedido(pedido, request.user)
        except ValueError as exc:
            transaction.set_rollback(True)
            return HttpResponseBadRequest(str(exc))
    elif field == "canal":
        canal = _safe_text(request.POST.get("value"))
        if canal not in dict(Pedido.Canal.choices):
            return HttpResponseBadRequest("Canal invalido.")
        pedido.canal = canal
        pedido.ifood = canal == Pedido.Canal.IFOOD
        update_fields = ["canal", "ifood"]
        if pedido.ifood:
            pedido.forma_pagamento = Pedido.FormaPagamento.IFOOD
            update_fields.append("forma_pagamento")
        elif pedido.forma_pagamento == Pedido.FormaPagamento.IFOOD:
            pedido.forma_pagamento = Pedido.FormaPagamento.DINHEIRO
            update_fields.append("forma_pagamento")
        pedido.save(update_fields=update_fields)
        try:
            reprice_order_items_from_catalog(pedido)
            recalculate_order_totals(pedido)
            sync_movimentacao_caixa_pedido(pedido, request.user)
        except ValueError as exc:
            transaction.set_rollback(True)
            return HttpResponseBadRequest(str(exc))
    elif field == "terminal":
        terminal = TerminalCaixa.objects.filter(id=request.POST.get("value"), ativo=True).first()
        if not terminal:
            return HttpResponseBadRequest("Terminal invalido.")
        pedido.terminal = terminal
        pedido.save(update_fields=["terminal"])
        sync_movimentacao_caixa_pedido(pedido, request.user)
    elif field == "frete_gratis":
        frete_gratis = _post_bool(request.POST.get("value"))
        pedido.frete_gratis = frete_gratis
        if frete_gratis:
            pedido.valor_frete = Decimal("0.00")
        elif pedido.distancia_km > 0:
            pedido.valor_frete, _ = _calcular_frete_por_distancia(pedido.distancia_km)
        else:
            pedido.valor_frete = Decimal("0.00")
        pedido.save(update_fields=["frete_gratis", "valor_frete"])
        recalculate_order_totals(pedido)
        sync_movimentacao_caixa_pedido(pedido, request.user)
    elif field == "tipo_coleta":
        tipo_coleta = _safe_text(request.POST.get("value"))
        if tipo_coleta not in dict(Pedido.TipoColeta.choices):
            return HttpResponseBadRequest("Tipo de coleta invalido.")
        if tipo_coleta == Pedido.TipoColeta.ENTREGA:
            return HttpResponseBadRequest("Informe o endereco de entrega para alterar o tipo.")
        pedido.tipo_coleta = Pedido.TipoColeta.RETIRADA
        pedido.rua = ""
        pedido.numero_endereco = ""
        pedido.bairro = ""
        pedido.cidade = "Rio Verde"
        pedido.estado = "GO"
        pedido.endereco_formatado = "Retirada no local"
        pedido.endereco = "Retirada no local"
        pedido.latitude = None
        pedido.longitude = None
        pedido.complemento = ""
        pedido.lote_quadra = ""
        pedido.ponto_referencia = ""
        pedido.distancia_km = Decimal("0.00")
        pedido.valor_frete = Decimal("0.00")
        pedido.frete_gratis = False
        if pedido.status == Pedido.Status.SAIU_ENTREGA:
            pedido.status = Pedido.Status.FINALIZADO
        pedido.save(update_fields=[
            "tipo_coleta",
            "rua",
            "numero_endereco",
            "bairro",
            "cidade",
            "estado",
            "endereco_formatado",
            "endereco",
            "latitude",
            "longitude",
            "complemento",
            "lote_quadra",
            "ponto_referencia",
            "distancia_km",
            "valor_frete",
            "frete_gratis",
            "status",
        ])
        recalculate_order_totals(pedido)
        sync_movimentacao_caixa_pedido(pedido, request.user)
        sync_customer_from_order(pedido)
    elif field == "observacao_geral":
        pedido.observacao_geral = _safe_text(request.POST.get("value"))
        pedido.save(update_fields=["observacao_geral"])
    else:
        return HttpResponseBadRequest("Campo invalido.")
    return JsonResponse(_pedido_modal_payload(pedido))


@staff_member_required(login_url="/admin/login/")
@require_POST
@transaction.atomic
def atualizar_entrega_pedido(request, pedido_id):
    if not _user_can_manage_order_payment(request.user):
        return HttpResponseBadRequest("Usuario sem permissao para editar pedido.")
    pedido = get_object_or_404(Pedido, id=pedido_id)
    rua = _safe_text(request.POST.get("rua"))
    numero = _safe_text(request.POST.get("numero"))
    bairro = _safe_text(request.POST.get("bairro"))
    cidade = _safe_text(request.POST.get("cidade")) or "Rio Verde"
    estado = _safe_text(request.POST.get("estado")) or "GO"
    endereco_formatado = _safe_text(request.POST.get("endereco_formatado"))
    frete_gratis = _post_bool(request.POST.get("frete_gratis"))
    destination_result = _destination_result_from_values(request.POST)
    lote_quadra = _safe_text(request.POST.get("lote_quadra"))
    endereco = _build_order_address(rua, numero, bairro, cidade, estado, lote_quadra) or endereco_formatado
    common_fields = {
        "tipo_coleta": Pedido.TipoColeta.ENTREGA,
        "rua": rua,
        "numero_endereco": numero,
        "bairro": bairro,
        "cidade": cidade,
        "estado": estado,
        "endereco": endereco or pedido.endereco,
        "endereco_formatado": endereco_formatado or endereco or pedido.endereco_formatado,
        "complemento": _safe_text(request.POST.get("complemento")),
        "lote_quadra": lote_quadra,
        "ponto_referencia": _safe_text(request.POST.get("ponto_referencia")),
    }
    if not destination_result:
        for field, value in common_fields.items():
            setattr(pedido, field, value)
        update_fields = list(common_fields.keys())
        was_frete_gratis = pedido.frete_gratis
        pedido.frete_gratis = frete_gratis
        update_fields.append("frete_gratis")
        if frete_gratis:
            pedido.valor_frete = Decimal("0.00")
            update_fields.append("valor_frete")
        elif was_frete_gratis and pedido.distancia_km > 0:
            pedido.valor_frete, _ = _calcular_frete_por_distancia(pedido.distancia_km)
            update_fields.append("valor_frete")
        pedido.save(update_fields=update_fields)
        if "valor_frete" in update_fields:
            recalculate_order_totals(pedido)
        sync_movimentacao_caixa_pedido(pedido, request.user)
        sync_customer_from_order(pedido)
        payload = _pedido_modal_payload(pedido)
        payload["frete_recalculado"] = False
        return JsonResponse(payload)
    origin_result = _resolve_saved_origin_result()
    if not origin_result:
        return HttpResponseBadRequest("Configure a origem de entrega antes de recalcular.")
    duration_seconds, distance_meters = _fetch_route_summary(
        origin_result["lat"],
        origin_result["lng"],
        destination_result["lat"],
        destination_result["lng"],
    )
    if duration_seconds is None or distance_meters is None:
        return HttpResponseBadRequest("Nao foi possivel calcular a rota para o ponto de entrega.")
    distancia_km = Decimal(str(round(max(distance_meters / 1000.0, 0.0), 2)))
    valor_frete, _ = _calcular_frete_por_distancia(distancia_km)

    for field, value in common_fields.items():
        setattr(pedido, field, value)
    pedido.endereco_formatado = destination_result.get("label") or common_fields["endereco_formatado"]
    pedido.latitude = _to_decimal(destination_result.get("lat"))
    pedido.longitude = _to_decimal(destination_result.get("lng"))
    pedido.distancia_km = distancia_km
    pedido.frete_gratis = frete_gratis
    pedido.valor_frete = Decimal("0.00") if frete_gratis else valor_frete
    pedido.save(update_fields=[
        "rua",
        "numero_endereco",
        "bairro",
        "cidade",
        "estado",
        "endereco",
        "endereco_formatado",
        "latitude",
        "longitude",
        "complemento",
        "lote_quadra",
        "ponto_referencia",
        "tipo_coleta",
        "distancia_km",
        "valor_frete",
        "frete_gratis",
    ])
    recalculate_order_totals(pedido)
    sync_movimentacao_caixa_pedido(pedido, request.user)
    sync_customer_from_order(pedido)
    payload = _pedido_modal_payload(pedido)
    payload["frete_recalculado"] = True
    return JsonResponse(payload)


@staff_member_required(login_url="/admin/login/")
def ajustes_admin(request):
    ajustes_aba = (_safe_text(request.GET.get("aba")) or "geral").lower()
    if ajustes_aba not in {"geral", "frete", "google", "whatsapp", "pagamento", "ifood", "contabil", "usuarios", "api", "lista_impressao", "importacao"}:
        ajustes_aba = "geral"

    _ensure_default_user_groups()
    config = ConfiguracaoEntrega.get_solo()
    origem = _saved_origin_snapshot(config)
    origem_resolution = _resolve_saved_origin_result(config) or _blank_origin_result()
    faixa_rows = _serialize_faixas_for_form()
    feedback = None
    feedback_kind = "success"
    preview = None
    legacy_import_preview = None
    legacy_import_payload = ""
    legacy_import_default_fee = "10.00"
    google_maps_status = _google_maps_status()

    if request.GET.get("saved") == "1":
        feedback = "Ajustes atualizados."

    if request.method == "POST":
        action = _safe_text(request.POST.get("action"))
        action_tabs = {
            "save_geral": "geral",
            "create_data_fechada": "geral",
            "toggle_data_fechada": "geral",
            "save_frete": "frete",
            "test_frete": "frete",
            "save_google": "google",
            "save_whatsapp": "whatsapp",
            "save_pagamento": "pagamento",
            "save_ifood": "ifood",
            "create_terminal": "contabil",
            "update_terminal": "contabil",
            "create_categoria": "contabil",
            "update_categoria": "contabil",
            "create_banco": "contabil",
            "update_banco": "contabil",
            "create_categoria_conta": "contabil",
            "update_categoria_conta": "contabil",
            "save_contabil_bancos_pagamento": "contabil",
            "create_api_key": "api",
            "delete_api_key": "api",
            "create_user": "usuarios",
            "update_user": "usuarios",
            "create_group": "usuarios",
            "update_group": "usuarios",
            "delete_group": "usuarios",
            "preview_legacy_orders": "importacao",
            "import_legacy_orders": "importacao",
        }
        ajustes_aba = action_tabs.get(action, ajustes_aba)
        origem = {
            "endereco": _safe_text(request.POST.get("origem_endereco")) or origem["endereco"],
            "latitude": _safe_text(request.POST.get("origem_latitude")) or (_safe_text(origem["latitude"]) if origem["latitude"] is not None else ""),
            "longitude": _safe_text(request.POST.get("origem_longitude")) or (_safe_text(origem["longitude"]) if origem["longitude"] is not None else ""),
        }
        origem_resolution = (
            _origin_result_from_coordinates(origem["endereco"], origem["latitude"], origem["longitude"])
            if _safe_float(origem["latitude"]) is not None and _safe_float(origem["longitude"]) is not None
            else _blank_origin_result()
        )
        faixa_rows = _parse_faixa_rows(request.POST)

        if action == "save_geral":
            try:
                config.horario_abertura = _parse_optional_time(request.POST.get("horario_abertura"))
                config.horario_fechamento = _parse_optional_time(request.POST.get("horario_fechamento"))
                config.save()
                return redirect(f"{request.path}?saved=1&aba=geral")
            except ValueError as exc:
                feedback = str(exc)
                feedback_kind = "error"

        if action == "create_data_fechada":
            data_fechada = parse_date(_safe_text(request.POST.get("data_fechada")))
            if not data_fechada:
                feedback = "Informe uma data valida para fechar."
                feedback_kind = "error"
            else:
                DataFechada.objects.update_or_create(
                    data=data_fechada,
                    defaults={"motivo": _safe_text(request.POST.get("motivo_fechamento")), "ativo": True},
                )
                return redirect(f"{request.path}?saved=1&aba=geral")

        if action == "toggle_data_fechada":
            data_fechada = get_object_or_404(DataFechada, id=request.POST.get("data_fechada_id"))
            data_fechada.ativo = not data_fechada.ativo
            data_fechada.save(update_fields=["ativo", "atualizado_em"])
            return redirect(f"{request.path}?saved=1&aba=geral")

        if action == "save_frete":
            try:
                if _to_decimal(origem["latitude"]) is None or _to_decimal(origem["longitude"]) is None:
                    raise ValueError("Confirme a origem no mapa antes de salvar os ajustes.")
                config.origem_endereco = origem["endereco"]
                config.origem_latitude = _to_decimal(origem["latitude"])
                config.origem_longitude = _to_decimal(origem["longitude"])
                config.save()
                _save_faixa_rows(faixa_rows)
                return redirect(f"{request.path}?saved=1&aba=frete")
            except ValueError as exc:
                feedback = str(exc)
                feedback_kind = "error"

        if action == "save_google":
            config.google_maps_api_key = _safe_text(request.POST.get("google_maps_api_key"))
            config.google_maps_language = _safe_text(request.POST.get("google_maps_language")) or "pt-BR"
            config.google_maps_region = _safe_text(request.POST.get("google_maps_region")) or "BR"
            config.save()
            return redirect(f"{request.path}?saved=1&aba=google")

        if action == "save_whatsapp":
            numero = _normalize_whatsapp_number(request.POST.get("whatsapp_numero"))
            if numero and len(numero) < 12:
                feedback = "Informe o WhatsApp com DDI e DDD. Exemplo: 5564999999999."
                feedback_kind = "error"
            else:
                config.whatsapp_numero = numero
                config.save()
                return redirect(f"{request.path}?saved=1&aba=whatsapp")

        if action == "save_pagamento":
            config.pix_chave = _safe_text(request.POST.get("pix_chave"))
            config.save()
            return redirect(f"{request.path}?saved=1&aba=pagamento")

        if action == "save_ifood":
            try:
                config.taxa_ifood_percentual = _parse_percentual_0_100(
                    request.POST.get("taxa_ifood_percentual"),
                    "a taxa Ifood",
                )
                config.save()
                return redirect(f"{request.path}?saved=1&aba=ifood")
            except ValueError as exc:
                feedback = str(exc)
                feedback_kind = "error"

        if action == "create_terminal":
            try:
                nome = _safe_text(request.POST.get("terminal_nome"))
                if not nome:
                    raise ValueError("Informe o nome do terminal.")
                codigo_base = unicodedata.normalize("NFKD", nome).encode("ascii", "ignore").decode("ascii").lower()
                codigo_base = "-".join(part for part in codigo_base.replace("_", "-").split() if part)[:24] or "terminal"
                codigo = codigo_base
                suffix = 2
                while TerminalCaixa.objects.filter(codigo=codigo).exists():
                    codigo = f"{codigo_base}-{suffix}"[:30]
                    suffix += 1
                TerminalCaixa.objects.create(
                    nome=nome,
                    codigo=codigo,
                    ordem=(TerminalCaixa.objects.aggregate(max_ordem=Max("ordem")).get("max_ordem") or 0) + 10,
                )
                return redirect(f"{request.path}?saved=1&aba=contabil")
            except ValueError as exc:
                feedback = str(exc)
                feedback_kind = "error"

        if action == "create_categoria":
            try:
                nome = _safe_text(request.POST.get("categoria_nome"))
                if not nome:
                    raise ValueError("Informe o nome da categoria.")
                tipo_padrao = _safe_text(request.POST.get("tipo_padrao"))
                if tipo_padrao not in dict(CategoriaMovimentacaoCaixa.TipoPadrao.choices):
                    tipo_padrao = CategoriaMovimentacaoCaixa.TipoPadrao.AMBOS
                categoria, created = CategoriaMovimentacaoCaixa.objects.get_or_create(
                    nome=nome,
                    defaults={"tipo_padrao": tipo_padrao, "ativo": True},
                )
                if not created:
                    categoria.tipo_padrao = tipo_padrao
                    categoria.ativo = True
                    categoria.save(update_fields=["tipo_padrao", "ativo"])
                return redirect(f"{request.path}?saved=1&aba=contabil")
            except ValueError as exc:
                feedback = str(exc)
                feedback_kind = "error"

        if action == "update_terminal":
            try:
                terminal = get_object_or_404(TerminalCaixa, id=request.POST.get("terminal_id"))
                nome = _safe_text(request.POST.get("terminal_nome"))
                if not nome:
                    raise ValueError("Informe o nome do terminal.")
                terminal.nome = nome
                terminal.save(update_fields=["nome", "atualizado_em"])
                return redirect(f"{request.path}?saved=1&aba=contabil")
            except ValueError as exc:
                feedback = str(exc)
                feedback_kind = "error"

        if action == "update_categoria":
            try:
                categoria = get_object_or_404(CategoriaMovimentacaoCaixa, id=request.POST.get("categoria_id"))
                nome = _safe_text(request.POST.get("categoria_nome"))
                if not nome:
                    raise ValueError("Informe o nome da categoria.")
                if CategoriaMovimentacaoCaixa.objects.filter(nome__iexact=nome).exclude(id=categoria.id).exists():
                    raise ValueError("Ja existe uma categoria do caixa com este nome.")
                tipo_padrao = _safe_text(request.POST.get("tipo_padrao"))
                if tipo_padrao not in dict(CategoriaMovimentacaoCaixa.TipoPadrao.choices):
                    tipo_padrao = CategoriaMovimentacaoCaixa.TipoPadrao.AMBOS
                categoria.nome = nome
                categoria.tipo_padrao = tipo_padrao
                categoria.ativo = True
                categoria.save(update_fields=["nome", "tipo_padrao", "ativo"])
                return redirect(f"{request.path}?saved=1&aba=contabil")
            except ValueError as exc:
                feedback = str(exc)
                feedback_kind = "error"

        if action == "create_banco":
            try:
                nome = _safe_text(request.POST.get("banco_nome"))
                if not nome:
                    raise ValueError("Informe o nome do banco.")
                codigo_base = unicodedata.normalize("NFKD", nome).encode("ascii", "ignore").decode("ascii").lower()
                codigo_base = "-".join(part for part in codigo_base.replace("_", "-").split() if part)[:24] or "banco"
                codigo = codigo_base
                suffix = 2
                while BancoConta.objects.filter(codigo=codigo).exists():
                    codigo = f"{codigo_base}-{suffix}"[:30]
                    suffix += 1
                BancoConta.objects.create(
                    nome=nome,
                    codigo=codigo,
                    ordem=(BancoConta.objects.aggregate(max_ordem=Max("ordem")).get("max_ordem") or 0) + 10,
                )
                return redirect(f"{request.path}?saved=1&aba=contabil")
            except ValueError as exc:
                feedback = str(exc)
                feedback_kind = "error"

        if action == "update_banco":
            try:
                banco = get_object_or_404(BancoConta, id=request.POST.get("banco_id"))
                nome = _safe_text(request.POST.get("banco_nome"))
                if not nome:
                    raise ValueError("Informe o nome do banco.")
                banco.nome = nome
                banco.save(update_fields=["nome", "atualizado_em"])
                return redirect(f"{request.path}?saved=1&aba=contabil")
            except ValueError as exc:
                feedback = str(exc)
                feedback_kind = "error"

        if action == "create_categoria_conta":
            try:
                nome = _safe_text(request.POST.get("categoria_conta_nome"))
                if not nome:
                    raise ValueError("Informe o nome da categoria da conta.")
                tipo_padrao = _safe_text(request.POST.get("tipo_padrao_conta"))
                if tipo_padrao not in dict(CategoriaMovimentacaoConta.TipoPadrao.choices):
                    tipo_padrao = CategoriaMovimentacaoConta.TipoPadrao.AMBOS
                categoria, created = CategoriaMovimentacaoConta.objects.get_or_create(
                    nome=nome,
                    defaults={"tipo_padrao": tipo_padrao, "ativo": True},
                )
                if not created:
                    categoria.tipo_padrao = tipo_padrao
                    categoria.ativo = True
                    categoria.save(update_fields=["tipo_padrao", "ativo"])
                return redirect(f"{request.path}?saved=1&aba=contabil")
            except ValueError as exc:
                feedback = str(exc)
                feedback_kind = "error"

        if action == "update_categoria_conta":
            try:
                categoria = get_object_or_404(CategoriaMovimentacaoConta, id=request.POST.get("categoria_conta_id"))
                nome = _safe_text(request.POST.get("categoria_conta_nome"))
                if not nome:
                    raise ValueError("Informe o nome da categoria da conta.")
                if CategoriaMovimentacaoConta.objects.filter(nome__iexact=nome).exclude(id=categoria.id).exists():
                    raise ValueError("Ja existe uma categoria da conta com este nome.")
                tipo_padrao = _safe_text(request.POST.get("tipo_padrao_conta"))
                if tipo_padrao not in dict(CategoriaMovimentacaoConta.TipoPadrao.choices):
                    tipo_padrao = CategoriaMovimentacaoConta.TipoPadrao.AMBOS
                categoria.nome = nome
                categoria.tipo_padrao = tipo_padrao
                categoria.ativo = True
                categoria.save(update_fields=["nome", "tipo_padrao", "ativo"])
                return redirect(f"{request.path}?saved=1&aba=contabil")
            except ValueError as exc:
                feedback = str(exc)
                feedback_kind = "error"

        if action == "save_contabil_bancos_pagamento":
            banco_pix = BancoConta.objects.filter(id=request.POST.get("banco_pix"), ativo=True).first()
            banco_cartao = BancoConta.objects.filter(id=request.POST.get("banco_cartao"), ativo=True).first()
            config.banco_pix = banco_pix
            config.banco_cartao = banco_cartao
            config.save(update_fields=["banco_pix", "banco_cartao", "atualizado_em"])
            return redirect(f"{request.path}?saved=1&aba=contabil")

        if action == "create_api_key":
            if not _user_can_manage_order_payment(request.user):
                feedback = "Usuario sem permissao para criar chaves da API."
                feedback_kind = "error"
            else:
                nome = _safe_text(request.POST.get("nome"))
                if not nome:
                    feedback = "Informe um nome para a chave."
                    feedback_kind = "error"
                else:
                    _api_key, raw_key = PedidoApiKey.create_key(nome, request.user)
                    feedback = f"Chave criada. Copie agora: {raw_key}"
                    feedback_kind = "success"

        if action == "delete_api_key":
            if not _user_can_manage_order_payment(request.user):
                feedback = "Usuario sem permissao para excluir chaves da API."
                feedback_kind = "error"
            else:
                PedidoApiKey.objects.filter(id=request.POST.get("api_key_id")).delete()
                feedback = "Chave excluida."
                feedback_kind = "success"

        if action in {"create_user", "update_user", "create_group", "update_group", "delete_group"}:
            try:
                _handle_user_admin_action(request)
                return redirect(f"{request.path}?saved=1&aba=usuarios")
            except (PermissionError, ValueError) as exc:
                feedback = str(exc)
                feedback_kind = "error"

        if action == "preview_legacy_orders":
            legacy_import_default_fee = f"{legacy_money_decimal(request.POST.get('default_delivery_fee') or '10.00'):.2f}"
            upload = request.FILES.get("legacy_orders_file")
            if not upload:
                feedback = "Selecione o CSV limpo antes de revisar a importacao."
                feedback_kind = "error"
            elif upload.size > 2 * 1024 * 1024:
                feedback = "O arquivo deve ter ate 2 MB."
                feedback_kind = "error"
            else:
                try:
                    content = upload.read().decode("utf-8-sig")
                    default_fee = legacy_money_decimal(legacy_import_default_fee)
                    legacy_import_preview = build_legacy_import_preview(content, default_fee)
                    legacy_import_payload = signing.dumps(
                        {"content": content, "default_delivery_fee": legacy_import_default_fee},
                        compress=True,
                    )
                    feedback = "Revise a pre-importacao antes de aplicar."
                    feedback_kind = "success"
                except (UnicodeDecodeError, ValueError, csv.Error) as exc:
                    feedback = f"Nao foi possivel ler o CSV: {exc}"
                    feedback_kind = "error"

        if action == "import_legacy_orders":
            legacy_import_payload = _safe_text(request.POST.get("legacy_import_payload"))
            try:
                payload = signing.loads(legacy_import_payload, max_age=60 * 60)
                content = payload["content"]
                default_fee = legacy_money_decimal(payload.get("default_delivery_fee") or "10.00")
                result = import_clean_legacy_orders(content, default_fee)
                feedback = (
                    "Importacao concluida. "
                    f"Importados: {result.imported} | Pulados por duplicidade: {result.skipped} | Erros: {result.errors}"
                )
                feedback_kind = "success"
            except signing.BadSignature:
                feedback = "A pre-importacao expirou ou foi alterada. Envie o CSV novamente."
                feedback_kind = "error"
            except (ValueError, csv.Error) as exc:
                feedback = f"Nao foi possivel importar o CSV: {exc}"
                feedback_kind = "error"

        if action == "test_frete":
            destino_teste = _safe_text(request.POST.get("destino_teste"))
            if _safe_float(origem["latitude"]) is None or _safe_float(origem["longitude"]) is None:
                feedback = "Confirme a origem no mapa antes de calcular o teste."
                feedback_kind = "error"
            elif len(destino_teste) < 5 and not (_safe_text(request.POST.get("destino_teste_lat")) and _safe_text(request.POST.get("destino_teste_lng"))):
                feedback = "Confirme um destino de teste no mapa antes de calcular."
                feedback_kind = "error"
            else:
                origem_lat = origem_resolution["lat"]
                origem_lng = origem_resolution["lng"]
                destino_resolvido = _destination_result_from_values(request.POST)
                if not destino_resolvido:
                    feedback = "Não foi possível localizar o destino de teste informado."
                    feedback_kind = "error"
                else:
                    duration_seconds, distance_meters = _fetch_route_summary(
                        origem_lat,
                        origem_lng,
                        destino_resolvido["lat"],
                        destino_resolvido["lng"],
                    )
                    if duration_seconds is None or distance_meters is None:
                        feedback = "Não foi possível calcular a rota viaria para esse destino."
                        feedback_kind = "error"
                    else:
                        distance_km = round(max(distance_meters / 1000.0, 0.0), 2)
                        frete_valor, faixa_aplicada = _calcular_frete_por_distancia(
                            distance_km,
                            faixas=_preview_faixas_from_rows(faixa_rows),
                        )
                        preview = {
                            "destino_digitado": destino_teste,
                            "origem_endereco": origem_resolution["label"],
                            "origem_lat": origem_lat,
                            "origem_lng": origem_lng,
                            "origem_mode": origem_resolution["mode"],
                            "origem_precision": origem_resolution["precision"],
                            "origem_precision_label": origem_resolution["precision_label"],
                            "origem_tipo": origem_resolution["type"],
                            "destino_label": destino_resolvido.get("label") or destino_teste,
                            "destino_lat": destino_resolvido["lat"],
                            "destino_lng": destino_resolvido["lng"],
                            "destino_tipo": destino_resolvido.get("type") or "desconhecido",
                            "destino_precision": destino_resolvido.get("precision") or "pending",
                            "destino_precision_label": destino_resolvido.get("precision_label") or _address_precision_label("pending"),
                            "distance_km": distance_km,
                            "duration_minutes": int(math.ceil(duration_seconds / 60.0)),
                            "frete_valor": frete_valor,
                            "faixa_label": _format_faixa_label(faixa_aplicada),
                            "origem_nao_salva": not _origin_matches_saved_config(origem, config),
                        }

    ultimo_pedido = Pedido.objects.prefetch_related("itens").first()
    ultimo_pedido_auditoria = None
    if ultimo_pedido:
        subtotal_itens = ultimo_pedido.itens.aggregate(total=Sum("subtotal")).get("total") or Decimal("0.00")
        frete_esperado, faixa_atual = _calcular_frete_por_distancia(ultimo_pedido.distancia_km)
        ultimo_pedido_auditoria = {
            "pedido": ultimo_pedido,
            "subtotal_itens": subtotal_itens,
            "frete_esperado": frete_esperado,
            "faixa_label": _format_faixa_label(faixa_atual),
            "diferenca_frete": ultimo_pedido.valor_frete - frete_esperado,
            "total_recalculado": subtotal_itens + ultimo_pedido.valor_frete,
        }

    usuarios_admin_rows = _serialize_users_for_admin()
    usuarios_classes = _serialize_groups_for_admin()
    usuarios_metrics = {
        "total": len(usuarios_admin_rows),
        "active": sum(1 for usuario in usuarios_admin_rows if usuario["is_active"]),
        "inactive": sum(1 for usuario in usuarios_admin_rows if not usuario["is_active"]),
        "superusers": sum(1 for usuario in usuarios_admin_rows if usuario["is_superuser"]),
        "classes": usuarios_classes.count(),
    }

    return render(
        request,
        "pedidos/ajustes_admin.html",
        {
            "active": "ajustes",
            "ajustes_aba": ajustes_aba,
            "pedidos_badge": Pedido.objects.exclude(
                status__in=[Pedido.Status.RASCUNHO, Pedido.Status.FINALIZADO, Pedido.Status.CANCELADO]
            ).count(),
            "origem": origem,
            "origem_resolution": origem_resolution,
            "faixa_rows": faixa_rows,
            "feedback": feedback,
            "feedback_kind": feedback_kind,
            "preview": preview,
            "google_maps_status": google_maps_status,
            "whatsapp_numero": config.whatsapp_numero,
            "pix_chave": config.pix_chave,
            "taxa_ifood_percentual": f"{config.taxa_ifood_percentual:.2f}",
            "terminais": TerminalCaixa.objects.all().order_by("ordem", "nome"),
            "categorias": CategoriaMovimentacaoCaixa.objects.all().order_by("nome"),
            "bancos": BancoConta.objects.all().order_by("ordem", "nome"),
            "categorias_conta": CategoriaMovimentacaoConta.objects.all().order_by("nome"),
            "banco_pix_id": config.banco_pix_id,
            "banco_cartao_id": config.banco_cartao_id,
            "horario_abertura": config.horario_abertura.strftime("%H:%M") if config.horario_abertura else "",
            "horario_fechamento": config.horario_fechamento.strftime("%H:%M") if config.horario_fechamento else "",
            "datas_fechadas": DataFechada.objects.all().order_by("-data")[:20],
            "ultimo_pedido_auditoria": ultimo_pedido_auditoria,
            "usuarios_admin_rows": usuarios_admin_rows,
            "usuarios_classes": usuarios_classes,
            "usuarios_metrics": usuarios_metrics,
            "can_manage_users": request.user.is_superuser,
            "pedido_api_keys": _serialize_pedido_api_keys(),
            "can_manage_api_keys": _user_can_manage_order_payment(request.user),
            "lista_impressao_rows": _serialize_lista_impressao_admin(),
            "legacy_import_preview": legacy_import_preview,
            "legacy_import_payload": legacy_import_payload,
            "legacy_import_default_fee": legacy_import_default_fee,
        },
    )


@staff_member_required(login_url="/admin/login/")
def adicionais_admin(request):
    adicionais = Adicional.objects.all()
    adicional_edicao = None
    form = AdicionalForm()

    edit_id = request.GET.get("edit")
    if edit_id:
        adicional_edicao = get_object_or_404(Adicional, id=edit_id)
        form = AdicionalForm(instance=adicional_edicao)

    return render(
        request,
        "pedidos/adicionais_gestao.html",
        {
            "adicionais": adicionais,
            "form": form,
            "adicional_edicao": adicional_edicao,
            "form_modal_open": bool(adicional_edicao),
        },
    )


@staff_member_required(login_url="/admin/login/")
def outros_admin(request):
    return render(
        request,
        "pedidos/modulo_admin_placeholder.html",
        {
            "active": "outros",
            "titulo": "Outros",
            "descricao": "Area dedicada para itens e configuracoes complementares da operacao.",
        },
    )


@staff_member_required(login_url="/admin/login/")
def cupons_admin(request):
    feedback = ""
    feedback_kind = ""
    if request.method == "POST":
        action = _safe_text(request.POST.get("action"))
        cupom = Cupom.objects.filter(pk=request.POST.get("cupom_id")).first() if request.POST.get("cupom_id") else None
        if action in {"save", "toggle"}:
            cupom = cupom or Cupom()
            if action == "toggle":
                cupom.ativo = not cupom.ativo
                cupom.save(update_fields=["ativo", "atualizado_em"])
                return redirect("pedidos:cupons_admin")
            cupom.codigo = _normalize_coupon_code(request.POST.get("codigo"))
            cupom.descricao = _safe_text(request.POST.get("descricao"))
            cupom.tipo_desconto = request.POST.get("tipo_desconto") if request.POST.get("tipo_desconto") in dict(Cupom.TipoDesconto.choices) else Cupom.TipoDesconto.VALOR_FIXO
            cupom.valor = _money_decimal(request.POST.get("valor"))
            cupom.valor_minimo_pedido = _money_decimal(request.POST.get("valor_minimo_pedido"))
            uso_maximo = _safe_text(request.POST.get("uso_maximo_total"))
            cupom.uso_maximo_total = int(uso_maximo) if uso_maximo.isdigit() else None
            cupom.ativo = request.POST.get("ativo") == "on"
            if not cupom.codigo or cupom.valor <= 0:
                feedback = "Informe codigo e valor do desconto."
                feedback_kind = "error"
            else:
                try:
                    cupom.save()
                    return redirect("pedidos:cupons_admin")
                except Exception:
                    feedback = "Nao foi possivel salvar. Verifique se o codigo ja existe."
                    feedback_kind = "error"
    cupons = Cupom.objects.all().annotate(usos=Count("pedidos"))
    return render(request, "pedidos/cupons_admin.html", {"active": "cupons", "cupons": cupons, "feedback": feedback, "feedback_kind": feedback_kind})


@staff_member_required(login_url="/admin/login/")
def gestao_pratos(request):
    pratos = Prato.objects.filter(exclusivo_prato_pronto=False)
    prato_edicao = None
    form = PratoForm()

    edit_id = request.GET.get("edit")
    if edit_id:
        prato_edicao = get_object_or_404(
            Prato,
            id=edit_id,
            exclusivo_prato_pronto=False,
        )
        form = PratoForm(instance=prato_edicao)
    open_new_modal = request.GET.get("new") == "1" and not prato_edicao

    return render(
        request,
        "pedidos/pratos_gestao.html",
        {
            "pratos": pratos,
            "form": form,
            "prato_edicao": prato_edicao,
            "form_modal_open": bool(prato_edicao) or open_new_modal,
        },
    )


@staff_member_required(login_url="/admin/login/")
@require_POST
def salvar_prato(request):
    prato_id = request.POST.get("prato_id")
    prato = (
        get_object_or_404(Prato, id=prato_id, exclusivo_prato_pronto=False)
        if prato_id
        else None
    )
    form = PratoForm(request.POST, request.FILES, instance=prato)
    if form.is_valid():
        form.save()
        return redirect("pedidos:gestao_pratos")

    pratos = Prato.objects.filter(exclusivo_prato_pronto=False)
    return render(
        request,
        "pedidos/pratos_gestao.html",
        {
            "pratos": pratos,
            "form": form,
            "prato_edicao": prato,
            "form_modal_open": True,
        },
        status=400,
    )


@staff_member_required(login_url="/admin/login/")
@require_POST
def alternar_prato(request, prato_id):
    prato = get_object_or_404(
        Prato,
        id=prato_id,
        exclusivo_prato_pronto=False,
    )
    prato.ativo = not prato.ativo
    prato.save(update_fields=["ativo"])
    return redirect("pedidos:gestao_pratos")


def _catalog_action_response(request, redirect_name):
    redirect_url = reverse(redirect_name)
    if request.headers.get("x-requested-with") == "XMLHttpRequest":
        return JsonResponse({"ok": True, "redirect_url": redirect_url})
    return redirect(redirect_name)


@staff_member_required(login_url="/admin/login/")
@require_POST
def duplicar_prato(request, prato_id):
    prato = get_object_or_404(
        Prato,
        id=prato_id,
        exclusivo_prato_pronto=False,
    )
    Prato.objects.create(
        nome=f"Copia de {prato.nome}"[:120],
        descricao=prato.descricao,
        variacoes=prato.variacoes,
        imagem=prato.imagem,
        preco=prato.preco,
        preco_balcao=prato.preco_balcao,
        preco_site=prato.preco_site,
        preco_ifood=prato.preco_ifood,
        ativo=False,
        dias_disponiveis=prato.dias_disponiveis,
    )
    return _catalog_action_response(request, "pedidos:gestao_pratos")


@staff_member_required(login_url="/admin/login/")
@require_POST
def excluir_prato(request, prato_id):
    prato = get_object_or_404(
        Prato,
        id=prato_id,
        exclusivo_prato_pronto=False,
    )
    prato.delete()
    return _catalog_action_response(request, "pedidos:gestao_pratos")


def _delete_catalog_image(model, object_id):
    item = get_object_or_404(model, id=object_id)
    if not item.imagem:
        return JsonResponse(
            {
                "ok": True,
                "message": "Item sem imagem.",
                "placeholder": settings.STATIC_URL + "img/placeholder-prato.svg",
            }
        )

    item.imagem.delete(save=False)
    item.imagem = None
    item.save(update_fields=["imagem"])
    return JsonResponse(
        {
            "ok": True,
            "message": "Imagem excluida.",
            "placeholder": settings.STATIC_URL + "img/placeholder-prato.svg",
        }
    )


@staff_member_required(login_url="/admin/login/")
@require_POST
def excluir_imagem_prato(request, prato_id):
    prato = get_object_or_404(
        Prato,
        id=prato_id,
        exclusivo_prato_pronto=False,
    )
    return _delete_catalog_image(Prato, prato.id)


@staff_member_required(login_url="/admin/login/")
def gestao_bebidas(request):
    bebidas = Bebida.objects.all()
    bebida_edicao = None
    form = BebidaForm()

    edit_id = request.GET.get("edit")
    if edit_id:
        bebida_edicao = get_object_or_404(Bebida, id=edit_id)
        form = BebidaForm(instance=bebida_edicao)

    return render(
        request,
        "pedidos/bebidas_gestao.html",
        {
            "bebidas": bebidas,
            "form": form,
            "bebida_edicao": bebida_edicao,
            "form_modal_open": bool(bebida_edicao),
        },
    )


@staff_member_required(login_url="/admin/login/")
@require_POST
def salvar_bebida(request):
    bebida_id = request.POST.get("bebida_id")
    bebida = get_object_or_404(Bebida, id=bebida_id) if bebida_id else None
    form = BebidaForm(request.POST, request.FILES, instance=bebida)
    if form.is_valid():
        form.save()
        return redirect("pedidos:gestao_bebidas")

    bebidas = Bebida.objects.all()
    return render(
        request,
        "pedidos/bebidas_gestao.html",
        {
            "bebidas": bebidas,
            "form": form,
            "bebida_edicao": bebida,
            "form_modal_open": True,
        },
        status=400,
    )


@staff_member_required(login_url="/admin/login/")
@require_POST
def alternar_bebida(request, bebida_id):
    bebida = get_object_or_404(Bebida, id=bebida_id)
    bebida.ativo = not bebida.ativo
    bebida.save(update_fields=["ativo"])
    return redirect("pedidos:gestao_bebidas")


@staff_member_required(login_url="/admin/login/")
@require_POST
def duplicar_bebida(request, bebida_id):
    bebida = get_object_or_404(Bebida, id=bebida_id)
    Bebida.objects.create(
        nome=f"Copia de {bebida.nome}"[:120],
        descricao=bebida.descricao,
        imagem=bebida.imagem,
        preco=bebida.preco,
        preco_balcao=bebida.preco_balcao,
        preco_site=bebida.preco_site,
        preco_ifood=bebida.preco_ifood,
        ativo=False,
        ordem=bebida.ordem,
    )
    return _catalog_action_response(request, "pedidos:gestao_bebidas")


@staff_member_required(login_url="/admin/login/")
@require_POST
def excluir_bebida(request, bebida_id):
    bebida = get_object_or_404(Bebida, id=bebida_id)
    bebida.delete()
    return _catalog_action_response(request, "pedidos:gestao_bebidas")


@staff_member_required(login_url="/admin/login/")
@require_POST
def excluir_imagem_bebida(request, bebida_id):
    return _delete_catalog_image(Bebida, bebida_id)


@staff_member_required(login_url="/admin/login/")
@require_POST
def salvar_adicional(request):
    adicional_id = request.POST.get("adicional_id")
    adicional = get_object_or_404(Adicional, id=adicional_id) if adicional_id else None
    form = AdicionalForm(request.POST, request.FILES, instance=adicional)
    if form.is_valid():
        form.save()
        return redirect("pedidos:adicionais_admin")

    adicionais = Adicional.objects.all()
    return render(
        request,
        "pedidos/adicionais_gestao.html",
        {
            "adicionais": adicionais,
            "form": form,
            "adicional_edicao": adicional,
            "form_modal_open": True,
        },
        status=400,
    )


@staff_member_required(login_url="/admin/login/")
@require_POST
def alternar_adicional(request, adicional_id):
    adicional = get_object_or_404(Adicional, id=adicional_id)
    adicional.ativo = not adicional.ativo
    adicional.save(update_fields=["ativo"])
    return redirect("pedidos:adicionais_admin")


@staff_member_required(login_url="/admin/login/")
@require_POST
def duplicar_adicional(request, adicional_id):
    adicional = get_object_or_404(Adicional, id=adicional_id)
    Adicional.objects.create(
        nome=f"Copia de {adicional.nome}"[:120],
        descricao=adicional.descricao,
        imagem=adicional.imagem,
        preco=adicional.preco,
        preco_balcao=adicional.preco_balcao,
        preco_site=adicional.preco_site,
        preco_ifood=adicional.preco_ifood,
        ativo=False,
        ordem=adicional.ordem,
    )
    return _catalog_action_response(request, "pedidos:adicionais_admin")


@staff_member_required(login_url="/admin/login/")
@require_POST
def excluir_adicional(request, adicional_id):
    adicional = get_object_or_404(Adicional, id=adicional_id)
    adicional.delete()
    return _catalog_action_response(request, "pedidos:adicionais_admin")


@staff_member_required(login_url="/admin/login/")
@require_POST
def excluir_imagem_adicional(request, adicional_id):
    return _delete_catalog_image(Adicional, adicional_id)


@require_POST
def atualizar_status_pedido(request, pedido_id):
    pedido = get_object_or_404(Pedido, id=pedido_id)
    status = request.POST.get("status")
    if pedido.status == Pedido.Status.AGUARDANDO_APROVACAO and status == Pedido.Status.NOVO:
        status = Pedido.Status.EM_PREPARO
    if pedido.tipo_coleta == Pedido.TipoColeta.RETIRADA and status == Pedido.Status.SAIU_ENTREGA:
        status = Pedido.Status.FINALIZADO
    if status not in dict(Pedido.Status.choices):
        return HttpResponseBadRequest("Status invalido.")
    pedido.status = status
    update_fields = ["status"]
    if status == Pedido.Status.PAGAMENTO_RECEBIDO and not pedido.pagamento_recebido_em:
        pedido.pagamento_recebido_em = timezone.now()
        update_fields.append("pagamento_recebido_em")
    try:
        pedido.save(update_fields=update_fields)
    except ValueError as exc:
        return HttpResponseBadRequest(str(exc))
    sync_movimentacao_caixa_pedido(pedido, request.user if getattr(request, "user", None) and request.user.is_authenticated else None)
    if request.headers.get("x-requested-with") == "XMLHttpRequest":
        return JsonResponse({"ok": True, "status": pedido.status_label_contextual})
    return redirect("pedidos:cozinha_pedidos")


@staff_member_required(login_url="/admin/login/")
@require_POST
def alternar_entregador_pedido(request, pedido_id):
    pedido = get_object_or_404(Pedido, id=pedido_id)
    pedido.entregador_solicitado = not pedido.entregador_solicitado
    pedido.save(update_fields=["entregador_solicitado"])
    if request.headers.get("x-requested-with") == "XMLHttpRequest":
        return JsonResponse({"ok": True, "entregador_solicitado": pedido.entregador_solicitado})
    return redirect("pedidos:cozinha_pedidos")


@staff_member_required(login_url="/admin/login/")
@require_POST
def alternar_pagamento_recebido_pedido(request, pedido_id):
    pedido = get_object_or_404(Pedido, id=pedido_id)
    value = _safe_text(request.POST.get("pagamento_recebido")).lower()
    if value in {"sim", "true", "1"}:
        should_mark_paid = True
    elif value in {"nao", "false", "0"}:
        should_mark_paid = False
    else:
        should_mark_paid = not pedido.pagamento_recebido
    pedido.pagamento_recebido_em = timezone.now() if should_mark_paid else None
    pedido.save(update_fields=["pagamento_recebido_em"])
    payload = {
        "ok": True,
        "pagamento_recebido": pedido.pagamento_recebido,
        "pagamento_recebido_em": _format_local_datetime(pedido.pagamento_recebido_em, "%H:%M") if pedido.pagamento_recebido_em else "",
    }
    if request.headers.get("x-requested-with") == "XMLHttpRequest":
        return JsonResponse(payload)
    return redirect("pedidos:cozinha_pedidos")


@staff_member_required(login_url="/admin/login/")
@require_GET
def api_pedidos_cozinha(request):
    pedidos = Pedido.objects.prefetch_related("itens").all()[:20]
    data = []
    for pedido in pedidos:
        data.append(
            {
                "id": pedido.id,
                "numero": pedido.numero,
                "cliente": pedido.nome_cliente,
                "telefone": pedido.telefone,
                "rua": pedido.rua,
                "numero_endereco": pedido.numero_endereco,
                "bairro": pedido.bairro,
                "cidade": pedido.cidade,
                "estado": pedido.estado,
                "endereco": pedido.endereco,
                "endereco_formatado": pedido.endereco_formatado,
                "lote_quadra": pedido.lote_quadra,
                "complemento": pedido.complemento,
                "ponto_referencia": pedido.ponto_referencia,
                "latitude": float(pedido.latitude) if pedido.latitude is not None else None,
                "longitude": float(pedido.longitude) if pedido.longitude is not None else None,
                "google_maps_route_url": pedido.google_maps_route_url,
                "has_coordinates": pedido.has_coordinates,
                "enviar_talheres": pedido.enviar_talheres,
                "distancia_km": float(pedido.distancia_km) if pedido.distancia_km is not None else 0,
                "valor_frete": f"R$ {pedido.valor_frete:.2f}".replace(".", ","),
                "observacao_geral": pedido.observacao_geral,
                "status": pedido.status,
                "status_label": pedido.status_label_contextual,
                "tipo_coleta": pedido.tipo_coleta,
                "turno": pedido.turno.codigo if pedido.turno_id else TurnoAtendimento.Codigo.PRINCIPAL,
                "turno_label": turno_label_pedido(pedido),
                "orientacao_turno": pedido.orientacao_turno_snapshot,
                "stage_labels": pedido.stage_labels,
                "item_type_counts": pedido.item_type_counts,
                "horario": _format_local_datetime(pedido.criado_em, "%H:%M"),
                "total": f"R$ {pedido.total:.2f}".replace(".", ","),
                "itens": [
                    {
                        "nome": item.nome_prato_snapshot,
                        "variacao": item.variacao_nome_snapshot,
                        "quantidade": item.quantidade,
                        "observacao": item.observacao,
                        "subtotal": f"R$ {item.subtotal:.2f}".replace(".", ","),
                    }
                    for item in pedido.itens.all()
                ],
            }
        )
    return JsonResponse({"pedidos": data, "status_choices": list(Pedido.Status.choices)})

